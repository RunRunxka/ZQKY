"""V00 · B3-A5 规模（200×100）、CSV 全链、超限 TABLE_TOO_LARGE。

覆盖：
  1. `read_score_sheet` 直读自造 200 行 × 100 叶（物理 201×102）XLSX：数秒级、不截断；
  2. 真装配全链：200 人次 × 100 叶上传 → 确认 → 矩阵分页（记录耗时）；
  3. CSV 输入全链（0/missing/absent 三态 + 矩阵）；
  4. 超限：XLSX/CSV 行 2001、列 601 → 明确 `TABLE_TOO_LARGE`（上传 API 与直读两条路径）。

运行（在 apps/api 下）：
  .venv/Scripts/python.exe -X utf8 <abs>/p24_scale_csv.py
"""

from __future__ import annotations

import csv
import io
import sys
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import b3lib as B  # noqa: E402

ROWS_200 = 200
LEAVES_100 = 100


def main() -> int:
    from openpyxl import Workbook

    from app.core.exceptions import AppError
    from app.services.tabular import read_score_sheet

    verdict = B.Verdict("p24_scale_csv")

    # ======================================================== 1. 直读 200×100 不截断
    wide = Path(__file__).parent / "_scale_200x100.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "成绩"
    sheet.append(["学号", "姓名"] + [f"Q{i}" for i in range(1, LEAVES_100 + 1)])
    for index in range(1, ROWS_200 + 1):
        sheet.append([f"{index:04d}", f"S{index}"] + [1] * LEAVES_100)
    workbook.save(wide)
    started = time.perf_counter()
    sheets = read_score_sheet(wide.read_bytes())
    elapsed = time.perf_counter() - started
    verdict.check("read_score_sheet 200×100 在 10s 内", elapsed < 10.0, {"seconds": round(elapsed, 3)})
    verdict.note(f"read_score_sheet 200 行×100 叶（物理 201×102）耗时 {elapsed:.3f}s")
    grid = sheets[0]
    verdict.expect("物理 max_row=201", grid.max_row, 201)
    verdict.expect("物理 max_column=102", grid.max_column, 102)
    last = grid.rows[-1][-1]
    verdict.expect("末行末列未被截断", (last.row, last.column, last.text), (201, 102, "1"))
    verdict.expect("首行首列（表头）", grid.rows[0][0].text, "学号")
    verdict.expect("行数完整", len(grid.rows), 201)

    # ======================================================== 2. 200×100 全链
    with B.Harness(tag="scale") as harness:
        leaves = tuple((f"Q{i}", 100) for i in range(1, LEAVES_100 + 1))
        paper = harness.seed_confirmed_paper(tag="scale", leaves=leaves)
        klass = harness.create_class(code="C-scale")
        joined = B.days_before(30)
        participants = []
        for index in range(1, ROWS_200 + 1):
            student = harness.create_student(
                name=f"S{index}", student_no=f"{index:04d}", class_id=klass["id"], joined_on=joined
            )
            participants.append(harness.participant(student["id"], klass["id"]))
        detail = harness.create_assessment(
            paper=paper,
            class_id=klass["id"],
            participants=participants,
            submission_id="as-scale",
            title="规模施测",
        )
        assessment_id = detail["assessment"]["assessmentId"]
        path = Path(harness.settings.data_dir) / "scale-200x100.xlsx"
        B.write_xlsx(
            path,
            ["学号", "姓名"] + [f"Q{i}" for i in range(1, LEAVES_100 + 1)],
            [
                [f"{index:04d}", f"S{index}"] + [(index % 3) * 0.5 for _ in range(LEAVES_100)]
                for index in range(1, ROWS_200 + 1)
            ],
        )
        upload_started = time.perf_counter()
        upload = harness.upload_scores(assessment_id, path)
        upload_seconds = time.perf_counter() - upload_started
        verdict.expect("规模上传 201", upload.status_code, 201)
        view = upload.json()
        verdict.expect("规模预览 resolvedRowCount=200", view["resolvedRowCount"], 200)
        verdict.expect("规模预览 missingCellCount=0", view["missingCellCount"], 0)
        verdict.expect("规模预览 rowCount=200", view["rowCount"], 200)

        confirm_started = time.perf_counter()
        confirm = harness.confirm_import(
            view["importId"],
            {
                "expectedImportRevision": view["revision"],
                "expectedAssessmentRevision": detail["assessment"]["revision"],
                "baseScoreRevisionId": None,
                "previewVersion": view["previewVersion"],
                "submissionId": "scale-confirm",
                "absences": [],
            },
        )
        confirm_seconds = time.perf_counter() - confirm_started
        verdict.expect("规模确认 200", confirm.status_code, 200)
        revision_id = confirm.json()["revisionId"]

        matrix_started = time.perf_counter()
        page = harness.score_matrix(revision_id, offset=0, limit=50).json()
        matrix_seconds = time.perf_counter() - matrix_started
        verdict.expect("规模矩阵 total=200", page["total"], 200)
        verdict.expect("规模矩阵 items=100", len(page["items"]), 100)
        verdict.expect("规模矩阵首屏 50 行", len(page["rows"]), 50)
        verdict.check(
            "规模矩阵每人 totalUnits 非空（全 recorded）",
            all(row["participant"]["totalUnits"] is not None for row in page["rows"]),
            page["rows"][0]["participant"],
        )
        tail = harness.score_matrix(revision_id, offset=150, limit=50).json()
        verdict.expect("尾页 50 行", len(tail["rows"]), 50)
        names = {
            row["participant"]["name"]
            for pagex in (
                page,
                harness.score_matrix(revision_id, offset=50, limit=50).json(),
                harness.score_matrix(revision_id, offset=100, limit=50).json(),
                tail,
            )
            for row in pagex["rows"]
        }
        verdict.expect("四页无重复且覆盖全部 200 人次", len(names), 200)
        verdict.note(
            f"200×100 全链耗时：上传解析 {upload_seconds:.3f}s、确认 {confirm_seconds:.3f}s、"
            f"矩阵首屏 {matrix_seconds:.3f}s"
        )
        verdict.check("上传解析 < 30s", upload_seconds < 30.0, upload_seconds)
        verdict.check("确认 < 30s", confirm_seconds < 30.0, confirm_seconds)

    # ======================================================== 3. CSV 输入全链
    with B.Harness(tag="csv") as harness:
        scene = harness.sample_scene(
            tag="csv",
            leaves=(("Q1", 200), ("Q2", 300), ("Q3", 500)),
            students=(("A", "0001", "present"), ("B", "0002", "present"), ("C", "0003", "absent")),
        )
        assessment_id = scene["assessment"]["assessmentId"]
        csv_path = B.write_csv(
            Path(harness.settings.data_dir) / "scores.csv",
            ("学号", "姓名", "Q1", "Q2", "Q3"),
            (("0001", "A", 2, 2, 5), ("0002", "B", 2, 3, ""), ("0003", "C", "", "", "")),
        )
        upload = harness.upload_scores(assessment_id, csv_path, media_type="text/csv")
        verdict.expect("CSV 上传 201", upload.status_code, 201)
        view = upload.json()
        verdict.expect("CSV 工作表名=CSV", (view["mapping"] or {}).get("workSheet"), "CSV")
        verdict.expect("CSV 自动映射 Q1..Q3", len((view["mapping"] or {}).get("itemColumns", [])), 3)
        verdict.expect("CSV missingCellCount=1", view["missingCellCount"], 1)
        rows = harness.import_rows(view["importId"]).json()["items"]
        verdict.expect("CSV 物理行号=文件行号（2/3/4）", [row["rowNo"] for row in rows], [2, 3, 4])
        confirm = harness.confirm_import(
            view["importId"],
            {
                "expectedImportRevision": view["revision"],
                "expectedAssessmentRevision": harness.assessment(assessment_id)["assessment"]["revision"],
                "baseScoreRevisionId": None,
                "previewVersion": view["previewVersion"],
                "submissionId": "csv-confirm",
                "absences": [{"classId": scene["class"]["id"], "participantIds": [scene["byName"]["C"]["participantId"]]}],
                "missing": {"participantIds": [scene["byName"]["B"]["participantId"]], "cellCount": 1},
            },
        )
        verdict.expect("CSV 确认 200", confirm.status_code, 200)
        page = harness.score_matrix(confirm.json()["revisionId"], limit=200).json()
        table = B.matrix_map(page)
        verdict.expect("CSV 矩阵 A.totalUnits=900", table["A"]["participant"]["totalUnits"], 900)
        verdict.expect("CSV 矩阵 B 的 Q3=missing", table["B"]["cells"][scene["paper"]["itemIds"][2]]["status"], "missing")
        verdict.expect("CSV 矩阵 C 全 absent", {table["C"]["cells"][i]["status"] for i in scene["paper"]["itemIds"]}, {"absent"})
        verdict.expect("CSV 矩阵 missingCellCount=1", page["missingCellCount"], 1)

    # ======================================================== 4. 超限 TABLE_TOO_LARGE
    # 4a. 直读：行 2001
    big_rows = Path(__file__).parent / "_too_many_rows.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["a", "b"])
    for index in range(2001):
        sheet.append([index, index])
    workbook.save(big_rows)
    try:
        read_score_sheet(big_rows.read_bytes())
        verdict.check("直读 2002 行 XLSX → TABLE_TOO_LARGE", False, "未报错")
    except AppError as exc:
        verdict.expect("直读 2002 行 XLSX → TABLE_TOO_LARGE", exc.code, "TABLE_TOO_LARGE")

    # 4b. 直读：列 601
    big_cols = Path(__file__).parent / "_too_many_cols.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([f"c{i}" for i in range(601)])
    sheet.append([1] * 601)
    workbook.save(big_cols)
    try:
        read_score_sheet(big_cols.read_bytes())
        verdict.check("直读 601 列 XLSX → TABLE_TOO_LARGE", False, "未报错")
    except AppError as exc:
        verdict.expect("直读 601 列 XLSX → TABLE_TOO_LARGE", exc.code, "TABLE_TOO_LARGE")

    # 4c. 直读：CSV 行 2001 / 列 601
    csv_rows = io.StringIO()
    writer = csv.writer(csv_rows)
    writer.writerow(["a", "b"])
    for index in range(2001):
        writer.writerow([index, index])
    try:
        read_score_sheet(csv_rows.getvalue().encode("utf-8"))
        verdict.check("直读 2002 行 CSV → TABLE_TOO_LARGE", False, "未报错")
    except AppError as exc:
        verdict.expect("直读 2002 行 CSV → TABLE_TOO_LARGE", exc.code, "TABLE_TOO_LARGE")

    csv_cols = io.StringIO()
    writer = csv.writer(csv_cols)
    writer.writerow([f"c{i}" for i in range(601)])
    writer.writerow([1] * 601)
    try:
        read_score_sheet(csv_cols.getvalue().encode("utf-8"))
        verdict.check("直读 601 列 CSV → TABLE_TOO_LARGE", False, "未报错")
    except AppError as exc:
        verdict.expect("直读 601 列 CSV → TABLE_TOO_LARGE", exc.code, "TABLE_TOO_LARGE")

    # 4d. 上传 API：2002 行 CSV → 422 TABLE_TOO_LARGE（不落批次）
    with B.Harness(tag="limit") as harness:
        scene = harness.sample_scene(tag="limit")
        assessment_id = scene["assessment"]["assessmentId"]
        before = harness.count("score_imports")
        big_csv = Path(harness.settings.data_dir) / "rows-too-many.csv"
        with big_csv.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["学号", "姓名", "Q1"])
            for index in range(2001):
                writer.writerow([f"{index:04d}", f"S{index}", 1])
        upload = harness.upload_scores(
            assessment_id, big_csv, file_name="rows-too-many.csv", media_type="text/csv"
        )
        verdict.expect("上传 2002 行 CSV → 422", upload.status_code, 422)
        verdict.expect("上传 2002 行 CSV 错误码", upload.json().get("code"), "TABLE_TOO_LARGE")
        verdict.expect("超限上传不落批次", harness.count("score_imports"), before)
        verdict.note(f"上传超限响应：{upload.json().get('message')}")

    return verdict.finish(path=HERE / "p24_scale_csv.json")


if __name__ == "__main__":
    sys.exit(main())
