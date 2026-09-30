"""V00 自建 DOCX 样本（供 V6 解析 / V7 渲染探针共用；不依赖实现者的样本）。

样本含：前导材料段、含 vMerge 纵向合并 + gridSpan 横向合并 + 表头行的表格、
行内 OMML、独立 OMML、内嵌图片、未知对象（w:object）、空段落与 body 级未知元素。
"""
from __future__ import annotations

import struct
import zlib
from pathlib import Path

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"

MATH_INLINE = f'<m:oMath xmlns:m="{M}"><m:r><m:t>x</m:t></m:r><m:r><m:t>+1</m:t></m:r></m:oMath>'
MATH_BLOCK = f'<m:oMath xmlns:m="{M}"><m:f><m:num><m:r><m:t>1</m:t></m:r></m:num><m:den><m:r><m:t>2</m:t></m:r></m:den></m:f></m:oMath>'

TBL_XML = f'''<w:tbl xmlns:w="{W}">
  <w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>
  <w:tblGrid><w:gridCol w:w="2000"/><w:gridCol w:w="2000"/><w:gridCol w:w="2000"/></w:tblGrid>
  <w:tr>
    <w:trPr><w:tblHeader w:val="true"/></w:trPr>
    <w:tc><w:tcPr><w:gridSpan w:val="2"/></w:tcPr><w:p><w:r><w:t>表头跨两列</w:t></w:r></w:p></w:tc>
    <w:tc><w:p><w:r><w:t>表头C</w:t></w:r></w:p></w:tc>
  </w:tr>
  <w:tr>
    <w:tc><w:tcPr><w:vMerge w:val="restart"/></w:tcPr><w:p><w:r><w:t>纵向起始</w:t></w:r></w:p></w:tc>
    <w:tc><w:p><w:r><w:t>r2c2</w:t></w:r></w:p></w:tc>
    <w:tc><w:p><w:r><w:t>r2c3</w:t></w:r></w:p></w:tc>
  </w:tr>
  <w:tr>
    <w:tc><w:tcPr><w:vMerge/></w:tcPr><w:p><w:r><w:t>纵向续接文本</w:t></w:r></w:p></w:tc>
    <w:tc><w:p><w:r><w:t>r3c2</w:t></w:r></w:p></w:tc>
    <w:tc><w:p><w:r><w:t>r3c3</w:t></w:r></w:p></w:tc>
  </w:tr>
</w:tbl>'''

SDT_XML = f'<w:sdt xmlns:w="{W}"><w:sdtPr/></w:sdt>'
OBJECT_XML = f'<w:object xmlns:w="{W}"/>'


def tiny_png(width: int = 4, height: int = 4, rgb: tuple[int, int, int] = (200, 30, 30)) -> bytes:
    """自造最小 PNG（无外部依赖）。"""
    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))

    def chunk(tag: bytes, payload: bytes) -> bytes:
        data = tag + payload
        return struct.pack(">I", len(payload)) + data + struct.pack(">I", zlib.crc32(data) & 0xFFFFFFFF)

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def build_sample(path: Path, *, formula_latex: bool = False) -> dict:
    """写一个自建 DOCX 样本，返回其结构说明（供断言引用）。"""
    from docx import Document
    from docx.oxml import parse_xml
    from docx.shared import Emu

    image_bytes = tiny_png()
    document = Document()
    document.add_paragraph("阅读下面的材料")
    document.add_paragraph("材料正文第一段")
    document.add_paragraph("材料正文第二段")
    # 表格（原始 XML 注入，精确控制 vMerge/gridSpan/续接文本）
    # 插到 sectPr 之前 = 紧接上面三段（与 Word 的真实 body 顺序一致）
    document.element.body.insert_element_before(parse_xml(TBL_XML), "w:sectPr")
    inline_p = document.add_paragraph("1. 题干文字")
    inline_p._p.append(parse_xml(MATH_INLINE))
    formula_p = document.add_paragraph()
    formula_p._p.append(parse_xml(MATH_BLOCK))
    image_p = document.add_paragraph("图片前的文字")
    # 无 extent 尺寸由 wp:extent 决定：1 英寸 = 914400 EMU = 96 px
    image_p.add_run().add_picture(__import__("io").BytesIO(image_bytes), width=Emu(914400))
    object_p = document.add_paragraph("未知对象同行")
    object_p._p.append(parse_xml(OBJECT_XML))
    document.add_paragraph("2. 第二题")
    document.add_paragraph("")
    document.element.body.insert_element_before(parse_xml(SDT_XML), "w:sectPr")
    document.save(str(path))
    return {
        "image_bytes": image_bytes,
        "image_sha256": __import__("hashlib").sha256(image_bytes).hexdigest(),
        "blocks": {
            "material": ["p1", "p2", "p3"],
            "table": "t4",
            "inline_stem": "p5",
            "inline_formula": "p5-1",
            "standalone_formula": "p6",
            "image_text": "p7",
            "image": "p7-1",
            "object_text": "p8",
            "second_question": "p9",
        },
        "table_cells": [
            ("表头跨两列", True, 1, 2),
            ("表头C", True, 1, 1),
            ("纵向起始\n纵向续接文本", False, 2, 1),
            ("r2c2", False, 1, 1),
            ("r2c3", False, 1, 1),
            ("r3c2", False, 1, 1),
            ("r3c3", False, 1, 1),
        ],
    }


def source_omml_serials(path: Path) -> list[str]:
    """独立从源 docx 的 word/document.xml 抽出 m:oMath 的字节序列（文档顺序）。"""
    import zipfile

    from lxml import etree

    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    root = etree.fromstring(xml)
    return [
        etree.tostring(node, encoding="unicode")
        for node in root.iter(f"{{{M}}}oMath")
    ]
