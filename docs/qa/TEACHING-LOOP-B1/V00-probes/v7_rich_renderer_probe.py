"""V7 富内容渲染探针（TEACHING-LOOP B1 / V00）。

自建 DOCX（源图关系改名成哨兵 rId404）→ 解析 → 组装富内容 → 渲染学生版/教师版：
  产物可被 python-docx 重新打开；学生版无答案与解析（含答案图片）；
  共同材料只出现一次；图片关系重建（哨兵 rId 不出现、字节一致、新 part）；
  OMML 节点存在；LaTeX-only 经 math2docx 转换；非法 LaTeX → 422 FORMULA_CONVERSION_FAILED；
  表格 columnCount 存在/缺失两条渲染路径的行结构，矛盾 columnCount 明确 422。
"""
from __future__ import annotations

import hashlib
import io
import json
import shutil
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _probe_common import Probe, ensure_api_on_path, temp_data_root  # noqa: E402

p = Probe("v7_rich_renderer")
root = temp_data_root("v7")
ensure_api_on_path()

import _rich_fixture as fx  # noqa: E402

from app.contracts.teaching_loop import (  # noqa: E402
    FormulaBlock,
    RichAsset,
    ImageBlock,
    ParagraphBlock,
    RichContentV2,
    RichOrigin,
    TableBlock,
)
from app.core.exceptions import AppError  # noqa: E402
from app.services.assets.store import AssetStore  # noqa: E402
from app.services.rich_content import (  # noqa: E402
    group_shared_materials,
    omml_from_latex,
    parse_docx_rich,
    render_rich_document,
)

SENTINEL = "rId404"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def rename_image_relationship(path: Path, sentinel: str = SENTINEL) -> str:
    """把源 docx 的图片关系 id 改成哨兵值（改 document.xml + rels），返回原 id。"""
    with zipfile.ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    rels_name = "word/_rels/document.xml.rels"
    rels = entries[rels_name].decode("utf-8")
    import re

    match = re.search(
        r'<Relationship Id="([^"]+)" Type="[^"]*/image"[^>]*/>', rels
    )
    assert match, "样本里没有图片关系"
    original = match.group(1)
    rels = rels.replace(f'Id="{original}"', f'Id="{sentinel}"')
    entries[rels_name] = rels.encode("utf-8")
    entries["word/document.xml"] = (
        entries["word/document.xml"].decode("utf-8").replace(f'r:embed="{original}"', f'r:embed="{sentinel}"').encode("utf-8")
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return original


def image_rels(docx_bytes_or_path) -> dict[str, bytes]:
    """返回 {rId: 图片字节}（按包内 relationships 解析）。"""
    source = (
        io.BytesIO(docx_bytes_or_path)
        if isinstance(docx_bytes_or_path, (bytes, bytearray))
        else docx_bytes_or_path
    )
    with zipfile.ZipFile(source) as archive:
        rels = archive.read("word/_rels/document.xml.rels").decode("utf-8")
        import re

        found: dict[str, bytes] = {}
        for match in re.finditer(r'<Relationship Id="([^"]+)"[^>]*Target="([^"]+)"', rels):
            rid, target = match.group(1), match.group(2)
            if "/image" not in match.group(0):
                continue
            name = "word/" + target.lstrip("/")
            if target.startswith("/"):
                name = target.lstrip("/")
            found[rid] = archive.read(name)
        return found


def all_text(document) -> str:
    chunks: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)
    return "\n".join(chunks)


def omml_count(document) -> int:
    return sum(1 for _ in document.element.body.iter(f"{{{M_NS}}}oMath"))


source_path = root / "source.docx"
spec = fx.build_sample(source_path)
original_rid = rename_image_relationship(source_path)
assets = AssetStore(root / "assets")
parsed = parse_docx_rich(path=source_path, file_name="source.docx", assets=assets)
origin = RichOrigin(
    original_asset_id="orig-1",
    original_sha256=hashlib.sha256(source_path.read_bytes()).hexdigest(),
    source_locator={"page": 1},
)
materials, _unassigned = group_shared_materials(parsed.blocks, parsed.locators)
table_block = next(block for block in parsed.blocks if block.kind == "table")
image_block = next(block for block in parsed.blocks if block.kind == "image")
# 答案专用图片：不同字节 → 教师版必须多一个 media part，学生版不得有
answer_png = fx.tiny_png(width=3, height=3, rgb=(20, 190, 60))
answer_png_sha = hashlib.sha256(answer_png).hexdigest()
answer_stored = assets.store_original(answer_png, media_type="image/png", original_name="answer.png")

stem_table = table_block.model_copy(update={"id": "stem-table"})
first_formula = next(block for block in parsed.blocks if block.kind == "formula")
rich = RichContentV2(
    version=2,
    sharedMaterials=list(materials),
    stemBlocks=[
        ParagraphBlock(id="stem-1", kind="paragraph", text="题干：请计算。"),
        FormulaBlock(id="stem-2", kind="formula", ommlXml=first_formula.omml_xml),
        FormulaBlock(id="stem-3", kind="formula", latex=r"\frac{3}{4}"),
        stem_table,
        image_block.model_copy(update={"id": "stem-image"}),
    ],
    optionBlocks={"A": [ParagraphBlock(id="opt-a", kind="paragraph", text="选项 A 文本")]},
    answerBlocks=[
        ParagraphBlock(id="ans-1", kind="paragraph", text="答案：42（学生版不得出现）"),
        ImageBlock(id="ans-2", kind="image", assetId=answer_stored.blob_key, width=8, height=8),
    ],
    explanationBlocks=[ParagraphBlock(id="exp-1", kind="paragraph", text="解析：因为……（学生版不得出现）")],
    assets=[
        *parsed.assets,
        RichAsset(asset_id=answer_stored.blob_key, sha256=answer_png_sha, media_type="image/png"),
    ],
    origin=origin,
)

student_bytes = render_rich_document(rich=rich, assets=assets, variant="student", title="学生版")
teacher_bytes = render_rich_document(rich=rich, assets=assets, variant="teacher", title="教师版")

from docx import Document  # noqa: E402

student = Document(io.BytesIO(student_bytes))
teacher = Document(io.BytesIO(teacher_bytes))
p.check(
    "v7.1a 学生版/教师版产物可被 python-docx 重新打开（非空、结构可读）",
    len(student.paragraphs) > 0
    and len(teacher.paragraphs) > 0
    and len(student.tables) == 2
    and len(teacher.tables) == 2,
    json.dumps(
        {"studentParas": len(student.paragraphs), "teacherParas": len(teacher.paragraphs), "tables": len(student.tables)}
    ),
)
student_text = all_text(student)
teacher_text = all_text(teacher)
p.check(
    "v7.1b 学生版不含答案与解析文本",
    "答案：42" not in student_text and "解析：因为" not in student_text,
    json.dumps({"studentHasAnswer": "答案：42" in student_text, "studentHasExplanation": "解析：因为" in student_text}),
)
p.check(
    "v7.1c 教师版含答案与解析（顺序：答案在解析前）",
    "答案：42" in teacher_text
    and "解析：因为" in teacher_text
    and teacher_text.index("答案：42") < teacher_text.index("解析：因为"),
    f"answer@{teacher_text.find('答案：42')} explanation@{teacher_text.find('解析：因为')}",
)
def media_blobs(docx_bytes: bytes) -> list[bytes]:
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as archive:
        return [
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("word/media/")
        ]


student_media = media_blobs(student_bytes)
teacher_media = media_blobs(teacher_bytes)
p.check(
    "v7.1d 学生版不含答案图片（media=1、无答案图字节），教师版含（media=2）",
    len(student_media) == 1
    and len(teacher_media) == 2
    and all(blob != answer_png for blob in student_media)
    and any(blob == answer_png for blob in teacher_media),
    json.dumps(
        {
            "studentMedia": len(student_media),
            "teacherMedia": len(teacher_media),
            "studentHasAnswerImage": any(blob == answer_png for blob in student_media),
            "teacherHasAnswerImage": any(blob == answer_png for blob in teacher_media),
        }
    ),
)
p.check(
    "v7.1e 共同材料只出现一次（学生/教师版各 1 次标题、2 次正文各 1 次）",
    student_text.count("阅读下面的材料") == 1
    and student_text.count("材料正文第一段") == 1
    and teacher_text.count("阅读下面的材料") == 1,
    json.dumps(
        {"student": student_text.count("阅读下面的材料"), "teacher": teacher_text.count("阅读下面的材料")}
    ),
)

# ------------------------------------------------------------------ 图片关系
source_rels = image_rels(source_path)
output_rels = image_rels(student_bytes)
p.check(
    "v7.2a 源图片关系已改为哨兵值（样本自检）",
    list(source_rels) == [SENTINEL],
    json.dumps(list(source_rels)),
)
p.check(
    "v7.2b 产物不复用源 rId（哨兵 rId404 不出现），图片关系为新建 part",
    SENTINEL not in output_rels
    and len(output_rels) == 1
    and list(source_rels.values())[0] == list(output_rels.values())[0] == spec["image_bytes"],
    json.dumps({"outputRids": list(output_rels), "bytesMatch": list(output_rels.values())[0] == spec["image_bytes"]}),
)
p.check(
    "v7.2c 产物图片字节与源一致（内容寻址不变）",
    list(output_rels.values())[0] == spec["image_bytes"],
    f"sha256={hashlib.sha256(list(output_rels.values())[0]).hexdigest()[:16]}…",
)
p.check(
    "v7.2d 源文件未被修改（渲染不写回源件）",
    hashlib.sha256(source_path.read_bytes()).hexdigest() == origin.original_sha256,
    "source untouched",
)
image_count_student = len(Document(io.BytesIO(student_bytes)).inline_shapes)
p.check(
    "v7.2e 图片尺寸按块字段还原（96px → 914400 EMU）",
    image_count_student == 1
    and int(Document(io.BytesIO(student_bytes)).inline_shapes[0].width) == 914400,
    json.dumps(
        {
            "shapes": image_count_student,
            "width": int(Document(io.BytesIO(student_bytes)).inline_shapes[0].width) if image_count_student else None,
        }
    ),
)

# ------------------------------------------------------------------ 公式
p.check(
    "v7.3a 产物含 OMML 节点（原样插入 + LaTeX 转换各 1+）",
    omml_count(teacher) >= 2 and omml_count(student) >= 2,
    json.dumps({"teacher": omml_count(teacher), "student": omml_count(student)}),
)
latex_xml = omml_from_latex(r"\frac{3}{4}")
p.check(
    "v7.3b math2docx 转换产物是合法 m:oMath XML（根节点正确）",
    latex_xml.strip().startswith(f"<m:oMath") and M_NS in latex_xml,
    latex_xml[:160],
)
bad = None
try:
    render_rich_document(
        rich=rich.model_copy(
            update={
                "stem_blocks": [
                    FormulaBlock(id="bad-1", kind="formula", latex=r"\frac{1}{")
                ]
            }
        ),
        assets=assets,
        variant="student",
    )
except AppError as exc:
    bad = exc
p.check(
    "v7.3c 非法 LaTeX → 422 FORMULA_CONVERSION_FAILED + details.fields 带块 id",
    bad is not None
    and bad.code == "FORMULA_CONVERSION_FAILED"
    and bad.status_code == 422
    and (bad.details or {}).get("fields") == ["bad-1"],
    f"{getattr(bad, 'code', None)} details={(bad.details if bad else None)}",
)

# ------------------------------------------------------------------ 表格两条渲染路径
doc_with = Document(io.BytesIO(render_rich_document(rich=rich, assets=assets, variant="student")))
table_with = doc_with.tables[0]
p.check(
    "v7.4a columnCount 存在：3x3 网格 + 合并单元格文本不丢",
    len(table_with.rows) == 3
    and len(table_with.columns) == 3
    and table_with.cell(0, 0).text == "表头跨两列"
    and "纵向起始" in table_with.cell(2, 0).text
    and "纵向续接文本" in table_with.cell(2, 0).text,
    json.dumps(
        {
            "rows": len(table_with.rows),
            "cols": len(table_with.columns),
            "c00": table_with.cell(0, 0).text,
            "c20": table_with.cell(2, 0).text,
        },
        ensure_ascii=False,
    ),
)
legacy_table = stem_table.model_copy(update={"column_count": None})
legacy_rich = rich.model_copy(
    update={
        "shared_materials": [],
        "stem_blocks": [legacy_table],
        "option_blocks": {},
        "answer_blocks": [],
        "explanation_blocks": [],
        "assets": [],
    }
)
legacy_doc = Document(io.BytesIO(render_rich_document(rich=legacy_rich, assets=assets, variant="student")))
legacy = legacy_doc.tables[0]
p.check(
    "v7.4b columnCount 缺失（旧数据）：启发式仍还原 3x3 行结构且文本不丢",
    len(legacy.rows) == 3
    and len(legacy.columns) == 3
    and legacy.cell(0, 0).text == "表头跨两列"
    and "纵向起始" in legacy.cell(2, 0).text,
    json.dumps({"rows": len(legacy.rows), "cols": len(legacy.columns)}, ensure_ascii=False),
)
contradict = None
try:
    render_rich_document(
        rich=legacy_rich.model_copy(
            update={"stem_blocks": [stem_table.model_copy(update={"column_count": 5})]}
        ),
        assets=assets,
        variant="student",
    )
except AppError as exc:
    contradict = exc
p.check(
    "v7.4c columnCount 与单元格序列矛盾 → 422 INVALID_REQUEST（不猜表格）",
    contradict is not None
    and contradict.code == "INVALID_REQUEST"
    and contradict.status_code == 422
    and (contradict.details or {}).get("fields") == ["stem-table"],
    f"{getattr(contradict, 'code', None)} details={(contradict.details if contradict else None)}",
)

# ------------------------------------------------------------------ 其它
dup = None
try:
    render_rich_document(
        rich=rich.model_copy(
            update={"answer_blocks": [ParagraphBlock(id="stem-1", kind="paragraph", text="重复 id")]}
        ),
        assets=assets,
        variant="student",
    )
except AppError as exc:
    dup = exc
p.check(
    "v7.5a 同一块 id 在多处出现 → 422（不静默去重/重复输出）",
    dup is not None and dup.status_code == 422 and (dup.details or {}).get("fields") == ["stem-1"],
    f"{getattr(dup, 'code', None)} details={(dup.details if dup else None)}",
)
corrupt = None
blob = image_block.asset_id
blob_path = assets.path_of(blob)
original_bytes = blob_path.read_bytes()
try:
    blob_path.write_bytes(b"tampered-not-the-original")
    try:
        render_rich_document(rich=rich, assets=assets, variant="student")
    except AppError as exc:
        corrupt = exc
finally:
    blob_path.write_bytes(original_bytes)
p.check(
    "v7.5b 受管资产被篡改 → ASSET_CORRUPT（不返回可疑内容）",
    corrupt is not None and corrupt.code == "ASSET_CORRUPT",
    f"{getattr(corrupt, 'code', None)}",
)

sys.exit(p.finish())
