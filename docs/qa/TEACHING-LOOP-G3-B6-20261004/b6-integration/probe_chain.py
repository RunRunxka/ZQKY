"""New B6 no-TCP real endpoint chain and complete fixed-history conservation."""
from datetime import datetime, timezone
import contextlib
import copy
import hashlib
import io
import json
import os
from pathlib import Path
import sys
import time
import traceback

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[3]
sys.path.insert(0, str(HERE))
from runtime_fixture import TEACHER, FIVE, SIX, PATCH_TEXT, PROCESS_TEXT, fixed_reply, prior_seed_module, SEED

def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'))

def digest(value):
    return hashlib.sha256(canonical(value).encode('utf-8')).hexdigest()

def main():
    output = Path(sys.argv[1]).resolve()
    output.mkdir(exist_ok=False)
    data = Path(os.environ['ZQKY_DATA_DIR']).resolve()
    assert data.name == 'data' and data.parent.name.startswith('zqky-b5-b6-integration-')
    assert os.environ['ZQKY_ENV'] == 'test' and os.environ['PYTHONUTF8'] == '1'
    assert Path(os.environ['TEMP']).resolve() in data.parents
    sys.dont_write_bytecode = True
    sys.path.insert(0, str(ROOT / 'apps/api'))
    # All environment and new TEMP established before importing standard main.
    from app.core.config import Settings
    from app.core.secrets import SecretStore
    from app.main import create_app
    from app.core.sqlite import open_readonly
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    settings = Settings(host='127.0.0.1', port=8001, env='test', data_dir=data,
        credentials_file=None, qdrant_url='http://127.0.0.1:16333', embedding_base_url='http://127.0.0.1:9',
        textbook_source_dir=Path(os.environ['ZQKY_TEXTBOOK_SOURCE_DIR']),
        allowed_origins=frozenset({'http://127.0.0.1:5174'}))
    app = create_app(settings, secret_store=SecretStore())
    assert app.state.lesson_generation_service.evidence is app.state.rag_v2
    fixture = None
    started = time.perf_counter()
    record = dict(task='B6-INTEGRATION-v1', pid=os.getpid(), startedAt=datetime.now(timezone.utc).isoformat(),
        sampleRoot=str(data.parent), credentialsFile=None, appMainPreImportIsolation=True,
        tcpStarted=False, paidModel=False, businessFetchMocked=False, seedHelperSHA=hashlib.sha256(SEED.read_bytes()).hexdigest())
    def save(name, value):
        path = output / (name+'.json')
        assert not path.exists(), path
        path.write_text(json.dumps(value, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
        return dict(path=str(path), sha256=hashlib.sha256(path.read_bytes()).hexdigest())
    calls = []
    try:
        with TestClient(app, base_url='http://127.0.0.1:8001') as client:
            class Recorder:
                def __getattr__(self, name):
                    original = getattr(client, name)
                    if name not in {'get', 'post', 'patch', 'delete'}:
                        return original
                    def invoke(path, **kwargs):
                        response = original(path, **kwargs)
                        body = copy.deepcopy(kwargs.get('json'))
                        if isinstance(body, dict):
                            body.pop('apiKey', None)
                        try:
                            reply = response.json()
                        except ValueError:
                            reply = dict(binarySHA=hashlib.sha256(response.content).hexdigest(), bytes=len(response.content))
                        calls.append(dict(method=name, path=path, request=body, status=response.status_code, response=reply))
                        return response
                    return invoke
            api = Recorder()
            module = prior_seed_module()
            seed, fixture = module.seed_for_browser(app, output/'seed-raw.json', client=api, tag='b6-integration-fixed')
            fixture.reply_override = fixed_reply()
            fixture.install()
            def request(method, path, status=200, **kwargs):
                response = getattr(api, method)('/api/v1'+path, **kwargs)
                assert response.status_code == status, (path, response.status_code, response.text)
                return response.json()
            def ready(receipt):
                deadline = time.monotonic()+15
                while time.monotonic() < deadline:
                    job = request('get', '/workflow-jobs/'+receipt['job']['jobId'], params={'domain':'teaching'})
                    if job['state'] in {'succeeded','failed','cancelled','interrupted'}:
                        assert job['state'] == 'succeeded', job
                        return job
                    time.sleep(.03)
                raise AssertionError('New B6 bounded job wait expired')
            kp = seed['selectedKnowledgePointIds'][0]
            point = next(p for p in seed['originalFixed']['report']['knowledgePoints'] if p['knowledgePointId'] == kp)
            context = dict(analysisRunId=seed['analysisRunId'], selectedKnowledgePointIds=[kp])
            lesson = request('post', '/lesson-plans', 201, json=dict(submissionId='b6-new-lesson',
                subjectId='math', classId=seed['classId'], data=TEACHER, context=context, source='manual'))
            assert lesson['currentRevision']['contextSnapshot']['analysis']['knowledgePoints'] == [point]
            verify = request('post', '/lesson-plans/evidence/verify', json={
                'selection': seed['textbook']['selection'], 'slices': seed['textbook']['slices']})
            generation = dict(submissionId='b6-five-generate', baseRevisionId=lesson['currentRevisionId'],
                baseServerRevision=lesson['revision'], analysisRunId=seed['analysisRunId'], classId=seed['classId'],
                selectedKnowledgePointIds=[kp], requirements='基于固定教材和班级计数安排有理数加法复习。', durationMinutes=43,
                modelProfileId=seed['modelProfiles'][0]['profileId'], scopeSnapshot=verify['scopeSnapshot'],
                evidenceRefs=verify['evidenceRefs'], questionRevisionIds=[seed['questionRevisionId']], practiceRevisionIds=[seed['practiceRevisionId']])
            receipt = request('post', '/lesson-plans/'+lesson['lessonPlanId']+'/proposals', 202, json=generation)
            job = ready(receipt)
            proposal = request('get', '/lesson-plans/'+lesson['lessonPlanId']+'/proposals/'+job['result']['proposalId'])
            assert len(fixture.calls) == 1
            wire_payload = fixture.user_payload(fixture.calls[0]['body'])
            assert len(wire_payload['classSummary']['knowledgePoints']) == 1
            assert wire_payload['classSummary']['knowledgePoints'][0]['counts'] == dict(
                selectedCount=1, validCount=1, needsCount=1, incompleteCount=0, noEvidenceCount=0, fullCreditCount=0, numerator=1, denominator=1)
            assert [e['kind'] for e in wire_payload['evidence']] == ['textbook','question','practice']
            assert wire_payload['durationMinutes'] == 43
            with app.state.teaching.read_connection() as conn:
                row = conn.execute('SELECT * FROM lesson_generation_inputs WHERE job_id=?', (job['jobId'],)).fetchone()
                frozen = json.loads(row['frozen_json'])
            # Independent mapping from recorded immutable input, never stable_ids or apply/merge.
            frozen_hash = digest(frozen)
            expected_process = [dict(id='lp_'+hashlib.sha256((frozen_hash+'\0'+f'new:N{n}').encode()).hexdigest()[:32], **item)
                for n, item in enumerate(PROCESS_TEXT, 1)]
            expected_body = {**copy.deepcopy(TEACHER), **PATCH_TEXT, 'process':expected_process}
            save('handwritten-oracle', dict(teacher=TEACHER, selectedFields=FIVE, teacherFields=SIX,
                patchText=PATCH_TEXT, processText=PROCESS_TEXT, expectedBody=expected_body, minutes=[11,11,11,10],
                ruleCountsBefore=dict(selectedCount=1,validCount=1,needsCount=1,numerator=1,denominator=1),
                ruleCountsAfter=dict(selectedCount=1,validCount=1,needsCount=0,fullCreditCount=1,numerator=0,denominator=1)))
            apply_body = dict(submissionId='b6-five-apply', expectedRevision=lesson['revision'],
                baseRevisionId=lesson['currentRevisionId'], selectedFields=FIVE)
            applied = request('post', '/lesson-plans/'+lesson['lessonPlanId']+'/proposals/'+proposal['proposalId']+'/apply',
                json=apply_body)
            assert applied['currentRevision']['data'] == expected_body
            assert applied['currentRevision']['selectedFields'] == FIVE
            assert all(applied['currentRevision']['data'][field] == TEACHER[field] for field in SIX)
            assert applied['currentRevision']['source'] == 'ai_applied' and applied['revision'] == lesson['revision']+1
            assert [s['minutes'] for s in applied['currentRevision']['processMetadata']] == [11,11,11,10]
            assert request('get', '/lesson-plans/'+lesson['lessonPlanId']+'/revisions/'+lesson['currentRevisionId']) == lesson['currentRevision']
            assert request('get', '/lesson-plans/'+lesson['lessonPlanId']+'/revisions/'+applied['currentRevisionId']) == applied['currentRevision']
            assert len(request('get', '/lesson-plans/'+lesson['lessonPlanId']+'/revisions')['items']) == 2
            # Formal reviewed practice is distinct from the lesson exercises string.
            practice = request('post', '/practice-sets', 201, json=dict(submissionId='b6-new-practice',
                analysisRunId=seed['analysisRunId'], title='B6正式单KP练习', targetKnowledgePointIds=[kp], constraints={'count':1}))
            source_item = seed['originalFixed']['practice']['draftItems'][0]
            item = copy.deepcopy(source_item)
            item.update(itemKey='b6-single-kp-item', selectedKnowledgePointIds=[kp])
            item['itemStructure']['nodes'][0]['knowledgePointIds'] = [kp]
            practice = request('patch', '/practice-sets/'+practice['practiceSetId']+'/draft', json=dict(
                submissionId='b6-practice-save', expectedRevision=practice['revision'], items=[item], constraints={'count':1}))
            practice = request('post', '/practice-sets/'+practice['practiceSetId']+'/review', json=dict(
                submissionId='b6-practice-review', expectedRevision=practice['revision']))
            fixed = practice['currentRevision']['practiceRevisionId']
            assert practice['currentRevision']['state'] == 'reviewed'
            assert practice['currentRevision']['targetKnowledgePoints'] == [point]
            student = seed['originalFixed']['score']['participantSnapshot'][0]
            conversion = request('post', f"/practice-sets/{practice['practiceSetId']}/revisions/{fixed}/assessments", 201,
                json=dict(submissionId='b6-new-conversion', title='B6练习回流施测', heldOn='2026-10-04',
                    classIds=[seed['classId']], participants=[dict(studentId=student['studentId'],classId=seed['classId'],attendance='present',attemptNo=1)]))
            export_receipt = request('post', f"/practice-sets/{practice['practiceSetId']}/revisions/{fixed}/exports", 202,
                json=dict(submissionId='b6-score-template',variant='score_template',assessmentId=conversion['assessmentId']))
            exported = ready(export_receipt)
            artifact = request('get', '/export-artifacts/'+exported['result']['artifactId'])
            response = api.get(artifact['downloadUrl'])
            assert response.status_code == 200 and hashlib.sha256(response.content).hexdigest() == artifact['sha256']
            (output/'new-score-template.xlsx').write_bytes(response.content)
            workbook = load_workbook(io.BytesIO(response.content))
            sheet = workbook.active
            assert [sheet.cell(1,n).value for n in range(1,6)] == ['学号','姓名','出勤','人次','1']
            assert sheet.cell(2,1).value == student['studentNo']
            sheet.cell(2,5).value = 1.25  # Teacher entry; no model scoring.
            stream = io.BytesIO(); workbook.save(stream)
            (output/'teacher-filled-scores.xlsx').write_bytes(stream.getvalue())
            assessment = request('get', '/assessments/'+conversion['assessmentId'])
            imported = request('post', '/assessments/'+conversion['assessmentId']+'/score-imports', 201,
                files={'file':('b6-teacher-scores.xlsx',stream.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
            scored = request('post', '/score-imports/'+imported['importId']+'/confirm', json=dict(submissionId='b6-returned-score',
                expectedImportRevision=imported['revision'], expectedAssessmentRevision=assessment['assessment']['revision'],
                previewVersion=imported['previewVersion'], baseScoreRevisionId=None))
            returned = request('post', '/assessments/'+conversion['assessmentId']+'/analysis-runs', 202,
                json=dict(submissionId='b6-returned-report',scoreRevisionId=scored['revisionId'],
                    selectedParticipantIds=[assessment['participants'][0]['participantId']],ruleCode='any_loss_v1'))
            ready(returned)
            classes = request('get', '/analysis-runs/'+returned['runId']+'/classes')
            assert len(classes['items']) == 1
            row = classes['items'][0]
            for field, expected in dict(selectedCount=1,validCount=1,needsCount=0,fullCreditCount=1,numerator=0,denominator=1).items():
                assert row[field] == expected, (field,row)
            assert row['knowledgePoint'] == point and row['ratio'] == 0.0
            returned_evidence = request('get', '/analysis-runs/'+returned['runId']+'/evidence')
            assert len(returned_evidence['items']) == 1
            ev = returned_evidence['items'][0]
            assert ev['practiceRevisionId'] == fixed and ev['practiceItemId'] == practice['currentRevision']['items'][0]['practiceItemId']
            # Capture complete JSON responses plus full SQL columns, retaining raw JSON strings.
            paths = ['/score-revisions/'+seed['scoreRevisionId'], '/score-revisions/'+seed['scoreRevisionId']+'/matrix',
                '/analysis-runs/'+seed['analysisRunId'], '/analysis-runs/'+seed['analysisRunId']+'/classes',
                '/analysis-runs/'+seed['analysisRunId']+'/students', '/analysis-runs/'+seed['analysisRunId']+'/evidence',
                f"/practice-sets/{seed['practiceSetId']}/revisions/{seed['practiceRevisionId']}",
                '/lesson-plans/'+lesson['lessonPlanId']+'/revisions/'+lesson['currentRevisionId'],
                '/lesson-plans/'+lesson['lessonPlanId']+'/revisions/'+applied['currentRevisionId'],
                '/lesson-plans/'+lesson['lessonPlanId']+'/proposals/'+proposal['proposalId'],
                '/score-revisions/'+scored['revisionId'], '/score-revisions/'+scored['revisionId']+'/matrix',
                '/analysis-runs/'+returned['runId'], '/analysis-runs/'+returned['runId']+'/classes',
                '/analysis-runs/'+returned['runId']+'/students', '/analysis-runs/'+returned['runId']+'/evidence',
                f"/practice-sets/{practice['practiceSetId']}/revisions/{fixed}"]
            before_json = {path:request('get',path) for path in paths}
            def sql_snapshot():
                result = {}
                for cat_name in ('teaching','catalog','knowledge','question_bank'):
                    cat = getattr(app.state, cat_name)
                    with contextlib.closing(open_readonly(cat.db_path)) as conn:
                        names = [r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")]
                        if cat_name == 'teaching':
                            names = [n for n in names if n.startswith(('score_','student_item_','analysis_','practice_','paper_','lesson_')) or n in {'papers','assessments','assessment_classes','assessment_participants'}]
                        elif cat_name == 'knowledge':
                            names = [n for n in names if n in {'knowledge_point_revisions','knowledge_aliases','textbook_knowledge_links'}]
                        elif cat_name == 'question_bank':
                            names = [n for n in names if n in {'question_revisions','question_knowledge_links'}]
                        result[cat_name] = {}
                        for table in names:
                            columns = [r['name'] for r in conn.execute(f'PRAGMA table_info("{table}")')]
                            values = [dict(row) for row in conn.execute(f'SELECT * FROM "{table}"')]
                            values.sort(key=canonical)
                            result[cat_name][table] = dict(sql=f'SELECT * FROM "{table}"', columns=columns, rows=values)
                        assert conn.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
                        assert conn.execute('PRAGMA foreign_key_check').fetchall() == []
                return result
            before_sql = sql_snapshot()
            assert before_sql['question_bank']['question_revisions']['rows']
            before_sql_receipt = save('before-mutations-full-sql',before_sql)
            before_json_receipt = save('before-mutations-full-json',before_json)
            # Source labels/maps live in complete context/source_metadata/frozen_json columns above.
            changed = []
            question = request('get','/questions/'+seed['questionId'])
            content = copy.deepcopy(question['content'])
            old_paragraph = content['richContent']['stemBlocks'][0]['text']
            assert old_paragraph in content['stemMarkdown']
            content['stemMarkdown'] = content['stemMarkdown'].replace(old_paragraph, 'B6题库新题面，与旧固定事实不同', 1)
            content['richContent']['stemBlocks'][0]['text'] = 'B6题库新题面，与旧固定事实不同'
            edited = request('patch','/questions/'+seed['questionId'],json=dict(expectedRevision=question['revision'],
                content=content,metadata=question['metadata']))
            assert edited['content']['stemMarkdown'] != question['content']['stemMarkdown']
            changed.append('question-bank-new-revision')
            klass = request('get','/classes/'+seed['classId'])
            renamed = request('patch','/classes/'+seed['classId'],json=dict(expectedRevision=klass['revision'],name='B6新班名不改旧标签'))
            assert renamed['name'] != klass['name']; changed.append('class-renamed')
            live_student = request('get','/students/'+student['studentId'])
            renamed_student = request('patch','/students/'+student['studentId'],json=dict(
                expectedRevision=live_student['revision'],name='B6名单新显示名'))
            assert renamed_student['name'] != live_student['name']; changed.append('roster-student-display-name-renamed')
            point_live = request('get','/knowledge-points/'+kp)
            archived = request('post','/knowledge-points/'+kp+'/archive',json=dict(expectedRevision=point_live['revision']))
            assert archived['status'] == 'archived'; changed.append('selected-knowledge-point-archived')
            calls_before = len(fixture.calls)
            invalid_request = {**generation, 'submissionId':'b6-invalid-after-mutations',
                'baseRevisionId':applied['currentRevisionId'], 'baseServerRevision':applied['revision']}
            rejection = request('post','/lesson-plans/'+lesson['lessonPlanId']+'/proposals',409,json=invalid_request)
            assert rejection['code'] == 'KNOWLEDGE_ARCHIVED'
            assert len(fixture.calls) == calls_before == 1
            old_generate_replay = request('post','/lesson-plans/'+lesson['lessonPlanId']+'/proposals',202,json=generation)
            assert old_generate_replay['replayed'] is True
            assert old_generate_replay['inputHash'] == receipt['inputHash'] and old_generate_replay['job']['jobId'] == receipt['job']['jobId']
            old_apply_replay = request('post','/lesson-plans/'+lesson['lessonPlanId']+'/proposals/'+proposal['proposalId']+'/apply',json=apply_body)
            assert old_apply_replay == {**applied, 'replayed':True}
            assert len(fixture.calls) == calls_before
            after_json = {path:request('get',path) for path in paths}
            after_sql = sql_snapshot()
            assert before_json == after_json, 'Full old/new fixed HTTP snapshots changed after mutable catalog edits'
            # Question edit appends a revision, but every preexisting SQL row must be identical.
            conservation = {}
            for cat, tables in before_sql.items():
                conservation[cat] = {}
                for table, expected in tables.items():
                    actual = after_sql[cat][table]
                    assert expected['columns'] == actual['columns']
                    assert all(row in actual['rows'] for row in expected['rows']), (cat,table)
                    if cat != 'question_bank':
                        assert expected['rows'] == actual['rows'], (cat,table)
                    conservation[cat][table] = dict(columns=len(expected['columns']),beforeRows=len(expected['rows']),
                        afterRows=len(actual['rows']), allOldFullRowsEqual=True)
            after_sql_receipt = save('after-mutations-full-sql',after_sql)
            after_json_receipt = save('after-mutations-full-json',after_json)
            chain = dict(seedSource='immutable seed helper, new execution', singleClassId=seed['classId'], singleKnowledge=point,
                sourceScoreId=seed['scoreRevisionId'], sourceReportId=seed['analysisRunId'], lesson=lesson, verified=verify,
                generationReceipt=receipt, proposal=proposal, applied=applied, formalPractice=practice, conversion=conversion,
                template=artifact, teacherScore=scored, returnedReport=request('get','/analysis-runs/'+returned['runId']),
                returnedClasses=classes, returnedEvidence=returned_evidence, mutations=changed, invalidNewCall=rejection,
                oldGenerationReplay=old_generate_replay,oldApplyReplay=old_apply_replay,
                providerCallsBeforeInvalid=calls_before, providerCallsAfterInvalid=len(fixture.calls), fixedConservation=conservation,
                beforeSQL=before_sql_receipt,afterSQL=after_sql_receipt,beforeJSON=before_json_receipt,afterJSON=after_json_receipt)
            save('chain-result',chain)
            record.update(status='TECHNICAL_CHAIN_PASS_TEACHER_REVIEW_PENDING',fiveFieldsApplied=True,teacherSixPreserved=True,
                newScoreUnits=125,needsBefore=1,needsAfter=0,providerCalls=1,invalidProviderDelta=0,
                fixedCompleteJSON=len(paths),oldFullSQLRows=sum(len(t['rows']) for cat in before_sql.values() for t in cat.values()),
                fixedConservation=conservation)
        record['testClientClosed'] = True
    except BaseException:
        record.update(status='FAILED',firstFailure=traceback.format_exc())
        raise
    finally:
        if fixture:
            fixture.restore(); record['transportRestored'] = not fixture.originals
        save('endpoint-calls', calls)
        record.update(finishedAt=datetime.now(timezone.utc).isoformat(),durationMs=round((time.perf_counter()-started)*1000,3),sampleRetained=data.exists())
        save('RESULT',record)
        print(json.dumps({key:record.get(key) for key in ('status','pid','durationMs','sampleRoot','fixedCompleteJSON','oldFullSQLRows')}, ensure_ascii=False),flush=True)

if __name__ == '__main__':
    main()
