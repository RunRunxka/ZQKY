"""Independent stdlib ZIP oracle and real downloaded XLSX score entry.

Invoked later by browser after real downloads. Own B4-V00 output roots only.
Never imports app.main or any production renderer/aggregation.
"""
import argparse
import hashlib
import json
import zipfile
import struct
import zlib
from pathlib import Path
from xml.etree import ElementTree as ET


def owned(path):
    value=Path(path).resolve()
    qa=Path(__file__).resolve().parent
    if not value.is_relative_to(qa) or "browser-artifacts-" not in str(value):
        raise ValueError("Only this batch's new browser artifacts are writable/readable here")
    return value


def verify(files, metadata):
    from openpyxl import load_workbook
    result={}
    parts={}
    for kind,path in files.items():
        data=owned(path).read_bytes()
        with zipfile.ZipFile(owned(path)) as archive:
            assert archive.testzip() is None
            entries=[dict(name=i.filename,crc=i.CRC,size=i.file_size,sha256=hashlib.sha256(archive.read(i)).hexdigest()) for i in archive.infolist()]
            parts[kind]={i.filename:archive.read(i) for i in archive.infolist()}
        expected=metadata[kind]
        assert len(data)==expected["byteSize"]
        assert hashlib.sha256(data).hexdigest()==expected["sha256"]
        result[kind]=dict(size=len(data),sha256=hashlib.sha256(data).hexdigest(),entries=entries)
    for value in parts["student"].values():
        assert b"B4_TEACHER_ANSWER_731" not in value and b"B4_TEACHER_EXPLANATION_947" not in value
    teacher=b"".join(parts["teacher"].values())
    assert b"B4_TEACHER_ANSWER_731" in teacher and b"B4_TEACHER_EXPLANATION_947" in teacher
    ns={"w":"http://schemas.openxmlformats.org/wordprocessingml/2006/main","m":"http://schemas.openxmlformats.org/officeDocument/2006/math"}
    for variant in ("student","teacher"):
        document=ET.fromstring(parts[variant]["word/document.xml"])
        text="".join(document.itertext())
        assert "16(1)" in text and "2.50" in text and "2.00" in text
        assert document.findall(".//w:tbl",ns), "real table must remain"
        assert document.findall(".//m:oMath",ns), "real OMML must remain"
        assert any(name.startswith("word/media/") for name in parts[variant]), "real image bytes must remain"
        assert text.count("共同材料：保持完整上下文。") == 1
    workbook=load_workbook(owned(files["score_template"]))
    sheet=workbook["成绩"]
    assert [sheet.cell(1,n).value for n in (5,6)]==["16(1)","2"]
    assert sheet.max_row==5
    for row in range(2,6):
        assert isinstance(sheet.cell(row,1).value,str) and sheet.cell(row,1).value.startswith("00")
        assert sheet.cell(row,1).data_type=="s" and sheet.cell(row,2).data_type=="s"
        assert sheet.cell(row,5).value is None and sheet.cell(row,6).value is None
    workbook.close()
    return result


def fill(source,target):
    from openpyxl import load_workbook
    workbook=load_workbook(owned(source))
    sheet=workbook["成绩"]
    # Actual generated roster: A0+1.5; B2.5+missing; Cabsent; D2.5+2.
    # Scores go into the downloaded blank template, not a replacement workbook.
    expected={"合成甲-B4":(0,1.5),"合成乙-B4":(2.5,None),"合成丙-B4":(None,None),"合成丁-B4":(2.5,2)}
    for row in range(2,sheet.max_row+1):
        values=expected[sheet.cell(row,2).value]
        for column,value in enumerate(values,5):sheet.cell(row,column,value)
    workbook.save(owned(target));workbook.close()


def pixels(paths):
    def decode(path):
        data=owned(path).read_bytes();assert data[:8]==b"\x89PNG\r\n\x1a\n"
        offset=8;compressed=b"";header=None
        while offset<len(data):
            length=struct.unpack(">I",data[offset:offset+4])[0];kind=data[offset+4:offset+8];payload=data[offset+8:offset+8+length]
            assert zlib.crc32(kind+payload)&0xffffffff==struct.unpack(">I",data[offset+8+length:offset+12+length])[0]
            if kind==b"IHDR":header=struct.unpack(">IIBBBBB",payload)
            if kind==b"IDAT":compressed+=payload
            offset+=12+length
        width,height,depth,color,compression,filtering,interlace=header
        assert depth==8 and color in (2,6) and compression==filtering==interlace==0
        channels=3 if color==2 else 4;stride=width*channels;raw=zlib.decompress(compressed);prior=bytearray(stride);result=[]
        def paeth(a,b,c):
            p=a+b-c;distances=[abs(p-a),abs(p-b),abs(p-c)]
            return (a,b,c)[distances.index(min(distances))]
        for row in range(height):
            start=row*(stride+1);mode=raw[start];current=bytearray(raw[start+1:start+1+stride])
            for i in range(stride):
                left=current[i-channels] if i>=channels else 0;up=prior[i];corner=prior[i-channels] if i>=channels else 0
                delta=(0,left,up,(left+up)//2,paeth(left,up,corner))[mode]
                current[i]=(current[i]+delta)&255
            result.extend(tuple(current[x:x+3]) for x in range(0,stride,channels));prior=current
        return width,height,result
    decoded=[decode(p) for p in paths[:4]]
    assert len({(w,h) for w,h,_ in decoded})==1
    normal=sum(a!=b for a,b in zip(decoded[0][2],decoded[1][2]));reduced=sum(a!=b for a,b in zip(decoded[2][2],decoded[3][2]))
    assert normal>0 and reduced>0, "Both actual hover states must visibly change rendered pixels"
    # Semantic style/frame timing is independently checked by Playwright; PNG
    # oracle proves actual changed RGB pixels rather than matchMedia alone.
    owned(paths[4]).write_text(json.dumps(dict(width=decoded[0][0],height=decoded[0][1],normalChangedRGBPixels=normal,reducedChangedRGBPixels=reduced,files=[dict(path=p,sha256=hashlib.sha256(owned(p).read_bytes()).hexdigest()) for p in paths[:4]]),indent=2),encoding="utf-8")


if __name__=="__main__":
    parser=argparse.ArgumentParser();parser.add_argument("action",choices=["verify","fill","pixels","db"]);parser.add_argument("args",nargs="+")
    options=parser.parse_args()
    if options.action=="fill":
        fill(*options.args)
    elif options.action=="db":
        from db_audit import verify as verify_db
        verify_db(options.args[0],owned(options.args[1]),owned(options.args[2]))
    elif options.action=="pixels":
        pixels(options.args)
    else:
        manifest=json.loads(owned(options.args[0]).read_text(encoding="utf-8"))
        result=verify(manifest["files"],manifest["metadata"])
        owned(options.args[1]).write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
