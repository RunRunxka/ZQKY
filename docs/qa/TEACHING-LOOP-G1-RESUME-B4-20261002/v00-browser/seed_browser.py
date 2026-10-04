"""Fresh G1R seed: unchanged 2/3/5 score paper, plus real DOCX OMML review input.

CTRL owns execution. This script never starts a listener and refuses reused data.
All app imports follow explicit, verified OS-temp ZQKY_DATA_DIR and test env.
"""
from __future__ import annotations

import hashlib
import json
import os
import sys
import tempfile
from io import BytesIO
from pathlib import Path

root = Path(os.environ["ZQKY_DATA_DIR"]).resolve()
temporary = Path(tempfile.gettempdir()).resolve()
repo = Path.cwd().resolve()
assert os.environ["ZQKY_ENV"] == "test"
assert root.is_absolute() and root.is_relative_to(temporary) and root != temporary
assert root.name == "data" and root.parent.name.startswith("zqky-g1-resume-browser-")
assert not root.exists() or not any(root.iterdir()), "Seed requires a fresh empty data root"
assert not root.is_relative_to(repo) and ".local-data" not in str(root)
root.parent.mkdir(parents=True, exist_ok=True)
sys.path.insert(0, str(repo / "apps/api"))
sys.path.insert(0, str(repo / "tests/fixtures"))

from docx import Document  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from tests.scores_support import ScoresHarness  # noqa: E402
from teaching_loop_docx import build_complete_paper  # noqa: E402

namespace = "http://schemas.openxmlformats.org/officeDocument/2006/math"
cases = [
    {"name": "default-pipe", "properties": "", "expectedText": "(x|y)", "expectedOperators": ["(", "|", ")"]},
    {"name": "explicit-comma", "properties": '<m:sepChr m:val=","/>', "expectedText": "(x,y)", "expectedOperators": ["(", ",", ")"]},
    {"name": "explicit-pipe", "properties": '<m:begChr m:val="["/><m:endChr m:val="]"/><m:sepChr m:val="|"/>', "expectedText": "[x|y]", "expectedOperators": ["[", "|", "]"]},
]
document = Document(BytesIO(build_complete_paper()))
anchor = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("2."))
for case in cases:
    paragraph = document.add_paragraph(f'R08 {case["name"]}：两个参数 x 与 y，保留原始 OMML 分隔符。')
    paragraph._p.append(parse_xml(
        f'<m:oMath xmlns:m="{namespace}"><m:d><m:dPr>{case["properties"]}</m:dPr>'
        '<m:e><m:r><m:t>x</m:t></m:r></m:e><m:e><m:r><m:t>y</m:t></m:r></m:e>'
        '</m:d></m:oMath>'
    ))
    anchor._p.addprevious(paragraph._p)
stream = BytesIO()
document.save(stream)
content = stream.getvalue()
rich_path = root.parent / "g1-resume-rich-paper.docx"
rich_path.write_bytes(content)

with ScoresHarness(root.parent) as harness:
    assert harness.settings.credentials_file is None
    paper = harness.seed_confirmed_paper(
        tag="g1resume-browser", leaves=(("Q1", 200), ("Q2", 300), ("Q3", 500)),
        title="G1 固定验证卷",
    )

metadata = {
    "dataDir": str(root), "paperId": paper.paper_id, "paperRevisionId": paper.revision_id,
    "title": paper.title, "richPaperFile": str(rich_path),
    "richPaperSha256": hashlib.sha256(content).hexdigest(),
    "delimiters": [{key: case[key] for key in ("name", "expectedText", "expectedOperators")} for case in cases],
    "credentialsFile": None, "scoreLeaves": [["Q1", 200], ["Q2", 300], ["Q3", 500]],
}
destination = repo / "docs/qa/TEACHING-LOOP-G1-RESUME-B4-20261002/v00-browser/browser-seed.json"
assert not destination.exists(), "Refusing to replace prior seed metadata"
destination.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(metadata, ensure_ascii=False))
