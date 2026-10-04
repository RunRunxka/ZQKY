"""B4-V00-P: independent T80 API/DB/archive oracles on real four catalogs."""
import copy
from contextlib import contextmanager
from datetime import UTC, datetime, timedelta
import hashlib
import io
import json
import sqlite3
import pytest
from openpyxl import Workbook, load_workbook
from probe_support import Scene, scene, http, wait_job, zip_oracle, record
from app.core.exceptions import AppError
from app.services.jobs.engine import FrozenJob, JobContext


def test_constraints_multikp_gap_unknown_difficulty_unknown_original_type(scene):
    s=scene
    # Only the fixed source original has this literal surface. It has no recorded type.
    with s.catalog.read_connection() as conn:
        source=json.loads(conn.execute("SELECT content_json FROM paper_items WHERE paper_revision_id=?",(s.seed['analysis']['paperRevisionId'],)).fetchone()[0])
    assert "type" not in source or source["type"] is None
    duplicates=[s.question(source['stemBlocks'][0]['text'],qtype=t,rich=source) for t in
        ("single_choice","multiple_choice","fill_blank","true_false","short_answer","other")]
    unknown=s.question("独立未知难度",difficulty="unspecified")
    accepted=s.question("独立适合难题",difficulty="hard")
    s.question("独立适合难题",difficulty="hard")  # same surface; deduplicate must constrain count
    practice=s.create(count=3,questionTypes=["short_answer"],difficulties=["hard"])
    before=http(s.client,"get","/practice-sets/"+practice['practiceSetId'])
    result=http(s.client,"post","/practice-sets/"+practice['practiceSetId']+"/suggestions",json={
        "expectedRevision":0,"constraints":{"count":3,"questionTypes":["short_answer"],"difficulties":["hard"]}})
    assert result['selectedCount']==1 and result['requestedCount']==3
    assert result['items'][0]['questionType']=="short_answer" and result['items'][0]['difficulty']=="hard"
    assert set(result['coverage'])==set(s.kps) and all(n==1 for n in result['coverage'].values())
    assert result['gaps'] and all(isinstance(x,str) for x in result['gaps'])
    assert http(s.client,"get","/practice-sets/"+practice['practiceSetId'])==before
    all_types=http(s.client,"post","/practice-sets/"+practice['practiceSetId']+"/suggestions",json={
        "expectedRevision":0,"constraints":{"count":100}})
    chosen={x['questionRevisionId'] for x in all_types['items']}
    assert not chosen.intersection(q['questionRevisionId'] for q in duplicates)
    assert unknown['questionRevisionId'] not in chosen
    including=http(s.client,"post","/practice-sets/"+practice['practiceSetId']+"/suggestions",json={
        "expectedRevision":0,"constraints":{"count":100,"includeUnknownDifficulty":True}})
    assert unknown['questionRevisionId'] in {x['questionRevisionId'] for x in including['items']}
    for invalid in ({"count":1,"questionTypes":["invented"]},{"count":1,"difficulties":["invented"]},
                    {"count":1,"questionTypes":["short_answer","short_answer"]}):
        http(s.client,"post","/practice-sets/"+practice['practiceSetId']+"/suggestions",422,json={"expectedRevision":0,"constraints":invalid})
    record("constraints-oracle",dict(originalType="unrecorded",excludedOriginalTypes=len(duplicates),selected=result,gaps=result['gaps'],unknownIncluded=True))


@pytest.mark.parametrize('opposite',['type','difficulty','unknown-difficulty','original','deduplicate'])
def test_manual_fixed_question_does_not_bypass_explicit_frozen_constraints(scene,opposite):
    s=scene
    count=2 if opposite=='deduplicate' else 1
    constraint=dict(questionTypes=['short_answer'],difficulties=['hard']) if opposite!='original' else {}
    rich=None;text='独立手动反约束题'
    if opposite=='original':
        with s.catalog.read_connection() as conn:
            rich=json.loads(conn.execute('SELECT content_json FROM paper_items WHERE paper_revision_id=?',(s.seed['analysis']['paperRevisionId'],)).fetchone()[0])
        text=rich['stemBlocks'][0]['text']
    q=s.question(text,qtype='fill_blank' if opposite=='type' else 'short_answer',
        difficulty='easy' if opposite=='difficulty' else 'unspecified' if opposite=='unknown-difficulty' else 'hard',rich=rich)
    practice=s.create(count=count,**constraint)
    items=[s.item(q)]
    if opposite=='deduplicate':items.append(s.item(s.question(text,difficulty='hard'),number='17',ordinal=2))
    response=s.client.patch('/api/v1/practice-sets/'+practice['practiceSetId']+'/draft',json=dict(expectedRevision=0,items=items,constraints=practice['currentRevision']['constraints']))
    if response.status_code==422:
        assert response.json()['code'] and http(s.client,'get','/practice-sets/'+practice['practiceSetId'])['revision']==0
        record('manual-constraint-'+opposite,dict(rejectedAt='save',error=response.json(),constraints=practice['currentRevision']['constraints']))
    else:
        assert response.status_code==200,response.text
        saved=response.json()
        assert saved['currentRevision']['constraints']==practice['currentRevision']['constraints']
        error=s.review(saved,422)
        current=http(s.client,'get','/practice-sets/'+practice['practiceSetId'])
        assert current['currentRevision']['state']=='draft'
        record('manual-constraint-'+opposite,dict(rejectedAt='review',error=error,constraints=current['currentRevision']['constraints']))


def multileaf(s):
    first=s.question("独立分层富题")
    second=s.question("重排至首的题")
    nodes=[dict(nodeKey="container",parentNodeKey=None,questionNo="16",ordinal=1,isScored=False,maxScore=None,knowledgePointIds=[],sourceBlockIds=["stem"]),
           dict(nodeKey="a",parentNodeKey="container",questionNo="16(1)",ordinal=2,isScored=True,maxScore="0.75",knowledgePointIds=[s.kps[0]],sourceBlockIds=["f","t"]),
           dict(nodeKey="b",parentNodeKey="container",questionNo="16(2)",ordinal=3,isScored=True,maxScore="0.50",knowledgePointIds=[s.kps[1]],sourceBlockIds=["i","a","b"])]
    practice=s.create(count=2)
    return s.review(s.save(practice,[s.item(first,ordinal=7,nodes=nodes),s.item(second,number="自定义-17",ordinal=2)]))


@pytest.mark.parametrize('invalid',['duplicate-full-number','unassigned-option','zero-max','missing-kp','missing-parent'])
def test_structure_rejects_invalid_explicit_leaf_and_no_half_draft(scene,invalid):
    s=scene;q=s.question('独立显式结构');practice=s.create(count=2 if invalid=='duplicate-full-number' else 1)
    first=s.item(q);items=[first];node=first['itemStructure']['nodes'][0]
    if invalid=='duplicate-full-number':items.append(s.item(s.question('第二独立显式结构'),number='16',ordinal=2))
    elif invalid=='unassigned-option':node['sourceBlockIds'].remove('b')
    elif invalid=='zero-max':node['maxScore']='0'
    elif invalid=='missing-kp':node['knowledgePointIds']=[]
    else:node['parentNodeKey']='missing-parent'
    error=s.save(practice,items,422)
    assert error['details']['issues'] and error['code']
    current=http(s.client,'get','/practice-sets/'+practice['practiceSetId'])
    assert current==practice
    record('structure-invalid-'+invalid,dict(error=error,noDraftMutation=True))


def test_cas_multileaf_reorder_fixed_old_revision_and_real_return_mapping(scene):
    s=scene
    q=s.question("独立CAS")
    empty=s.create()
    saved=s.save(empty,[s.item(q)])
    stale=s.save(empty,[s.item(q)],409)
    assert stale['code']=='REVISION_CONFLICT'
    practice=multileaf(s)
    expected=[("自定义-17",125), ("16",None), ("16(1)",75), ("16(2)",50)]
    rev=practice['currentRevision'];rid=rev['practiceRevisionId']
    assert [(i['questionNo'],i['maxScoreUnits']) for i in rev['items']]==expected
    old=http(s.client,"get",s.path(practice))
    new=http(s.client,"post",f"/practice-sets/{practice['practiceSetId']}/revisions",201,
        json={"submissionId":"new-fixed-copy","sourceRevisionId":rid})
    assert new['currentRevision']['state']=='draft' and new['currentRevision']['version']==2
    assert http(s.client,"get",s.path(practice))==old
    conversion=s.convert(practice)
    with s.catalog.read_connection() as conn:
        paper=[dict(r) for r in conn.execute("SELECT * FROM paper_items WHERE paper_revision_id=? ORDER BY ordinal",(conversion['paperRevisionId'],))]
        mappings=[dict(r) for r in conn.execute("SELECT * FROM practice_paper_item_mappings WHERE conversion_id=?",(conversion['conversionId'],))]
        assert len(mappings)==len(rev['items'])==4
        bypractice={i['practiceItemId']:i for i in rev['items']};bypaper={i['id']:i for i in paper}
        mappedids={m['practice_item_id']:m['paper_item_id'] for m in mappings}
        assert set(mappedids)==set(bypractice) and len(set(mappedids.values()))==4
        for mapping in mappings:
            original=bypractice[mapping['practice_item_id']];actual=bypaper[mapping['paper_item_id']]
            assert (actual['question_no'],actual['max_score_units'])==(original['questionNo'],original['maxScoreUnits'])
            assert actual['question_revision_id']==original['questionRevisionId']
            assert actual['parent_item_id']==mappedids.get(original['parentItemId'])
            assert json.loads(actual['content_json'])==original['content']
            expectedkp={(x['knowledgePointId'],x['knowledgeRevisionId'],x['name'],x['role']) for x in original['knowledgePoints']}
            actualkp={tuple(x) for x in conn.execute("SELECT knowledge_point_id,knowledge_revision_id,knowledge_name_snapshot,role FROM paper_item_knowledge WHERE item_id=?",(actual['id'],))}
            assert actualkp==expectedkp
    assessment=http(s.client,"get","/assessments/"+conversion['assessmentId'])
    participant=assessment['participants'][0]
    workbook=Workbook();sheet=workbook.active
    sheet.append(["学号","姓名","自定义-17","16(1)","16(2)"])
    sheet.append([participant['studentNoSnapshot'],participant['nameSnapshot'],0,.75,.25])
    data=io.BytesIO();workbook.save(data)
    imported=http(s.client,"post",f"/assessments/{conversion['assessmentId']}/score-imports",201,
        files={'file':('return.xlsx',data.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    score=http(s.client,"post","/score-imports/"+imported['importId']+"/confirm",json=dict(submissionId="independent-return-score",
        expectedImportRevision=imported['revision'],expectedAssessmentRevision=assessment['assessment']['revision'],previewVersion=imported['previewVersion'],baseScoreRevisionId=None))
    report=http(s.client,"post",f"/assessments/{conversion['assessmentId']}/analysis-runs",202,json=dict(submissionId="independent-return-analysis",
        scoreRevisionId=score['revisionId'],selectedParticipantIds=[participant['participantId']],ruleCode="any_loss_v1"))
    wait_job(s.client,report)
    evidence=http(s.client,"get","/analysis-runs/"+report['runId']+"/evidence")
    assert evidence['total']==3
    seen=set()
    literal_scores={'自定义-17':0,'16(1)':75,'16(2)':25}
    for row in evidence['items']:
        assert row['runId']==report['runId'] and row['scoreRevisionId']==score['revisionId']
        assert row['paperRevisionId']==conversion['paperRevisionId'] and row['practiceRevisionId']==rid
        original=bypractice[row['practiceItemId']]
        assert row['itemId']==mappedids[original['practiceItemId']]
        # v1.3 requires the complete final number in return evidence; never double-prefix 16(1).
        assert row['itemPath']==original['questionNo']
        assert row['scoreUnits']==literal_scores[original['questionNo']]
        assert row['maxScoreUnits']==original['maxScoreUnits'] and row['status']=='recorded'
        assert row['sourceLocator']['practiceItemId']==original['practiceItemId']
        assert row['sourceLocator']['practiceRevisionId']==rid
        actual_content=dict(row['content']);source_blocks=actual_content.pop('sourceBlocks')
        assert actual_content==original['content'] and row['knowledgePoints']==original['knowledgePoints']
        with s.catalog.read_connection() as conn:
            expected_blocks=[dict(x) for x in conn.execute('SELECT * FROM paper_source_blocks WHERE paper_revision_id=? ORDER BY ordinal',(conversion['paperRevisionId'],))]
        expected_blocks=[dict(blockId=b['id'],ordinal=b['ordinal'],block=json.loads(b['block_json']),sourceLocator=json.loads(b['locator_json']),
            disposition=b['disposition'],itemId=b['item_id'],excludeReason=b['exclude_reason']) for b in expected_blocks]
        assert source_blocks==expected_blocks
        seen.add(row['itemId'])
    assert len(seen)==3 and http(s.client,"get",s.path(practice))==old
    record("multileaf-real-return",dict(expectedNumbers=expected,conversion=conversion,mappings=mappings,newScore=score,newReport=report,evidence=evidence,oldFixedUnchanged=True))


@pytest.mark.parametrize("sql",[
    "UPDATE practice_revisions SET title_snapshot='changed' WHERE id=?",
    "UPDATE practice_selections SET reason_json='{}' WHERE practice_revision_id=?",
    "DELETE FROM practice_items WHERE practice_revision_id=?",
    "DELETE FROM practice_item_knowledge WHERE practice_revision_id=?",
    "INSERT INTO practice_selections SELECT 'extra',practice_revision_id,'extra',99,question_id,question_revision_id,question_content_hash,content_snapshot_json,metadata_snapshot_json,rich_assets_json,reason_json,source_snapshot_json,answer_state FROM practice_selections WHERE practice_revision_id=?",
    "INSERT INTO practice_items SELECT 'extra',practice_revision_id,selection_id,'extra',NULL,'99',99,is_scored,max_score_units,content_json,source_locator_json FROM practice_items WHERE practice_revision_id=?",
    "INSERT INTO practice_item_knowledge SELECT item_id,practice_revision_id,'extra','extra','extra','math','primary' FROM practice_item_knowledge WHERE practice_revision_id=? LIMIT 1",
])
def test_reviewed_children_database_sealed(scene,sql):
    s=scene;practice=s.seed['practice'];before=http(s.client,"get",s.path(practice))
    with pytest.raises(sqlite3.IntegrityError,match="IMMUTABLE_REVISION"):
        with s.catalog.write_transaction() as conn:conn.execute(sql,(practice['currentRevision']['practiceRevisionId'],))
    assert http(s.client,"get",s.path(practice))==before


def test_review_io_outside_lock_archive_interleave_rechecked(scene,monkeypatch):
    s=scene;q=s.question("独立归档交错富题");practice=s.save(s.create(),[s.item(q)])
    original=s.service.read_question_asset;events=[]
    original_recheck=s.service._active_refs
    def recheck(*args):
        assert s.app.state.publication_coordinator.busy
        events.append('inside-lock-reference-recheck')
        return original_recheck(*args)
    monkeypatch.setattr(s.service,'_active_refs',recheck)
    def asset_read(aid):
        assert not s.app.state.publication_coordinator.busy
        events.append('outside-lock-read')
        s.app.state.question_bank_service.delete_question(q['questionId'])
        return original(aid)
    monkeypatch.setattr(s.service,"read_question_asset",asset_read)
    error=s.review(practice,409)
    assert error['code']=='PRACTICE_REFERENCE_CHANGED'
    assert events and 'inside-lock-reference-recheck' in events
    actual=http(s.client,"get","/practice-sets/"+practice['practiceSetId'])
    assert actual['revision']==practice['revision'] and actual['currentRevision']['state']=='draft'
    with s.catalog.read_connection() as conn:
        assert conn.execute("SELECT count(*) FROM command_submissions WHERE submission_id=?",(s.last_review['submissionId'],)).fetchone()[0]==0
    record("review-interleaving",dict(events=events,error=error,current=actual))


@pytest.mark.parametrize("same_bytes",[True,False])
def test_whole_docx_private_parts_material_asset_identity(scene,same_bytes):
    s=scene
    rich=copy.deepcopy(s.rich);private,png=s.image("PRIVATE_ANSWER_IMAGE")
    rich['assets'].append(private);rich['answerBlocks'].append(dict(id="private-image",kind="image",assetId=private['assetId'],width=10,height=10))
    rich['sharedMaterials'][0]['blocks'].append(dict(id='material-image',kind='image',assetId=rich['assets'][0]['assetId'],width=10,height=10))
    first=s.question("第一独立完整题干",rich=rich)
    rich2=copy.deepcopy(rich)
    if same_bytes:
        # Equal material, distinct IDs: identity must follow content.
        rich2['sharedMaterials'][0]['id']='other-material-id';rich2['sharedMaterials'][0]['blocks'][0]['id']='other-block-id'
        declaration=copy.deepcopy(rich['assets'][0]);declaration['assetId']=declaration['sha256']
        payload=s.app.state.asset_store.read('blobs/'+declaration['sha256'])
        blob_id,byte_size=s.app.state.question_bank_service.blobs.write(payload)
        assert blob_id==declaration['sha256'] and byte_size==len(payload)
        with s.catalog.write_transaction() as conn:
            s.app.state.file_assets.create_in(conn,kind='attachment',asset_id=declaration['assetId'],blob_key='blobs/'+declaration['sha256'],
                sha256=declaration['sha256'],byte_size=len(payload),original_name='alias.png',media_type='image/png')
    else:declaration,_=s.image('DISTINCT_MATERIAL_IMAGE')
    rich2['assets'].append(declaration)
    rich2['sharedMaterials'][0]['blocks'][1]['assetId']=declaration['assetId']
    second=s.question("第二独立完整题干",rich=rich2)
    practice=s.review(s.save(s.create(count=2),[s.item(first,number="16",ordinal=1),s.item(second,number="17",ordinal=2)]))
    packages={}
    for variant in ('student','teacher'):
        receipt=s.export(practice,variant);meta,data=s.artifact(practice,receipt)
        parts,document,ns,text=zip_oracle(data)
        all_parts=b'\n'.join(parts.values())
        if variant=='student':
            for marker in (b'ANS_PRIVATE',b'EXPL_PRIVATE',b'PRIVATE_ANSWER_IMAGE',png):assert marker not in all_parts
        else:
            for marker in (b'ANS_PRIVATE',b'EXPL_PRIVATE',b'PRIVATE_ANSWER_IMAGE'):assert marker in all_parts
        assert text.count('共享材料唯一出现')==(1 if same_bytes else 2)
        assert '第一独立完整题干' in text and '第二独立完整题干' in text
        assert '选项甲' in text and '选项乙' in text
        assert document.findall('.//m:f',ns) and document.findall('.//w:gridSpan',ns)
        media={name:hashlib.sha256(payload).hexdigest() for name,payload in parts.items() if name.startswith('word/media/')}
        assert media
        assert all(('题号 '+n) in text for n in ('16','17'))
        packages[variant]=dict(artifact=meta,parts={n:hashlib.sha256(b).hexdigest() for n,b in parts.items()},media=media,allPartsPrivateScan='pass')
    record('whole-zip-'+str(same_bytes),packages)


def test_template_freezes_actual_roster_at_accept_safe_text_and_full_numbers(scene,monkeypatch):
    s=scene;practice=multileaf(s);conversion=s.convert(practice)
    engine=s.app.state.job_engine;queued=[];originalschedule=engine.schedule
    monkeypatch.setattr(engine,'schedule',lambda *a,**kw:queued.append((a,kw)))
    receipt=s.export(practice,'score_template',assessment=conversion['assessmentId'])
    accepted_body=copy.deepcopy(s.last_export)
    student=http(s.client,'post','/students',201,json=dict(name='=ORACLE()',studentNo='000007',classId=s.seed['classId'],joinedOn='2026-01-01'))
    detail=http(s.client,'get','/assessments/'+conversion['assessmentId'])
    http(s.client,'post',f"/assessments/{conversion['assessmentId']}/participants",json=dict(submissionId='supplement-roster',
        expectedRevision=detail['assessment']['revision'],participants=[dict(studentId=student['id'],classId=s.seed['classId'],attendance='absent',attemptNo=1)]))
    monkeypatch.setattr(engine,'schedule',originalschedule)
    # Public registered queued retry must schedule the accepted frozen job.
    retry=http(s.client,'post','/workflow-jobs/'+receipt['job']['jobId']+'/retry',json={'domain':'teaching'})
    assert retry['jobId']==receipt['job']['jobId']
    meta,data=s.artifact(practice,receipt)
    workbook=load_workbook(io.BytesIO(data));sheet=workbook['成绩']
    assert list(sheet.values)[0]==('学号','姓名','出勤','人次','自定义-17','16(1)','16(2)')
    assert sheet.max_row==2 and sheet.cell(2,1).value=='00001'
    assert sheet.cell(2,1).data_type=='s' and sheet.cell(2,1).number_format=='@'
    assert all(sheet.cell(2,n).value is None for n in (5,6,7))
    fixed=list(workbook['固定映射'].values)
    assert any(conversion['paperRevisionId'] in row for row in fixed)
    again=http(s.client,'post',s.path(practice)+'/exports',202,json=accepted_body)
    assert again['replayed'] and again['exportId']==receipt['exportId']
    second=s.export(practice,'score_template',assessment=conversion['assessmentId'])
    assert not second['reused'] and second['inputHash']!=receipt['inputHash']
    _,data2=s.artifact(practice,second)
    new=load_workbook(io.BytesIO(data2))['成绩']
    assert new.max_row==3
    added=[n for n in range(2,new.max_row+1) if new.cell(n,1).value==student['studentNo']]
    assert len(added)==1
    added_row=added[0]
    assert new.cell(added_row,1).value=='000007' and new.cell(added_row,2).value=='=ORACLE()'
    assert new.cell(added_row,1).data_type==new.cell(added_row,2).data_type=='s'
    assert new.cell(added_row,1).number_format==new.cell(added_row,2).number_format=='@'
    assert tuple(new.cell(added_row,n).value for n in range(3,8))==('absent',1,None,None,None)
    assert list(new.values)[0]==('学号','姓名','出勤','人次','自定义-17','16(1)','16(2)')
    record('actual-roster-template',dict(accepted=receipt,queueCaptured=len(queued),publicRetry=retry,oldRows=list(sheet.values),fixedMapping=fixed,newRows=list(new.values)))


def test_all_success_replay_before_reference_checks_and_collision(scene):
    s=scene;q=s.question('独立重放题');created=s.create();create_body=copy.deepcopy(s.last_create)
    practice=s.review(s.save(created,[s.item(q)]));review_body=copy.deepcopy(s.last_review)
    conversion=s.convert(practice);conversion_body=copy.deepcopy(s.last_conversion)
    export=s.export(practice);export_body=copy.deepcopy(s.last_export);s.artifact(practice,export)
    s.app.state.question_bank_service.delete_question(q['questionId'])
    with s.app.state.knowledge.write_transaction() as conn:conn.execute("UPDATE knowledge_points SET status='archived' WHERE id=?",(s.kps[0],))
    copies=[http(s.client,'post','/practice-sets',201,json=create_body),
        http(s.client,'post',f"/practice-sets/{practice['practiceSetId']}/review",json=review_body),
        http(s.client,'post',s.path(practice)+'/assessments',201,json=conversion_body),
        http(s.client,'post',s.path(practice)+'/exports',202,json=export_body)]
    assert all(v['replayed'] for v in copies)
    assert copies[0]['practiceSetId']==created['practiceSetId'] and copies[2]['conversionId']==conversion['conversionId'] and copies[3]['exportId']==export['exportId']
    conflict=dict(conversion_body,title='异包')
    error=http(s.client,'post',s.path(practice)+'/assessments',409,json=conflict)
    assert error['code']=='SUBMISSION_CONFLICT'
    fresh=dict(conversion_body,submissionId='fresh-after-archive')
    assert http(s.client,'post',s.path(practice)+'/assessments',409,json=fresh)['code']=='KNOWLEDGE_ARCHIVED'
    record('success-replay',dict(copies=copies,collision=error))


@pytest.mark.parametrize('stage',['after-t30','mapping-insert'])
def test_actual_conversion_fault_rolls_all_half_objects(scene,monkeypatch,stage):
    s=scene;practice=s.reviewed(text='独立转换故障')
    tables=('papers','paper_revisions','paper_items','paper_item_knowledge','paper_source_blocks','assessments','assessment_classes','assessment_participants',
            'practice_conversions','practice_paper_item_mappings','file_assets','command_submissions')
    before=s.counts(tables)
    if stage=='after-t30':
        original=s.service.assessment_service.create_in
        def fault(conn,payload):
            original(conn,payload)
            raise AppError('独立故障注入',code='QA_CONVERSION_FAULT',status_code=500)
        monkeypatch.setattr(s.service.assessment_service,'create_in',fault)
        error=s.convert(practice,expected=500)
        assert error['code']=='QA_CONVERSION_FAULT'
    else:
        original=s.catalog.write_transaction
        @contextmanager
        def denied():
            with original() as conn:
                conn.set_authorizer(lambda action,table,*args:sqlite3.SQLITE_DENY if action==sqlite3.SQLITE_INSERT and table=='practice_paper_item_mappings' else sqlite3.SQLITE_OK)
                yield conn
        monkeypatch.setattr(s.catalog,'write_transaction',denied)
        from app.contracts.b4 import PracticeConversionRequest
        body=PracticeConversionRequest(submissionId='mapping-fault',title='故障转换',heldOn='2026-10-02',classIds=[s.seed['classId']],participants=[dict(studentId=s.seed['studentId'],classId=s.seed['classId'])])
        with pytest.raises(sqlite3.DatabaseError):s.service.convert(practice['practiceSetId'],practice['currentRevision']['practiceRevisionId'],body)
    assert s.counts(tables)==before
    record('conversion-fault-'+stage,dict(before=before,after=s.counts(tables),stage=stage))


def test_export_publish_fault_public_retry_frozen_and_atomic(scene,monkeypatch):
    s=scene;practice=s.reviewed(text='独立导出故障');engine=s.app.state.job_engine
    originalschedule=engine.schedule;monkeypatch.setattr(engine,'schedule',lambda *a,**kw:None)
    receipt=s.export(practice);store=engine.store('teaching');record_before=store.get(receipt['job']['jobId']);frozen=copy.deepcopy(record_before.frozen_input)
    baseline=s.counts(('export_artifacts','file_assets'))
    originalcreate=s.service.file_assets.create_in
    def fail(conn,**kw):
        result=originalcreate(conn,**kw)
        if kw.get('kind')=='export':raise RuntimeError('independent fault after file metadata')
        return result
    monkeypatch.setattr(s.service.file_assets,'create_in',fail)
    spec=s.app.state.job_executors.spec_for('teaching','export')
    failed=s.client.portal.call(engine.run_job,'teaching',record_before.job_id,spec.factory(record_before))
    assert failed.state=='failed' and s.counts(('export_artifacts','file_assets'))==baseline
    assert store.get(failed.job_id).frozen_input==frozen
    monkeypatch.setattr(s.service.file_assets,'create_in',originalcreate)
    monkeypatch.setattr(engine,'schedule',originalschedule)
    retry=http(s.client,'post','/workflow-jobs/'+failed.job_id+'/retry',json={'domain':'teaching'})
    assert retry['jobId']==failed.job_id
    finished=wait_job(s.client,receipt)
    assert finished['attempt']==failed.attempt+1
    meta,_=s.artifact(practice,receipt)
    assert s.counts(('export_artifacts','file_assets'))=={k:v+1 for k,v in baseline.items()}
    assert store.get(failed.job_id).frozen_input==frozen
    record('export-publish-fault',dict(baseline=baseline,failed=failed.view().model_dump(by_alias=True),retry=retry,finished=finished,artifact=meta,frozenInput=frozen))


def test_export_asset_read_failure_public_retry_uses_same_frozen_revision(scene,monkeypatch):
    s=scene;practice=s.reviewed(text='独立资产故障导出');engine=s.app.state.job_engine;store=engine.store('teaching')
    originalschedule=engine.schedule;monkeypatch.setattr(engine,'schedule',lambda *a,**kw:None)
    receipt=s.export(practice);record_before=store.get(receipt['job']['jobId']);frozen=copy.deepcopy(record_before.frozen_input)
    baseline=s.counts(('export_artifacts','file_assets'))
    sha=practice['currentRevision']['items'][0]['content']['assets'][0]['sha256'];path=s.app.state.asset_store.path_of('blobs/'+sha);original=path.read_bytes()
    try:
        path.write_bytes(b'own test corrupt source asset')
        spec=s.app.state.job_executors.spec_for('teaching','export')
        failed=s.client.portal.call(engine.run_job,'teaching',record_before.job_id,spec.factory(record_before))
        assert failed.state=='failed' and s.counts(('export_artifacts','file_assets'))==baseline
    finally:path.write_bytes(original)
    monkeypatch.setattr(engine,'schedule',originalschedule)
    retry=http(s.client,'post','/workflow-jobs/'+record_before.job_id+'/retry',json={'domain':'teaching'})
    meta,data=s.artifact(practice,receipt)
    assert store.get(record_before.job_id).frozen_input==frozen
    assert hashlib.sha256(data).hexdigest()==meta['sha256']
    record('export-source-asset-fault',dict(failed=failed.view().model_dump(by_alias=True),retry=retry,artifact=meta,frozenInputUnchanged=True))


@pytest.mark.parametrize('terminal',['cancel','lease-expired'])
def test_real_export_output_cannot_publish_after_cancel_or_original_lease_loss(scene,monkeypatch,terminal):
    s=scene;practice=s.reviewed(text='独立迟到导出'+terminal);engine=s.app.state.job_engine;store=engine.store('teaching')
    monkeypatch.setattr(engine,'schedule',lambda *a,**kw:None)
    receipt=s.export(practice);jid=receipt['job']['jobId'];lease=store.claim(jid);current=store.get(jid)
    frozen=FrozenJob(job_id=jid,domain=current.domain,kind=current.kind,attempt=lease.attempt,input=current.frozen_input,model_snapshot=current.model_snapshot,input_hash=current.input_hash)
    spec=s.app.state.job_executors.spec_for('teaching','export')
    outcome=s.client.portal.call(spec.factory(current),frozen,JobContext(store,jid,lease))
    assert outcome.publish is not None
    baseline=s.counts(('export_artifacts','file_assets'))
    if terminal=='cancel':
        store.request_cancel(jid)
        result=store.complete(jid,lease,result=outcome.result,publish=outcome.publish)
        assert result.state=='cancelled'
    else:
        expiry=datetime.fromisoformat(lease.expires_at.replace('Z','+00:00'))+timedelta(seconds=1)
        monkeypatch.setattr(store,'_now',lambda:expiry.astimezone(UTC).isoformat().replace('+00:00','Z'))
        assert jid in store.reconcile_interrupted()
        second=store.retry(jid);newlease=store.claim(jid)
        with pytest.raises(AppError) as lost:store.complete(jid,lease,result=outcome.result,publish=outcome.publish)
        assert lost.value.code=='LEASE_LOST'
        assert store.get(jid).attempt==newlease.attempt and newlease.token!=lease.token
        store.request_cancel(jid);store.mark_cancelled(jid,newlease)
    assert s.counts(('export_artifacts','file_assets'))==baseline
    record('late-export-'+terminal,dict(originalLease=lease,after=store.get(jid).view().model_dump(by_alias=True),baseline=baseline,metadataAfter=s.counts(('export_artifacts','file_assets'))))


def test_artifact_owner_actual_blob_hash_and_job_state_gate(scene):
    s=scene;meta=s.seed['exports'][0]
    from app.services.export_artifacts import ExportArtifactsService
    wrong=ExportArtifactsService(s.catalog,assets=s.app.state.asset_store,file_assets=s.app.state.file_assets,owner_id='different-owner')
    with pytest.raises(AppError) as absent:wrong.download(meta['artifactId'])
    assert absent.value.code=='EXPORT_ARTIFACT_NOT_FOUND'
    path=s.app.state.asset_store.path_of('blobs/'+meta['sha256']);original=path.read_bytes()
    try:
        path.write_bytes(b'corrupt managed export')
        response=s.client.get(meta['downloadUrl'])
        assert response.status_code==500 and response.json()['code']
    finally:path.write_bytes(original)
    assert hashlib.sha256(s.client.get(meta['downloadUrl']).content).hexdigest()==meta['sha256']
    with s.catalog.write_transaction() as conn:
        conn.execute("UPDATE workflow_jobs SET state='failed' WHERE id=(SELECT job_id FROM practice_exports WHERE id=?)",(meta['exportId'],))
    try:
        response=s.client.get('/api/v1/export-artifacts/'+meta['artifactId'])
        assert response.status_code==500 and response.json()['code']=='EXPORT_ARTIFACT_CORRUPT'
        response=s.client.get(meta['downloadUrl'])
        assert response.status_code==500 and response.json()['code']=='EXPORT_ARTIFACT_CORRUPT'
    finally:
        with s.catalog.write_transaction() as conn:
            conn.execute("UPDATE workflow_jobs SET state='succeeded' WHERE id=(SELECT job_id FROM practice_exports WHERE id=?)",(meta['exportId'],))
