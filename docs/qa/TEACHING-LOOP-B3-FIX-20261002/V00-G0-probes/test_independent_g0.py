"""Independent frozen-candidate acceptance. No product writes or network servers.

Environment isolation precedes all app/main and test helper imports. Helpers are
used for setup only; fault interleavings and assertions here are independently
constructed and assert correct product behaviour.
"""
from __future__ import annotations

import asyncio
import atexit
from dataclasses import replace
import hashlib
import json
import os
from pathlib import Path
import sqlite3
import sys
import tempfile
import threading
import time

_temporary = tempfile.TemporaryDirectory(prefix="zqky-v00-g0-bootstrap-")
atexit.register(_temporary.cleanup)
os.environ["ZQKY_DATA_DIR"] = str(Path(_temporary.name) / "data")
os.environ["ZQKY_ENV"] = "test"
os.environ.pop("ZQKY_CREDENTIALS_FILE", None)
WORKSPACE = Path(__file__).resolve().parents[4]
sys.path.insert(0, str(WORKSPACE / "apps" / "api"))

import pytest
from app.core.exceptions import AppError
from app.core.migrations import REGISTERED_MIGRATIONS, apply_migrations
from app.core.migrations.base import Migration
from app.core.migrations.teaching import _ASSESSMENTS_REBUILD
from app.core.sqlite import connect
from app.contracts.papers import PaperConfirmRequest, PaperDraftPatchRequest
from app.services.jobs.engine import JobEngine, JobOutcome, FrozenJob
from app.services.model_runtime import fingerprint_of_handle
from app.services.papers.proposals import ProposalRunner
from tests.papers_support import PapersHarness, build_paper_docx, items_payload, make_handle, FakeProvider
from tests.test_question_generation import GenerationHarness, _generation_request, generation_reply
from tests.test_question_bank import LOCAL_PROFILE, FakeLLMProvider
from tests.test_question_bank_confirm import ONE_QUESTION_DOC, reviewed_draft, confirm
from tests.test_retry_dispatch import RetryHarness, _wait_job
from tests.test_question_lease_batches import LeaseFixture
from tests.assessments_support import AssessmentsHarness


def emit(case, **facts):
    print(json.dumps({"case": case, **facts}, ensure_ascii=False))


def snapshot_db(path):
    connection = connect(path)
    try:
        tables = [r[0] for r in connection.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")]
        rows = {t: sorted([tuple(r) for r in connection.execute(f'SELECT * FROM "{t}"')], key=repr) for t in tables}
        return rows
    finally:
        connection.close()


def prepare_paper(tmp_path, *, unknown=False):
    h = PapersHarness(tmp_path)
    s = h.service()
    imported = h.import_docx(s, build_paper_docx(tmp_path / 'paper.docx', with_unknown_object=unknown), title='FROZEN TITLE')
    point = h.add_point('INDEPENDENT', '独立知识点')
    s.patch_draft(imported.paper.paper_id, PaperDraftPatchRequest.model_validate({
        'expectedRevision': 0,
        'items': items_payload(imported.revision.items, knowledge={i.question_no: [point.point_id] for i in imported.revision.items if i.is_scored}),
    }))
    return h, s, imported, point


def test_public_retry_schedule_failure_is_recoverable_and_duplicate_clicks_do_not_repeat(tmp_path, monkeypatch):
    h = RetryHarness(tmp_path, replies=['INVALID JSON', generation_reply()])
    try:
        job_id = h.start_generation()
        failed = _wait_job(h.client, job_id)
        assert failed['state'] == 'failed' and failed['attempt'] == 1
        frozen = h.store().get(job_id)
        original = h.service.job_engine.schedule
        def unavailable(*args, **kwargs):
            raise RuntimeError('independent dispatch outage')
        monkeypatch.setattr(h.service.job_engine, 'schedule', unavailable)
        with pytest.raises(RuntimeError, match='independent dispatch outage'):
            h.retry(job_id)
        stranded = h.store().get(job_id)
        assert stranded.state == 'queued' and stranded.attempt == 1
        assert len(h.provider.calls) == 1
        monkeypatch.setattr(h.service.job_engine, 'schedule', original)
        response = h.retry(job_id)
        assert response.status_code == 200, response.text
        duplicate = h.retry(job_id)
        assert duplicate.status_code in (200, 409)
        final = _wait_job(h.client, job_id)
        assert final['state'] == 'succeeded' and final['attempt'] == 2
        assert len(h.provider.calls) == 2
        now = h.store().get(job_id)
        assert (now.frozen_input, now.model_snapshot, now.input_hash) == (frozen.frozen_input, frozen.model_snapshot, frozen.input_hash)
        assert h.retry(job_id).status_code == 409
        emit('RV01 dispatch fault recovery', initial=failed['state'], orphaned=stranded.state, final=final['state'], attempt=2, calls=2)
    finally:
        h.close()


def test_public_retry_after_actual_application_restart_is_explicit(tmp_path):
    h = RetryHarness(tmp_path, replies=['INVALID JSON'])
    job_id = h.start_generation()
    assert _wait_job(h.client, job_id)['state'] == 'failed'
    h.store().retry(job_id)
    frozen = h.store().get(job_id)
    h.close()
    reopened = RetryHarness(tmp_path, replies=[generation_reply()])
    try:
        time.sleep(.05)
        persisted = reopened.store().get(job_id)
        assert persisted.state == 'queued' and persisted.attempt == 1
        assert reopened.provider.calls == []
        assert reopened.retry(job_id).status_code == 200
        final = _wait_job(reopened.client, job_id)
        assert final['state'] == 'succeeded' and final['attempt'] == 2
        assert len(reopened.provider.calls) == 1
        after = reopened.store().get(job_id)
        assert (after.frozen_input, after.model_snapshot, after.input_hash) == (frozen.frozen_input, frozen.model_snapshot, frozen.input_hash)
        emit('RV01 application restart', autoReplayCalls=0, explicitCalls=1, finalAttempt=2)
    finally:
        reopened.close()


@pytest.mark.parametrize('loss', ['wrong_token','wrong_attempt','missing_token','boolean_attempt','expired_exact','expired','cancel_requested','interrupted','cancelled','failed','succeeded','takeover'])
def test_intermediate_batch_and_failure_checkpoint_zero_write_for_every_invalid_identity(tmp_path, loss):
    f = LeaseFixture(tmp_path)
    job = f.job()
    lease = f.store.claim(job.job_id)
    token, attempt = lease.token, lease.attempt
    if loss == 'wrong_token': token = 'other-token'
    elif loss == 'wrong_attempt': attempt += 1
    elif loss == 'missing_token': token = ''
    elif loss == 'boolean_attempt': attempt = True
    elif loss in ('expired_exact', 'expired'):
        f.clock.advance(30 if loss == 'expired_exact' else 31)
    elif loss == 'cancel_requested': f.store.request_cancel(job.job_id)
    elif loss == 'takeover':
        f.clock.advance(31)
        fresh = f.store.claim(job.job_id)
        current, created, failure = f.catalog.record_organize_batch(job.job_id, batch_index=0, next_batch_index=1,
            draft_id=f.draft.draft_id, base_draft_revision=f.draft.revision, proposed_content=dict(f.draft.content),
            lease_attempt=fresh.attempt, lease_token=fresh.token, now=f.clock.now())
        assert created is not None and failure is None
    else:
        with f.catalog.write_transaction() as c:
            c.execute('UPDATE question_jobs SET state=? WHERE id=?', (loss, job.job_id))
    before = snapshot_db(f.catalog.db_path)
    _, created, _ = f.catalog.record_organize_batch(job.job_id, batch_index=9, next_batch_index=10,
        draft_id=f.draft.draft_id, base_draft_revision=f.draft.revision, proposed_content=dict(f.draft.content),
        lease_attempt=attempt, lease_token=token, now=f.clock.now())
    assert created is None
    f.catalog.record_organize_failure(job.job_id, code='LATE_ERROR', message='independent late failure',
        lease_attempt=attempt, lease_token=token, now=f.clock.now())
    assert snapshot_db(f.catalog.db_path) == before
    emit('RV05/R09 zero write', loss=loss, byteEquivalentRows=True)


@pytest.mark.parametrize('finalizer', ['failed', 'cancelled'])
def test_late_terminal_finalizers_cannot_revive_expired_execution(tmp_path, finalizer):
    f = LeaseFixture(tmp_path)
    j = f.job(); old = f.store.claim(j.job_id)
    f.clock.advance(30)
    before = snapshot_db(f.catalog.db_path)
    if finalizer == 'failed': f.store.fail_if_current_lease(j.job_id, old, code='LATE', message='late')
    else: f.store.mark_cancelled(j.job_id, old)
    assert snapshot_db(f.catalog.db_path) == before


@pytest.mark.parametrize('fault', ['ordinary','cancel','takeover'])
@pytest.mark.parametrize('domain',['question','knowledge','teaching'])
@pytest.mark.parametrize('error_kind',['app_error','internal'])
def test_publish_partial_business_write_rolls_back_before_original_lease_failure_cas(tmp_path, monkeypatch, fault,domain,error_kind):
    h = GenerationHarness(tmp_path, with_client=False)
    store = h.service.job_engine.store(domain)
    catalog=getattr(h.app.state,{'question':'question_bank','knowledge':'knowledge','teaching':'teaching'}[domain])
    table={'question':'question_jobs','knowledge':'knowledge_jobs','teaching':'workflow_jobs'}[domain]
    with catalog.write_transaction() as c: c.execute('CREATE TABLE independent_sentinel(value TEXT)')
    j = store.create(kind={'question':'generate','knowledge':'suggestion','teaching':'export'}[domain], frozen_input={'probe': True})
    original = store.fail_if_current_lease
    leases = []
    def finalizer(job_id, lease, **kwargs):
        leases.append(lease)
        with catalog.write_transaction() as c:
            assert c.execute('SELECT count(*) FROM independent_sentinel').fetchone()[0] == 0
            if fault == 'takeover':
                c.execute(f"UPDATE {table} SET lease_expires_at='2000-01-01T00:00:00Z' WHERE id=?", (job_id,))
        if fault == 'cancel': store.request_cancel(job_id)
        if fault == 'takeover':
            fresh = store.claim(job_id)
            with catalog.write_transaction() as c:
                c.execute(f'UPDATE {table} SET checkpoint_json=? WHERE id=?', (json.dumps({'fresh': fresh.attempt}), job_id))
        return original(job_id, lease, **kwargs)
    monkeypatch.setattr(store, 'fail_if_current_lease', finalizer)
    async def executor(frozen, ctx):
        def publish(c):
            c.execute("INSERT INTO independent_sentinel VALUES ('partial')")
            if error_kind=='app_error': raise AppError('independent publication fault', code='PROBE_PUBLISH_FAILURE', status_code=422)
            raise RuntimeError('independent internal publication fault')
        return JobOutcome(result={'published': True}, publish=publish)
    result = asyncio.run(h.service.job_engine.run_job(domain, j.job_id, executor))
    assert len(leases) == 1 and leases[0].attempt == 1
    assert snapshot_db(catalog.db_path)['independent_sentinel'] == []
    if fault == 'ordinary': assert result.state == 'failed' and result.error_code == ('PROBE_PUBLISH_FAILURE' if error_kind=='app_error' else 'JOB_FAILED')
    elif fault == 'cancel': assert result.state == 'cancelled' and result.error_code is None
    else:
        assert result.state == 'running' and result.attempt == 2
        assert result.lease_token != leases[0].token and result.checkpoint == {'fresh': 2}
        assert result.error_code is None
    assert h.service.job_engine.active_jobs == 0
    emit('publication rollback/CAS', fault=fault, domain=domain,errorKind=error_kind,state=result.state, attempt=result.attempt, sentinelRows=0)


def test_generation_waited_slot_cancel_zero_provider_and_next_job_runs(tmp_path):
    async def run():
        p = FakeLLMProvider(replies=[generation_reply()])
        h = GenerationHarness(tmp_path, provider=p, with_client=False)
        engine = h.service.job_engine
        engine._ensure_primitives(); await engine._model.acquire()
        try:
            j = await h.service.create_generation_job(_generation_request([]))
            store = engine.store('question')
            for _ in range(200):
                if store.get(j.jobId).state == 'running': break
                await asyncio.sleep(.005)
            assert store.get(j.jobId).state == 'running' and p.calls == []
            store.request_cancel(j.jobId)
            task = engine._tracked[('question', j.jobId)]
            engine._model.release()
            result = await asyncio.wait_for(task, 3)
            assert result.state == 'cancelled' and p.calls == []
            assert h.row_count('question_imports') == 0
            next_job = await h.service.create_generation_job(_generation_request([]))
            next_task = engine._tracked[('question', next_job.jobId)]
            next_result = await asyncio.wait_for(next_task, 3)
            assert next_result.state == 'succeeded' and len(p.calls) == 1
            assert engine.active_jobs == 0
            emit('R10 waited cancel', cancelledCalls=0, nextCalls=1, nextState=next_result.state)
        finally:
            await engine.shutdown(); h.close()
    asyncio.run(run())


def test_generation_drift_after_waited_slot_fails_before_provider(tmp_path):
    async def run():
        p = FakeLLMProvider(replies=[generation_reply()]); h = GenerationHarness(tmp_path, provider=p, with_client=False)
        engine = h.service.job_engine; engine._ensure_primitives(); await engine._model.acquire()
        try:
            h.resolver.default = h.handle(model_id='frozen-A', provider=p)
            j = await h.service.create_generation_job(_generation_request([]))
            frozen = engine.store('question').get(j.jobId)
            h.resolver.default = h.handle(model_id='changed-B', provider=p)
            task = engine._tracked[('question', j.jobId)]; engine._model.release()
            result = await asyncio.wait_for(task, 3)
            assert result.state == 'failed' and result.error_code == 'MODEL_CONFIG_DRIFT'
            assert p.calls == [] and h.row_count('question_imports') == 0
            assert result.model_snapshot == frozen.model_snapshot
            emit('RV04 waited drift', providerCalls=0, error=result.error_code)
        finally: await engine.shutdown(); h.close()
    asyncio.run(run())


def test_paper_arbitrary_resolution_and_necessary_content_loss_are_rejected(tmp_path):
    h, s, imported, _point = prepare_paper(tmp_path, unknown=True)
    issue = next(i for i in imported.revision.issues if i.severity == 'blocking')
    before = snapshot_db(h.catalog.db_path)
    with pytest.raises(Exception):
        PaperDraftPatchRequest.model_validate({'expectedRevision':1,'issues':[{'issueId':issue.issue_id,'status':'resolved','resolution':{'anything':1}}]})
    with pytest.raises(AppError) as excluded:
        s.patch_draft(imported.paper.paper_id, PaperDraftPatchRequest.model_validate({'expectedRevision':1,
            'issues':[{'issueId':issue.issue_id,'status':'excluded','resolution':{'kind':'exclude','reason':'hide the required object'}}]}))
    assert excluded.value.status_code == 422
    assert snapshot_db(h.catalog.db_path) == before
    payload = items_payload(imported.revision.items, knowledge={i.question_no:[_point.point_id] for i in imported.revision.items if i.is_scored})
    for i in payload: i['content'] = {}
    # Incomplete editable drafts are allowed; empty scored content must fail the
    # confirmation gate. Use a clean import so open unsupported-object issues
    # cannot hide a missing-content bug.
    clean = tmp_path / 'clean'; clean.mkdir()
    h2, s2, imported2, point2 = prepare_paper(clean)
    payload = items_payload(imported2.revision.items, knowledge={i.question_no:[point2.point_id] for i in imported2.revision.items if i.is_scored})
    for i in payload: i['content'] = {}
    s2.patch_draft(imported2.paper.paper_id, PaperDraftPatchRequest.model_validate({'expectedRevision':1,'items':payload}))
    before2 = snapshot_db(h2.catalog.db_path)
    with pytest.raises(AppError) as empty:
        s2.confirm(imported2.paper.paper_id, PaperConfirmRequest(expectedRevision=2, submissionId='empty-must-not-publish'))
    assert empty.value.status_code == 422 and empty.value.code == 'ITEM_STEM_MISSING'
    assert snapshot_db(h2.catalog.db_path) == before2
    emit('RV02 required content', arbitraryRejected=True, excludedRejected=True, emptyRejected=True)


def test_paper_valid_supplement_and_confirmed_title_remain_historical(tmp_path):
    h, s, imported, point = prepare_paper(tmp_path, unknown=True)
    issue = next(i for i in imported.revision.issues if i.severity == 'blocking')
    target = next(b.block_id for b in imported.revision.blocks if b.kind=='paragraph' and b.disposition=='item')
    revised = s.patch_draft(imported.paper.paper_id, PaperDraftPatchRequest.model_validate({'expectedRevision':1,
        'issues':[{'issueId':issue.issue_id,'status':'resolved','resolution':{'kind':'supplement_text','targetBlockId':target,'text':'教师补录：右图为半径2的四分之一圆，阴影为圆内三角形之外。'}}]}))
    result = s.confirm(imported.paper.paper_id, PaperConfirmRequest(expectedRevision=2, submissionId='independent-title'))
    before = s.get_revision_content(imported.paper.paper_id, result.paper_revision_id).model_dump(mode='json')
    reader = s.confirmed_reader().read(result.paper_revision_id)
    s.patch_draft(imported.paper.paper_id, PaperDraftPatchRequest(expectedRevision=3, title='A NEW DRAFT TITLE'))
    assert s.get_revision_content(imported.paper.paper_id, result.paper_revision_id).model_dump(mode='json') == before
    assert s.confirmed_reader().read(result.paper_revision_id) == reader
    assert reader.title == 'FROZEN TITLE'
    assert '四分之一圆' in json.dumps(before,ensure_ascii=False)
    emit('RV02/RV11 supplemental historical snapshot', title=reader.title, unchanged=True)


def test_archive_after_draft_blocks_new_paper_confirm_and_zero_writes(tmp_path):
    h, s, imported, point = prepare_paper(tmp_path)
    h.archive_point(point.point_id)
    before = snapshot_db(h.catalog.db_path)
    with pytest.raises(AppError) as err:
        s.confirm(imported.paper.paper_id, PaperConfirmRequest(expectedRevision=1, submissionId='archived-block'))
    assert err.value.code == 'KNOWLEDGE_ARCHIVED'
    assert snapshot_db(h.catalog.db_path) == before


def test_publication_lock_spans_paper_validation_and_domain_commit(tmp_path, monkeypatch):
    h, s, imported, point = prepare_paper(tmp_path)
    entered, release, archived = threading.Event(), threading.Event(), threading.Event()
    original = s._apply_confirm
    def paused(c, **kwargs):
        entered.set(); assert release.wait(3)
        return original(c, **kwargs)
    monkeypatch.setattr(s, '_apply_confirm', paused)
    results, errors = [], []
    def publish():
        try: results.append(s.confirm(imported.paper.paper_id, PaperConfirmRequest(expectedRevision=1, submissionId='lock-span')))
        except Exception as e: errors.append(e)
    def archive():
        with h.coordinator.publication(operation='independent-archive'):
            h.archive_point(point.point_id); archived.set()
    t1=threading.Thread(target=publish); t1.start(); assert entered.wait(3)
    t2=threading.Thread(target=archive); t2.start()
    time.sleep(.05); assert not archived.is_set()
    release.set(); t1.join(3); t2.join(3)
    assert not t1.is_alive() and not t2.is_alive() and not errors
    assert results[0].state == 'confirmed' and archived.is_set()
    # Post-archive replay must return the old fact rather than rerun validation.
    replay=s.confirm(imported.paper.paper_id, PaperConfirmRequest(expectedRevision=1, submissionId='lock-span'))
    assert replay.replayed and replay.paper_revision_id == results[0].paper_revision_id
    emit('RV06 coordinator span', archiveBlockedDuringCommit=True, historicalReplay=True)


def test_question_archive_gate_and_inherited_cross_subject_links(tmp_path):
    h=GenerationHarness(tmp_path)
    try:
        point=h.create_point()
        detail=h.upload_sample(ONE_QUESTION_DOC, subjectId='math'); draft=detail['drafts'][0]
        linked=h.patch_links(draft['draftId'],draft['revision'],[{'knowledgePointId':point['pointId']}]); assert linked.status_code==200
        reviewed=h.patch_draft(linked.json()); assert reviewed.status_code==200, reviewed.text
        current=reviewed.json()
        before=snapshot_db(h.catalog.db_path)
        moved={**current['metadata'],'subjectId':'chinese'}
        rejected=h.client.patch('/api/v1/question-drafts/'+draft['draftId'],json={'expectedRevision':current['revision'],'content':current['content'],'metadata':moved})
        assert rejected.status_code==422 and rejected.json()['code']=='KNOWLEDGE_REFERENCE_INVALID'
        assert snapshot_db(h.catalog.db_path)==before
        p=h.client.get('/api/v1/knowledge-points/'+point['pointId']).json()
        archived=h.client.post('/api/v1/knowledge-points/'+point['pointId']+'/archive',json={'expectedRevision':p['revision']})
        assert archived.status_code==200
        refused=confirm(h,detail['importId'],[{'draftId':draft['draftId'],'expectedDraftRevision':current['revision']}],submission_id='new-after-archive')
        assert refused.status_code==409 and refused.json()['code']=='KNOWLEDGE_ARCHIVED'
        assert snapshot_db(h.catalog.db_path)==before
        emit('RV06/RV07 question guards', archiveCode=refused.json()['code'], inheritedCode=rejected.json()['code'])
    finally: h.close()


def test_assessment_date_invalidation_and_explicit_past_confirmation(tmp_path):
    with AssessmentsHarness(tmp_path) as h:
        paper=h.create_confirmed_paper(tag='independent-date'); klass=h.create_class(code='DATE')
        student=h.create_student(name='甲',student_no='010001',class_id=klass['id'],joined_on='2026-09-20')
        body=h.create_body(paper,submission_id='date-create',class_ids=[klass['id']],participants=[h.participant(student['id'],klass['id'])],held_on='2026-09-30')
        created=h.client.post('/api/v1/assessments',json=body); assert created.status_code==201, created.text
        a=created.json()['assessment']; memberships=h.memberships_snapshot()
        bad=h.client.patch('/api/v1/assessments/'+a['assessmentId'],json={'expectedRevision':a['revision'],'heldOn':'2026-09-01'})
        assert bad.status_code==422 and bad.json()['code']=='PARTICIPANT_CLASS_UNCONFIRMED'
        expected_detail={k:v for k,v in created.json().items() if k!='replayed'}
        assert h.client.get('/api/v1/assessments/'+a['assessmentId']).json()==expected_detail
        confirmed=h.client.post('/api/v1/assessments/'+a['assessmentId']+'/participants',json={'expectedRevision':a['revision'],'submissionId':'explicit-membership','participants':[h.participant(student['id'],klass['id'],attempt_no=1,class_confirmed=True,class_confirmation_note='名单今日导入，但教师确认过去本次考试在本班')]})
        assert confirmed.status_code==200, confirmed.text
        rev=confirmed.json()['assessment']['revision']
        good=h.client.patch('/api/v1/assessments/'+a['assessmentId'],json={'expectedRevision':rev,'heldOn':'2026-09-01'})
        assert good.status_code==200, good.text
        assert h.memberships_snapshot()==memberships
        emit('RV08 date', invalidStatus=422, explicitlyConfirmedStatus=200, membershipHistoryUnchanged=True)


def test_full_unassigned_rich_source_blocks_and_scoped_asset_bytes(tmp_path):
    from docx import Document
    from docx.oxml import parse_xml
    from tests.papers_support import PNG_BYTES
    from io import BytesIO
    document=Document(); document.add_paragraph('独立无题号原文，必须保留全部内容')
    table=document.add_table(rows=2,cols=3); table.cell(0,0).merge(table.cell(0,2)).text='独立合并表头'
    table.cell(1,0).text='甲'; table.cell(1,1).text='乙'; table.cell(1,2).text='丙'
    p=document.add_paragraph(); p._p.append(parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:f><m:num><m:r><m:t>x</m:t></m:r></m:num><m:den><m:r><m:t>2</m:t></m:r></m:den></m:f></m:oMath>'))
    document.add_picture(BytesIO(PNG_BYTES)); path=tmp_path/'unassigned.docx'; document.save(path)
    h=PapersHarness(tmp_path); s=h.service(); imported=h.import_docx(s,path)
    assert imported.revision.items==[]
    views={b.block_id:b for b in imported.revision.blocks}
    raw=h.raw_rows('SELECT id,block_json FROM paper_source_blocks WHERE paper_revision_id=?',[imported.paper.current_revision_id])
    assert all(views[r['id']].content==json.loads(r['block_json']) for r in raw)
    assert {'paragraph','table','formula','image'} <= {b.kind for b in views.values()}
    assert all(b.disposition=='unassigned' for b in views.values())
    tbl=next(b for b in views.values() if b.kind=='table'); assert any(c.get('colSpan',1)==3 for c in tbl.content['cells'])
    formula=next(b for b in views.values() if b.kind=='formula'); assert '<m:f>' in formula.content['ommlXml']
    image=next(b for b in views.values() if b.kind=='image')
    data,mime=s.get_revision_asset_content(imported.paper.paper_id,imported.paper.current_revision_id,image.content['assetId'])
    assert data==PNG_BYTES and mime=='image/png'
    unrelated=h.assets.store_original(PNG_BYTES+b'new',media_type='image/png',original_name='other.png')
    with pytest.raises(AppError): s.get_revision_asset_content(imported.paper.paper_id,imported.paper.current_revision_id,unrelated.blob_key)
    emit('RV03 full unassigned source', blocks=len(raw), exactStoredContent=True, imageBytes=True, unreferencedRejected=True)


def test_formal_question_subject_change_requires_explicit_link_replacement_and_keeps_old_revision(tmp_path):
    h=GenerationHarness(tmp_path)
    try:
        point=h.create_point(); detail=h.upload_sample(ONE_QUESTION_DOC,subjectId='math'); draft=detail['drafts'][0]
        linked=h.patch_links(draft['draftId'],draft['revision'],[{'knowledgePointId':point['pointId']}]); assert linked.status_code==200
        current=h.patch_draft(linked.json()).json()
        result=confirm(h,detail['importId'],[{'draftId':current['draftId'],'expectedDraftRevision':current['revision']}],submission_id='formal-links')
        assert result.status_code==200 and result.json()['failures']==[]
        qid=result.json()['confirmedQuestionIds'][0]; old=h.client.get('/api/v1/questions/'+qid).json(); before=snapshot_db(h.catalog.db_path)
        old_revision=h.catalog.get_question(qid).current_revision_id
        links=h.revision_link_rows(old_revision)
        body={'expectedRevision':old['revision'],'content':old['content'],'metadata':{**old['metadata'],'subjectId':'chinese'}}
        rejected=h.client.patch('/api/v1/questions/'+qid,json=body)
        assert rejected.status_code==422 and rejected.json()['code']=='KNOWLEDGE_REFERENCE_INVALID'
        assert snapshot_db(h.catalog.db_path)==before
        accepted=h.client.patch('/api/v1/questions/'+qid,json={**body,'knowledgeLinks':[]})
        assert accepted.status_code==200 and accepted.json()['knowledgeLinks']==[]
        assert h.revision_link_rows(old_revision)==links
        emit('RV07 formal subject', inheritedStatus=422, explicitClearStatus=200, historicalLinks=True)
    finally: h.close()


def test_paper_provider_drift_or_cancel_before_call_does_not_publish():
    p=FakeProvider(); old=make_handle(provider=p); changed=replace(old,model_id='changed',config=replace(old.config,modelId='changed'))
    async def resolver(_): return changed
    class Ctx:
        async def cancellation_requested(self): return False
    published=[]
    runner=ProposalRunner(resolve_model=resolver,publish_proposal=lambda conn,draft:published.append(draft))
    frozen=FrozenJob(job_id='probe',domain='teaching',kind='paper_mapping',attempt=1,input={'paperId':'p','paperRevisionId':'r','baseRevision':0,'subjectId':'math','modelProfileId':LOCAL_PROFILE,'allowedItemIds':[],'allowedKnowledgePointIds':[],'items':[]},model_snapshot={'profileId':LOCAL_PROFILE,'fingerprint':fingerprint_of_handle(old)},input_hash='probe')
    with pytest.raises(AppError) as err: asyncio.run(runner(frozen,Ctx()))
    assert err.value.code=='MODEL_CONFIG_DRIFT' and p.calls==[] and published==[]
    async def same(_): return old
    runner.resolve_model=same
    class CancelCtx:
        async def cancellation_requested(self): return True
    # The domain executor cooperatively returns an empty outcome; the shared
    # engine settles cancelled rather than publishing it.
    outcome=asyncio.run(runner(frozen,CancelCtx()))
    assert p.calls==[] and published==[]
    emit('RV04 paper drift/cancel', driftCode=err.value.code, providerCalls=0)


def test_running_restart_is_interrupted_without_automatic_model_replay(tmp_path):
    h=RetryHarness(tmp_path,replies=['INVALID JSON'])
    jid=h.start_generation(); assert _wait_job(h.client,jid)['state']=='failed'
    h.store().retry(jid); lease=h.store().claim(jid); assert lease.attempt==2
    h.close()
    reopened=RetryHarness(tmp_path,replies=[generation_reply()])
    try:
        current=reopened.store().get(jid)
        assert current.state=='interrupted' and current.attempt==2 and current.lease_token is None
        assert reopened.provider.calls==[]
        assert reopened.retry(jid).status_code==200
        final=_wait_job(reopened.client,jid)
        assert final['state']=='succeeded' and final['attempt']==3 and len(reopened.provider.calls)==1
        emit('RV01 running restart', restartState='interrupted', autoCalls=0, finalAttempt=3)
    finally: reopened.close()


@pytest.mark.parametrize('case',['drift','cancel'])
def test_knowledge_waited_slot_drift_or_cancel_zero_provider(tmp_path,case):
    from app.contracts.knowledge import KnowledgeSuggestionRequest
    from tests.test_knowledge_suggestions import FakeProvider as KPProvider, FakeResolver as KPResolver, make_handle as kp_handle, make_harness
    async def run():
        p=KPProvider(); original=kp_handle(provider=p); resolver=KPResolver({LOCAL_PROFILE:original})
        h=make_harness(tmp_path,model_resolver=resolver,with_client=False); engine=h.service.job_engine
        engine._ensure_primitives(); await engine._model.acquire()
        try:
            j=await h.service.create_suggestion_job(KnowledgeSuggestionRequest.model_validate(h.payload()))
            store=engine.store('knowledge'); frozen=store.get(j.job_id)
            for _ in range(200):
                if store.get(j.job_id).state=='running': break
                await asyncio.sleep(.005)
            assert p.calls==[]
            if case=='drift': resolver.profiles[LOCAL_PROFILE]=replace(original,model_id='changed',config=replace(original.config,modelId='changed'))
            else: store.request_cancel(j.job_id)
            task=engine._tracked[('knowledge',j.job_id)]; engine._model.release()
            result=await asyncio.wait_for(task,3)
            assert result.state==('failed' if case=='drift' else 'cancelled')
            if case=='drift': assert result.error_code=='MODEL_CONFIG_DRIFT'
            assert p.calls==[] and h.knowledge_count('knowledge_imports')==0 and h.knowledge_count('knowledge_import_rows')==0
            assert result.model_snapshot==frozen.model_snapshot and engine.active_jobs==0
            emit('RV04/R10 knowledge waited '+case,providerCalls=0,state=result.state,batches=0)
        finally: await engine.shutdown(); h.close()
    asyncio.run(run())


def test_formal_question_new_link_publication_blocks_archive_until_domain_commit(tmp_path,monkeypatch):
    h=GenerationHarness(tmp_path)
    release=threading.Event(); threads=[]
    try:
        point=h.create_point(); detail=h.upload_sample(ONE_QUESTION_DOC,subjectId='math')
        reviewed=reviewed_draft(h,detail['importId'])
        result=confirm(h,detail['importId'],[{'draftId':reviewed['draftId'],'expectedDraftRevision':reviewed['revision']}],submission_id='unlinked-formal')
        assert result.status_code==200 and result.json()['failures']==[]
        qid=result.json()['confirmedQuestionIds'][0]; current=h.client.get('/api/v1/questions/'+qid).json()
        p=h.client.get('/api/v1/knowledge-points/'+point['pointId']).json()
        entered,archived=threading.Event(),threading.Event(); outcomes={}
        original=h.catalog.patch_question
        def paused(*args,**kwargs):
            entered.set(); assert release.wait(3)
            return original(*args,**kwargs)
        monkeypatch.setattr(h.catalog,'patch_question',paused)
        def publish():
            outcomes['patch']=h.client.patch('/api/v1/questions/'+qid,json={'expectedRevision':current['revision'],'content':current['content'],'metadata':current['metadata'],'knowledgeLinks':[{'knowledgePointId':point['pointId']}]})
        def archive():
            outcomes['archive']=h.client.post('/api/v1/knowledge-points/'+point['pointId']+'/archive',json={'expectedRevision':p['revision']})
            archived.set()
        t1=threading.Thread(target=publish); threads.append(t1); t1.start(); assert entered.wait(3)
        t2=threading.Thread(target=archive); threads.append(t2); t2.start()
        time.sleep(.1); completed_before_commit=archived.is_set()
        release.set()
        for t in threads: t.join(3); assert not t.is_alive()
        emit('RV06 formal question new-link race',archiveFinishedBeforeDomainCommit=completed_before_commit,
            patchStatus=outcomes['patch'].status_code,archiveStatus=outcomes['archive'].status_code,
            links=h.client.get('/api/v1/questions/'+qid).json()['knowledgeLinks'])
        assert not completed_before_commit, 'PublicationCoordinator released after validation but before new formal link commit'
        assert outcomes['patch'].status_code==200 and outcomes['archive'].status_code==200
    finally:
        release.set()
        for t in threads: t.join(3)
        h.close()


def test_question_confirm_coordinator_spans_recheck_and_entire_domain_transaction(tmp_path,monkeypatch):
    h=GenerationHarness(tmp_path); release=threading.Event(); threads=[]
    try:
        point=h.create_point(); detail=h.upload_sample(ONE_QUESTION_DOC,subjectId='math'); draft=detail['drafts'][0]
        linked=h.patch_links(draft['draftId'],draft['revision'],[{'knowledgePointId':point['pointId']}]); assert linked.status_code==200
        current=h.patch_draft(linked.json()).json(); p=h.client.get('/api/v1/knowledge-points/'+point['pointId']).json()
        entered,archived=threading.Event(),threading.Event(); outcomes={}
        original=h.service._confirm_in_transaction
        def paused(*args,**kwargs):
            entered.set(); assert release.wait(3)
            return original(*args,**kwargs)
        monkeypatch.setattr(h.service,'_confirm_in_transaction',paused)
        def publish(): outcomes['confirm']=confirm(h,detail['importId'],[{'draftId':current['draftId'],'expectedDraftRevision':current['revision']}],submission_id='confirm-lock')
        def archive():
            outcomes['archive']=h.client.post('/api/v1/knowledge-points/'+point['pointId']+'/archive',json={'expectedRevision':p['revision']}); archived.set()
        t1=threading.Thread(target=publish); threads.append(t1); t1.start(); assert entered.wait(3)
        t2=threading.Thread(target=archive); threads.append(t2); t2.start()
        time.sleep(.1); completed_before_commit=archived.is_set(); release.set()
        for t in threads: t.join(3); assert not t.is_alive()
        assert not completed_before_commit
        assert outcomes['confirm'].status_code==200 and outcomes['confirm'].json()['failures']==[]
        assert outcomes['archive'].status_code==200
        replay=confirm(h,detail['importId'],[{'draftId':current['draftId'],'expectedDraftRevision':current['revision']}],submission_id='confirm-lock')
        assert replay.json()==outcomes['confirm'].json()
        emit('RV06 question confirm coordinator',archiveBlockedUntilCommit=True,replayAfterArchiveUnchanged=True)
    finally:
        release.set()
        for t in threads: t.join(3)
        h.close()


def test_organizer_actual_failure_uses_captured_lease_after_new_holder_commits(tmp_path,monkeypatch):
    from tests.test_model_drift_guards import _organize_contract
    h=GenerationHarness(tmp_path)
    try:
        detail=h.sample_detail(); draft=h.catalog.get_draft(detail['drafts'][0]['draftId'])
        p=FakeLLMProvider(replies=[AppError('independent old upstream failure',code='UPSTREAM_UNAVAILABLE',status_code=503)])
        handle=h.handle(provider=p); store=h.service.job_engine.store('question')
        job=store.create(kind='organize',frozen_input=_organize_contract(draft,fingerprint=fingerprint_of_handle(handle)))
        original=h.catalog.record_organize_failure; captured=[]; expected={}
        def paused(job_id,**kwargs):
            captured.append((kwargs['lease_attempt'],kwargs['lease_token']))
            with h.catalog.write_transaction() as c:
                c.execute("UPDATE question_jobs SET lease_expires_at='2000-01-01T00:00:00Z' WHERE id=?",(job_id,))
            fresh=store.claim(job_id)
            current,suggestion,failure=h.catalog.record_organize_batch(job_id,batch_index=0,next_batch_index=1,draft_id=draft.draft_id,base_draft_revision=draft.revision,proposed_content=dict(draft.content),lease_attempt=fresh.attempt,lease_token=fresh.token)
            assert suggestion is not None and failure is None
            expected.update(checkpoint=current.checkpoint,token=fresh.token)
            return original(job_id,**kwargs)
        monkeypatch.setattr(h.catalog,'record_organize_failure',paused)
        result=asyncio.run(h.service.job_engine.run_job('question',job.job_id,h.service._organize_executor_factory(model=handle),uses_model=True))
        assert len(captured)==1 and captured[0][0]==1 and captured[0][1]!=expected['token']
        assert result.state=='running' and result.attempt==2 and result.lease_token==expected['token']
        assert result.checkpoint==expected['checkpoint'] and result.checkpoint['nextBatchIndex']==1
        assert len(result.checkpoint['suggestionIds'])==1 and 'jobError' not in result.checkpoint
        assert len(h.catalog.list_suggestions(organization_job_id=job.job_id))==1
        emit('R09 actual service failure',oldAttempt=1,newAttempt=2,originalLeaseCaptured=True,newCheckpointPreserved=True)
    finally: h.close()
