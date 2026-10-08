"""知识点表格导入验收：预览持久化、映射、动作、回滚、重放（TEACHING-LOOP B1 / A6）。

XLSX 用 ``openpyxl`` 现造、CSV 用标准库现造；全部落在 ``tmp_path`` 隔离数据根。
覆盖：中文表头自动映射 / 手工映射与改映射、``0012`` 前导零保留、批内父节点拓扑写入、
update/ignore、阻断问题拒绝且零写入、``submissionId`` 重放、apply 中途失败整批回滚、
空白单元格不清空 vs ``clearFields`` 清空。
"""

from __future__ import annotations

import io
import json
import sqlite3
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.api.v1 import knowledge as knowledge_route
from app.main import create_app
from app.repositories.assets.file_assets import FileAssetsRepository
from app.repositories.knowledge.points import KnowledgePointRepository
from app.services.assets.store import AssetStore
from app.services.knowledge.imports import BLOCKING_ISSUE_CODES
from app.services.knowledge.service import build_knowledge_service
from app.services.publication import PublicationCoordinator
from tests.conftest import make_settings

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def blocking_count(view: dict) -> int:
    """详情视图没有 blockingIssueCount 字段（那属于批次摘要），这里按同一码表累计。"""
    codes = [issue["code"] for issue in view["issues"]]
    codes.extend(issue["code"] for row in view["rows"] for issue in row["issues"])
    return sum(1 for code in codes if code in BLOCKING_ISSUE_CODES)


def ensure_knowledge_router(app) -> None:
    paths = {getattr(route, "path", None) for route in app.router.routes}
    if "/api/v1/knowledge-points" not in paths:
        app.include_router(knowledge_route.router, prefix="/api/v1")
    routes = list(app.router.routes)
    catch = [
        route for route in routes if getattr(route, "name", "") == "feature_not_implemented"
    ]
    if catch:
        app.router.routes[:] = [route for route in routes if route not in catch] + catch


def make_xlsx(headers: list[str], rows: list[list[str]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_csv(headers: list[str], rows: list[list[str]]) -> bytes:
    lines = [",".join(headers)]
    for row in rows:
        lines.append(",".join(row))
    return ("\n".join(lines) + "\n").encode("utf-8-sig")


class ImportHarness:
    """知识点导入测试台：四库临时目录 + 挂载 B1 路由的真实服务。"""

    def __init__(self, tmp_path: Path) -> None:
        self.settings = make_settings(tmp_path / "data")
        self.app = create_app(self.settings)
        self.catalog = self.app.state.knowledge
        self.teaching = self.app.state.teaching
        self.assets = AssetStore(self.settings.assets_root)
        self.file_assets = FileAssetsRepository(self.teaching)
        self.coordinator = PublicationCoordinator()
        self.service = build_knowledge_service(
            self.catalog,
            asset_store=self.assets,
            file_assets=self.file_assets,
            evidence=None,
            coordinator=self.coordinator,
            model_resolver=None,
            job_engine=self.app.state.job_engine,
        )
        self.app.state.knowledge_service = self.service
        ensure_knowledge_router(self.app)
        self.client = TestClient(self.app, base_url="http://127.0.0.1:8001")
        self.client.__enter__()

    def close(self) -> None:
        self.client.__exit__(None, None, None)

    # -------------------------------------------------------------- 便捷方法

    def upload(
        self,
        data: bytes,
        *,
        name: str = "知识点.xlsx",
        subject_id: str | None = "math",
        mapping: dict[str, str] | None = None,
        sheet_name: str | None = None,
        media_type: str | None = None,
    ):
        fields: dict[str, str] = {}
        if subject_id is not None:
            fields["subjectId"] = subject_id
        if mapping is not None:
            fields["mappingJson"] = json.dumps(mapping, ensure_ascii=False)
        if sheet_name is not None:
            fields["sheetName"] = sheet_name
        default_media = XLSX_MEDIA if name.endswith(".xlsx") else "text/csv"
        return self.client.post(
            "/api/v1/knowledge-imports",
            files={"file": (name, data, media_type or default_media)},
            data=fields or None,
        )

    def upload_ok(self, data: bytes, **kwargs) -> dict:
        response = self.upload(data, **kwargs)
        assert response.status_code == 201, response.text
        return response.json()

    def preview(self, import_id: str) -> dict:
        response = self.client.get(f"/api/v1/knowledge-imports/{import_id}")
        assert response.status_code == 200, response.text
        return response.json()

    def confirm_body(
        self,
        import_id: str,
        *,
        actions=None,
        submission_id: str = "sub-1",
        expected_revision: int | None = None,
    ) -> dict:
        body: dict[str, object] = {
            "expectedRevision": (
                self.preview(import_id)["revision"]
                if expected_revision is None
                else expected_revision
            ),
            "submissionId": submission_id,
        }
        if actions is not None:
            body["actions"] = actions
        return body

    def confirm(self, import_id: str, body: dict):
        return self.client.post(f"/api/v1/knowledge-imports/{import_id}/confirm", json=body)

    def patch_import(self, import_id: str, **payload):
        body = {"expectedRevision": self.preview(import_id)["revision"], **payload}
        return self.client.patch(f"/api/v1/knowledge-imports/{import_id}", json=body)

    def create_point(self, code: str, **overrides) -> dict:
        body = {"subjectId": "math", "code": code, "name": f"知识点 {code}"}
        body.update(overrides)
        response = self.client.post("/api/v1/knowledge-points", json=body)
        assert response.status_code == 201, response.text
        return response.json()

    def point(self, point_id: str) -> dict:
        response = self.client.get(f"/api/v1/knowledge-points/{point_id}")
        assert response.status_code == 200, response.text
        return response.json()

    def points_by_code(self) -> dict[str, dict]:
        listing = self.client.get(
            "/api/v1/knowledge-points", params={"subjectId": "math", "limit": 200}
        ).json()
        return {item["code"]: item for item in listing["items"]}

    def count(self, db_path: Path, table: str) -> int:
        connection = sqlite3.connect(db_path)
        try:
            return int(connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0])
        finally:
            connection.close()

    def point_count(self) -> int:
        return self.count(self.catalog.db_path, "knowledge_points")

    def submission_count(self) -> int:
        return self.count(self.catalog.db_path, "knowledge_submissions")


@pytest.fixture()
def harness(tmp_path: Path):
    instance = ImportHarness(tmp_path)
    try:
        yield instance
    finally:
        instance.close()


# --------------------------------------------------------------------------- 预览


def test_xlsx_auto_mapping_preview_persists_rows(harness: ImportHarness) -> None:
    data = make_xlsx(
        ["学科", "编码", "名称", "说明", "父级", "别名"],
        [
            ["math", "0012", "有理数", "引入负数后的数系", "", "正数、负数"],
            ["math", "0013", "有理数加法", "同号相加", "0012", ""],
        ],
    )
    preview = harness.upload_ok(data)

    assert preview["source"] == "file"
    assert preview["state"] == "reviewing"
    assert preview["revision"] == 0
    assert preview["subjectId"] == "math"
    assert preview["mapping"] == {
        "subjectCode": "学科",
        "code": "编码",
        "name": "名称",
        "description": "说明",
        "parentCode": "父级",
        "aliases": "别名",
    }
    assert preview["headers"] == ["学科", "编码", "名称", "说明", "父级", "别名"]
    assert blocking_count(preview) == 0
    assert preview["fileAsset"]["kind"] == "attachment"
    assert preview["fileAsset"]["originalName"] == "知识点.xlsx"

    first, second = preview["rows"]
    assert first["rowNo"] == 1 and second["rowNo"] == 2
    assert first["code"] == "0012"  # 前导零保留
    assert first["aliases"] == ["正数", "负数"]
    assert first["decision"] is None and first["targetKnowledgePointId"] is None
    assert first["issues"] == []
    assert second["parentCode"] == "0012"
    assert second["issues"] == []

    # 预览持久化：详情与批次列表都能读回
    assert harness.preview(preview["importId"]) == preview
    listing = harness.client.get("/api/v1/knowledge-imports").json()
    assert listing["total"] == 1
    assert listing["items"][0]["rowCount"] == 2
    assert listing["items"][0]["blockingIssueCount"] == 0


def test_batch_issues_live_in_issues_json_and_mapping_is_pure(
    harness: ImportHarness,
) -> None:
    """批次级 issues 落 0003 的 issues_json 列；mapping_json 只有纯映射。"""
    data = make_xlsx(["学科", "编码", "名称", "额外列"], [["数学", "A1", "甲", "X"]])
    preview = harness.upload_ok(data)
    assert {issue["code"] for issue in preview["issues"]} == {
        "KNOWLEDGE_IMPORT_UNKNOWN_COLUMN"
    }

    connection = sqlite3.connect(harness.catalog.db_path)
    try:
        row = connection.execute(
            "SELECT mapping_json, issues_json FROM knowledge_imports WHERE id = ?",
            (preview["importId"],),
        ).fetchone()
    finally:
        connection.close()
    mapping_payload = json.loads(row[0])
    issues_payload = json.loads(row[1])
    assert mapping_payload == {
        "subjectCode": "学科",
        "code": "编码",
        "name": "名称",
    }
    assert "issues" not in mapping_payload and "mapping" not in mapping_payload
    assert [issue["code"] for issue in issues_payload] == ["KNOWLEDGE_IMPORT_UNKNOWN_COLUMN"]

    # 读回一致（批次级 issues 经新列持久化）
    assert harness.preview(preview["importId"])["issues"] == preview["issues"]


def test_legacy_mapping_layout_is_tolerated_on_read(harness: ImportHarness) -> None:
    """0003 之前的旧行（issues 塞在 mapping_json 保留键）仍可读出映射与 issues。"""
    data = make_xlsx(["学科", "编码", "名称", "额外列"], [["数学", "A1", "甲", "X"]])
    preview = harness.upload_ok(data)
    legacy_issue = {
        "code": "KNOWLEDGE_IMPORT_UNKNOWN_COLUMN",
        "column": "额外列",
        "message": "旧布局留下的问题。",
    }
    legacy_mapping = {"mapping": preview["mapping"], "issues": [legacy_issue]}
    connection = sqlite3.connect(harness.catalog.db_path)
    try:
        # issues_json 同时留一条同内容问题：读取时必须合并去重
        connection.execute(
            "UPDATE knowledge_imports SET mapping_json = ?, issues_json = ? WHERE id = ?",
            (
                json.dumps(legacy_mapping, ensure_ascii=False),
                json.dumps([legacy_issue], ensure_ascii=False),
                preview["importId"],
            ),
        )
        connection.commit()
    finally:
        connection.close()

    legacy = harness.preview(preview["importId"])
    assert legacy["mapping"] == preview["mapping"]
    # 旧 issues 与 issues_json 里的同内容问题合并去重后只剩一条
    assert len(legacy["issues"]) == 1
    issue = legacy["issues"][0]
    assert issue["code"] == "KNOWLEDGE_IMPORT_UNKNOWN_COLUMN"
    assert issue["column"] == "额外列"
    assert issue["message"] == "旧布局留下的问题。"
    summary = harness.client.get("/api/v1/knowledge-imports").json()["items"][0]
    assert summary["blockingIssueCount"] == 0


def test_csv_manual_mapping_and_patch_rederives_rows(harness: ImportHarness) -> None:
    data = make_csv(["代码", "知识点"], [["0012", "有理数"], ["0013", "有理数加法"]])
    preview = harness.upload_ok(data, name="知识点.csv", media_type="text/csv")
    assert preview["mapping"] == {}
    assert blocking_count(preview) == 4  # 两行各缺 code / name
    assert {issue["code"] for issue in preview["issues"]} == {
        "KNOWLEDGE_IMPORT_MISSING_COLUMN",
        "KNOWLEDGE_IMPORT_UNKNOWN_COLUMN",
    }
    assert [row["issues"][0]["code"] for row in preview["rows"]] == [
        "KNOWLEDGE_ROW_MISSING_CODE",
        "KNOWLEDGE_ROW_MISSING_CODE",
    ]

    patched = harness.patch_import(preview["importId"], mapping={"code": "代码", "name": "知识点"})
    assert patched.status_code == 200, patched.text
    body = patched.json()
    assert body["mapping"] == {"code": "代码", "name": "知识点"}
    assert body["revision"] == 1
    assert blocking_count(body) == 0
    assert [row["code"] for row in body["rows"]] == ["0012", "0013"]

    # 建批次时直接给手工映射（与自动映射合并）
    direct = harness.upload_ok(
        make_csv(["代码", "知识点"], [["0020", "数轴"]]),
        name="知识点2.csv",
        mapping={"code": "代码", "name": "知识点"},
    )
    assert direct["mapping"] == {"code": "代码", "name": "知识点"}
    assert direct["rows"][0]["code"] == "0020"

    # 改映射后仍可用同一确认接口入库
    confirmed = harness.confirm(
        preview["importId"],
        harness.confirm_body(
            preview["importId"],
            actions=[
                {"rowNo": 1, "decision": "create"},
                {"rowNo": 2, "decision": "create"},
            ],
        ),
    )
    assert confirmed.status_code == 200, confirmed.text
    assert [item["rowNo"] for item in confirmed.json()["created"]] == [1, 2]
    assert harness.points_by_code()["0013"]["name"] == "有理数加法"


def test_preview_keeps_blocking_issues_without_writing(harness: ImportHarness) -> None:
    data = make_xlsx(
        ["编码", "名称", "父级"],
        [
            ["", "缺编码", ""],
            ["R2", "父不存在", "NOPE"],
            ["R3", "正常", ""],
        ],
    )
    preview = harness.upload_ok(data)
    assert blocking_count(preview) == 2
    assert preview["rows"][0]["issues"][0] == {
        "row": 1,
        "column": "编码",
        "field": "code",
        "code": "KNOWLEDGE_ROW_MISSING_CODE",
        "message": "缺少知识点编码（code）。",
    }
    assert preview["rows"][1]["issues"][0]["code"] == "KNOWLEDGE_PARENT_INVALID"
    assert preview["rows"][1]["issues"][0]["field"] == "parentCode"
    assert preview["rows"][2]["issues"] == []
    summary = harness.client.get("/api/v1/knowledge-imports").json()["items"][0]
    assert summary["blockingIssueCount"] == 2

    refused = harness.confirm(
        preview["importId"],
        harness.confirm_body(
            preview["importId"],
            actions=[
                {"rowNo": 1, "decision": "create"},
                {"rowNo": 2, "decision": "create"},
                {"rowNo": 3, "decision": "create"},
            ],
        ),
    )
    assert refused.status_code == 422, refused.text
    body = refused.json()
    assert body["code"] == "KNOWLEDGE_IMPORT_BLOCKING_ISSUES"
    assert {issue["row"] for issue in body["details"]["issues"]} == {1, 2}
    # 零写入
    assert harness.point_count() == 0
    assert harness.submission_count() == 0
    unchanged = harness.preview(preview["importId"])
    assert unchanged["state"] == "reviewing" and unchanged["revision"] == 0


def test_unknown_and_mismatched_columns_are_non_blocking(harness: ImportHarness) -> None:
    data = make_xlsx(["学科", "编码", "名称", "额外列"], [["数学", "A1", "甲", "X"]])
    preview = harness.upload_ok(data)
    assert blocking_count(preview) == 0
    codes = {issue["code"] for issue in preview["issues"]}
    assert "KNOWLEDGE_IMPORT_UNKNOWN_COLUMN" in codes
    assert preview["rows"][0]["issues"][0]["code"] == "KNOWLEDGE_IMPORT_SUBJECT_MISMATCH"
    assert preview["rows"][0]["issues"][0]["field"] == "subjectCode"


# --------------------------------------------------------------------------- 确认


def test_confirm_create_update_ignore_with_topology(harness: ImportHarness) -> None:
    existing = harness.create_point("E1", name="既有", description="旧说明")
    data = make_xlsx(
        ["编码", "名称", "说明", "父级", "别名"],
        [
            ["C1", "子", "", "P1", ""],
            ["P1", "父", "", "", ""],
            ["E1", "改名后", "", "", "别名X"],
            ["I1", "忽略", "", "", ""],
        ],
    )
    preview = harness.upload_ok(data)
    rows = {row["code"]: row for row in preview["rows"]}
    assert rows["E1"]["targetKnowledgePointId"] == existing["id"]
    assert rows["E1"]["baseRevision"] == 0
    assert rows["E1"]["baseVersion"] == 1  # 预览冻结的内容修订版本
    assert rows["E1"]["decision"] == "update"  # 已存在同 code → 建议 update

    response = harness.confirm(
        preview["importId"],
        harness.confirm_body(
            preview["importId"],
            actions=[
                {"rowNo": 1, "decision": "create"},
                {"rowNo": 2, "decision": "create"},
                {"rowNo": 3, "decision": "update"},
                {"rowNo": 4, "decision": "ignore"},
            ],
        ),
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["state"] == "confirmed"
    assert result["replayed"] is False
    # 批内父节点先建（文件中子行在前，仍按拓扑序写入）
    assert [item["rowNo"] for item in result["created"]] == [2, 1]
    assert [item["rowNo"] for item in result["updated"]] == [3]
    assert result["ignored"] == [4]

    points = harness.points_by_code()
    assert points["C1"]["parentId"] == points["P1"]["id"]
    assert points["C1"]["parentCode"] == "P1"
    assert "I1" not in points
    updated = points["E1"]
    assert updated["name"] == "改名后"
    assert updated["description"] == "旧说明"  # 空白单元格不清空
    assert updated["aliases"] == ["别名X"]
    assert updated["version"] == 2 and updated["revision"] == 1
    assert harness.preview(preview["importId"])["state"] == "confirmed"
    assert harness.submission_count() == 1


def test_confirm_replay_returns_original_result(harness: ImportHarness) -> None:
    data = make_xlsx(["编码", "名称"], [["A1", "甲"], ["A2", "乙"]])
    preview = harness.upload_ok(data)
    body = harness.confirm_body(
        preview["importId"],
        actions=[
            {"rowNo": 1, "decision": "create"},
            {"rowNo": 2, "decision": "create"},
        ],
    )
    first = harness.confirm(preview["importId"], body)
    assert first.status_code == 200, first.text
    assert first.json()["replayed"] is False

    # 重放：同一 submissionId + 同一请求体（含同一 expectedRevision）
    replay = harness.confirm(preview["importId"], json.loads(json.dumps(body)))
    assert replay.status_code == 200, replay.text
    assert replay.json()["replayed"] is True
    assert replay.json()["created"] == first.json()["created"]
    assert harness.point_count() == 2
    assert harness.submission_count() == 1

    conflict = harness.confirm(
        preview["importId"],
        {**body, "expectedRevision": int(body["expectedRevision"]) + 1},
    )
    assert conflict.status_code == 409
    assert conflict.json()["code"] == "SUBMISSION_CONFLICT"


def test_confirm_rolls_back_whole_batch_on_mid_apply_failure(
    harness: ImportHarness, monkeypatch: pytest.MonkeyPatch
) -> None:
    data = make_xlsx(["编码", "名称"], [["B1", "甲"], ["B2", "乙"]])
    preview = harness.upload_ok(data)
    original = KnowledgePointRepository.create_point
    calls = {"count": 0}

    def flaky(self, conn, **kwargs):
        calls["count"] += 1
        if calls["count"] == 2:
            raise RuntimeError("boom")
        return original(self, conn, **kwargs)

    monkeypatch.setattr(KnowledgePointRepository, "create_point", flaky)
    body = harness.confirm_body(
        preview["importId"],
        actions=[
            {"rowNo": 1, "decision": "create"},
            {"rowNo": 2, "decision": "create"},
        ],
    )
    with pytest.raises(RuntimeError):
        harness.confirm(preview["importId"], body)
    assert calls["count"] == 2
    assert harness.point_count() == 0
    assert harness.submission_count() == 0
    assert harness.preview(preview["importId"])["state"] == "reviewing"


def test_update_row_with_blank_cells_keeps_and_clear_fields_clears(
    harness: ImportHarness,
) -> None:
    point = harness.create_point("U1", name="原名", description="原说明", aliases=["旧别名"])
    data = make_xlsx(["编码", "名称", "说明", "别名"], [["U1", "原名", "", ""]])
    preview = harness.upload_ok(data)
    response = harness.confirm(
        preview["importId"],
        harness.confirm_body(preview["importId"], actions=[{"rowNo": 1, "decision": "update"}]),
    )
    assert response.status_code == 200, response.text
    unchanged = harness.point(point["id"])
    assert unchanged["description"] == "原说明"
    assert unchanged["aliases"] == ["旧别名"]
    assert unchanged["version"] == 1 and unchanged["revision"] == 0

    cleared = harness.client.patch(
        f"/api/v1/knowledge-points/{point['id']}",
        json={"expectedRevision": 0, "clearFields": ["description", "aliases"]},
    )
    assert cleared.status_code == 200, cleared.text
    assert cleared.json()["description"] == ""
    assert cleared.json()["aliases"] == []


def test_confirm_requires_final_action_for_every_row(harness: ImportHarness) -> None:
    data = make_xlsx(["编码", "名称"], [["N1", "甲"]])
    preview = harness.upload_ok(data)
    response = harness.confirm(preview["importId"], harness.confirm_body(preview["importId"]))
    assert response.status_code == 422
    body = response.json()
    assert body["code"] == "KNOWLEDGE_ROW_INVALID"
    assert body["details"]["issues"][0]["row"] == 1
    assert harness.point_count() == 0


def test_import_list_state_filter(harness: ImportHarness) -> None:
    data = make_xlsx(["编码", "名称"], [["A1", "甲"]])
    preview = harness.upload_ok(data)
    reviewing = harness.client.get(
        "/api/v1/knowledge-imports", params={"state": "reviewing"}
    ).json()
    assert reviewing["total"] == 1
    confirmed = harness.client.get(
        "/api/v1/knowledge-imports", params={"state": "confirmed"}
    ).json()
    assert confirmed["total"] == 0
    ok = harness.confirm(
        preview["importId"],
        harness.confirm_body(preview["importId"], actions=[{"rowNo": 1, "decision": "create"}]),
    )
    assert ok.status_code == 200, ok.text
    assert (
        harness.client.get("/api/v1/knowledge-imports", params={"state": "confirmed"}).json()[
            "total"
        ]
        == 1
    )


def test_import_batch_cycle_is_blocking_with_parent_code_field(
    harness: ImportHarness,
) -> None:
    """批内 parentCode 互相指向形成环：确认被拒且可按 parentCode 定位（F1 第三条）。"""
    data = make_xlsx(
        ["编码", "名称", "父级"],
        [["A1", "甲", "A2"], ["A2", "乙", "A1"]],
    )
    preview = harness.upload_ok(data)
    assert blocking_count(preview) == 2
    row_issues = [issue for row in preview["rows"] for issue in row["issues"]]
    assert [issue["code"] for issue in row_issues] == ["KNOWLEDGE_CYCLE", "KNOWLEDGE_CYCLE"]
    assert {issue["field"] for issue in row_issues} == {"parentCode"}

    refused = harness.confirm(
        preview["importId"],
        harness.confirm_body(
            preview["importId"],
            actions=[
                {"rowNo": 1, "decision": "create"},
                {"rowNo": 2, "decision": "create"},
            ],
        ),
    )
    assert refused.status_code == 422, refused.text
    body = refused.json()
    assert body["code"] == "KNOWLEDGE_IMPORT_BLOCKING_ISSUES"
    assert {issue["code"] for issue in body["details"]["issues"]} == {"KNOWLEDGE_CYCLE"}
    assert {issue["field"] for issue in body["details"]["issues"]} == {"parentCode"}
    assert {issue["row"] for issue in body["details"]["issues"]} == {1, 2}
    assert harness.point_count() == 0


def test_patch_import_row_decision_bumps_revision(harness: ImportHarness) -> None:
    data = make_xlsx(["编码", "名称"], [["P1", "甲"], ["P2", "乙"]])
    preview = harness.upload_ok(data)
    patched = harness.patch_import(
        preview["importId"],
        rows=[
            {"rowNo": 1, "decision": "ignore"},
            {"rowNo": 2, "decision": "create"},
        ],
    )
    assert patched.status_code == 200, patched.text
    body = patched.json()
    assert body["revision"] == 1
    assert [row["decision"] for row in body["rows"]] == ["ignore", "create"]

    # 行 decision 已经足够：确认时无需再给 actions
    confirmed = harness.confirm(
        preview["importId"], harness.confirm_body(preview["importId"])
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["ignored"] == [1]
    assert [item["rowNo"] for item in confirmed.json()["created"]] == [2]
    assert harness.point_count() == 1


def test_multi_sheet_xlsx_uses_sheet_name(harness: ImportHarness) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    first = workbook.active
    first.title = "说明"
    first.append(["备注"])
    first.append(["封面页，不是知识点"])
    second = workbook.create_sheet("知识点")
    second.append(["编码", "名称"])
    second.append(["S1", "来自第二个工作表"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    default = harness.upload_ok(buffer.getvalue())
    assert default["headers"] == ["备注"]
    assert len(default["rows"]) == 1
    assert blocking_count(default) == 2  # 第一个工作表不是知识点表：缺 code/name
    assert any("工作表" in warning for warning in default["warnings"])

    chosen = harness.upload_ok(buffer.getvalue(), sheet_name="知识点")
    assert [row["code"] for row in chosen["rows"]] == ["S1"]
    assert chosen["rows"][0]["name"] == "来自第二个工作表"


def test_upload_requires_subject_and_supported_format(harness: ImportHarness) -> None:
    data = make_xlsx(["编码", "名称"], [["A1", "甲"]])
    missing_subject = harness.upload(data, subject_id=None)
    assert missing_subject.status_code == 422
    assert missing_subject.json()["code"] == "INVALID_REQUEST"
    unsupported = harness.upload(
        b"%PDF-1.4 stub", name="知识点.pdf", media_type="application/pdf"
    )
    assert unsupported.status_code == 422
    assert unsupported.json()["code"] == "UNSUPPORTED_DOCUMENT_FORMAT"


def test_upload_over_10_mib_is_413(harness: ImportHarness) -> None:
    from app.services.knowledge.service import MAX_UPLOAD_BYTES

    oversized = b"a" * (MAX_UPLOAD_BYTES + 1)
    response = harness.upload(oversized, name="知识点.csv", media_type="text/csv")
    assert response.status_code == 413
    assert response.json()["code"] == "DOCUMENT_TOO_LARGE"
    assert harness.count(harness.catalog.db_path, "knowledge_imports") == 0


# --------------------------------------------------------------------------- 放弃批次


def discard(harness: ImportHarness, import_id: str, *, expected_revision: int):
    return harness.client.post(
        f"/api/v1/knowledge-imports/{import_id}/discard",
        json={"expectedRevision": expected_revision},
    )


def test_discard_unconfirmed_import_keeps_records_and_zero_writes(
    harness: ImportHarness,
) -> None:
    """放弃未确认批次：state=cancelled + revision+1；记录/文件/预览行保留、零知识点写入。"""
    preview = harness.upload_ok(make_xlsx(["编码", "名称"], [["D1", "甲"], ["D2", "乙"]]))

    discarded = discard(harness, preview["importId"], expected_revision=preview["revision"])
    assert discarded.status_code == 200, discarded.text
    body = discarded.json()
    assert body["state"] == "cancelled"
    assert body["revision"] == preview["revision"] + 1
    # 批次记录、原始文件资产与预览行保留（审计），没有任何知识点写入
    assert harness.count(harness.catalog.db_path, "knowledge_imports") == 1
    assert harness.count(harness.catalog.db_path, "knowledge_import_rows") == 2
    assert harness.point_count() == 0
    # 原始文件资产保留（教学库 file_assets + assets/blobs 受管文件）
    assert harness.count(harness.teaching.db_path, "file_assets") == 1

    # 放弃后不可再 patch/confirm（cancelled 不在 EDITABLE/CONFIRMABLE 状态集合内，
    # 模块既有守卫按 422 INVALID_REQUEST 拒绝，与 failed 状态同一分支）
    patched = harness.patch_import(
        preview["importId"], rows=[{"rowNo": 1, "decision": "create"}]
    )
    assert patched.status_code == 422, patched.text
    assert patched.json()["code"] == "INVALID_REQUEST"
    assert "cancelled" in patched.json()["message"]

    # 放弃后不可再 confirm
    confirmed = harness.confirm(
        preview["importId"], harness.confirm_body(preview["importId"])
    )
    assert confirmed.status_code == 422, confirmed.text
    assert confirmed.json()["code"] == "INVALID_REQUEST"
    assert harness.point_count() == 0

    # 重复放弃幂等：不报错、不再递增 revision
    again = discard(harness, preview["importId"], expected_revision=body["revision"])
    assert again.status_code == 200, again.text
    assert again.json()["state"] == "cancelled"
    assert again.json()["revision"] == body["revision"]


def test_discard_confirmed_import_is_409(harness: ImportHarness) -> None:
    """已确认批次不可放弃：409 KNOWLEDGE_IMPORT_CONFIRMED，已写入的知识点保留。"""
    preview = harness.upload_ok(make_xlsx(["编码", "名称"], [["C1", "甲"]]))
    confirmed = harness.confirm(
        preview["importId"],
        harness.confirm_body(
            preview["importId"], actions=[{"rowNo": 1, "decision": "create"}]
        ),
    )
    assert confirmed.status_code == 200, confirmed.text
    current = harness.preview(preview["importId"])
    assert current["state"] == "confirmed"

    refused = discard(harness, preview["importId"], expected_revision=current["revision"])
    assert refused.status_code == 409, refused.text
    body = refused.json()
    assert body["code"] == "KNOWLEDGE_IMPORT_CONFIRMED"
    assert harness.point_count() == 1


def test_discard_import_revision_conflict_carries_current_revision(
    harness: ImportHarness,
) -> None:
    """expectedRevision 不符：409 REVISION_CONFLICT + details.currentRevision，状态不变。"""
    preview = harness.upload_ok(make_xlsx(["编码", "名称"], [["R1", "甲"]]))

    stale = discard(harness, preview["importId"], expected_revision=99)
    assert stale.status_code == 409, stale.text
    body = stale.json()
    assert body["code"] == "REVISION_CONFLICT"
    assert body["details"]["currentRevision"] == preview["revision"]

    after = harness.preview(preview["importId"])
    assert after["state"] == preview["state"]
    assert after["revision"] == preview["revision"]


def test_discard_import_route_not_found(harness: ImportHarness) -> None:
    """不存在的批次放弃：404 KNOWLEDGE_IMPORT_NOT_FOUND。"""
    response = discard(harness, "no-such-import", expected_revision=0)
    assert response.status_code == 404, response.text
    assert response.json()["code"] == "KNOWLEDGE_IMPORT_NOT_FOUND"
