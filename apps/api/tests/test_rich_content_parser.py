"""富内容解析（T10 / 验收项 A3）：块顺序与定位、表格合并、图片真实字节、OMML 原样、
未知对象可见问题、共同材料分组与 ``RichContentV2`` 组装。

样本全部在 ``tmp_path`` 程序化生成（``tests/rich_content_support.py``），不读正式教材目录、
不写正式 ``.local-data``、不联网。
"""

from __future__ import annotations

import hashlib

import pytest
from tests import rich_content_support as support
from tests.rich_content_support import (
    MATERIAL_BODY,
    MATERIAL_HEADING,
    PNG_BYTES,
    QUESTION_ONE,
    QUESTION_TWO_PREFIX,
    QUESTION_TWO_SUFFIX,
    TABLE_H_HEADER,
    TABLE_H_LEFT,
    TABLE_H_RIGHT,
    TABLE_V_MERGED,
    TABLE_V_RIGHT_BOTTOM,
    TABLE_V_RIGHT_TOP,
    TITLE_TEXT,
    build_rich_docx,
    rename_image_relationship,
    source_omml_strings,
)

from app.contracts.teaching_loop import (
    FormulaBlock,
    ImageBlock,
    ParagraphBlock,
    RichContentV2,
    RichOrigin,
    SharedMaterial,
    TableBlock,
    TableCell,
)
from app.core.exceptions import AppError
from app.services.assets.store import AssetStore
from app.services.rich_content import (
    group_shared_materials,
    is_material_heading_block,
    is_question_number_block,
    omml_from_latex,
    parse_docx_rich,
    rich_content_from_blocks,
)

EXPECTED_KINDS = [
    "paragraph",  # 1 标题
    "paragraph",  # 2 材料标题
    "paragraph",  # 3 材料正文
    "image",  # 4 材料配图
    "table",  # 5 横向合并表格
    "table",  # 6 纵向合并表格
    "paragraph",  # 7 题号一
    "paragraph",  # 8 题号二（公式前半段）
    "formula",  # 8-1 行内公式
    "paragraph",  # 8-2 题号二（公式后半段）
    "formula",  # 9 独立公式段
]


@pytest.fixture()
def sample(tmp_path):
    """样本 DOCX + 资产根 + 一次解析结果。"""
    docx_path = build_rich_docx(tmp_path / "sample.docx")
    store = AssetStore(tmp_path / "assets")
    parsed = parse_docx_rich(path=docx_path, file_name="sample.docx", assets=store)
    return docx_path, store, parsed


def _blocks_of(parsed, kind: str):
    return [block for block in parsed.blocks if block.kind == kind]


def _plain_paragraph(block_id: str, text: str) -> ParagraphBlock:
    return ParagraphBlock(id=block_id, kind="paragraph", text=text)


# --------------------------------------------------------------------------- 块顺序与定位


def test_blocks_follow_document_order_with_contiguous_locators(sample):
    _docx, _store, parsed = sample
    assert [block.kind for block in parsed.blocks] == EXPECTED_KINDS
    ids = [block.id for block in parsed.blocks]
    assert len(ids) == len(set(ids)), "块 id 必须唯一"
    assert parsed.block_kinds == {block.id: block.kind for block in parsed.blocks}
    assert len(parsed.locators) == len(parsed.blocks)

    starts = [parsed.locators[block_id]["blockStart"] for block_id in ids]
    assert starts == sorted(starts), "块顺序必须与原件 body 顺序一致"
    assert sorted(set(starts)) == list(range(1, 10)), "块序号 1 基、连续、无空洞"
    for block in parsed.blocks:
        locator = parsed.locators[block.id]
        assert locator["blockEnd"] == locator["blockStart"]
    # 同属第 8 个段落的三块共享同一来源坐标，顺序 = 文档顺序
    assert [parsed.locators[block_id]["blockStart"] for block_id in ("p8", "p8-1", "p8-2")] == [8, 8, 8]


def test_paragraph_text_and_inline_formula_split(sample):
    _docx, _store, parsed = sample
    texts = {block.id: block.text for block in _blocks_of(parsed, "paragraph")}
    assert TITLE_TEXT in texts.values()
    assert MATERIAL_HEADING in texts.values()
    assert MATERIAL_BODY in texts.values()
    assert QUESTION_ONE in texts.values()
    # 行内公式把段落切成 文本 → 公式 → 文本，三段都在且顺序正确
    assert texts["p8"] == QUESTION_TWO_PREFIX
    assert texts["p8-2"] == QUESTION_TWO_SUFFIX
    assert [
        block.id
        for block in parsed.blocks
        if parsed.locators[block.id]["blockStart"] == 8
    ] == ["p8", "p8-1", "p8-2"]
    # 空段落与"只有未知对象"的段落不产生块
    assert "" not in texts.values()
    assert len(_blocks_of(parsed, "formula")) == 2


# --------------------------------------------------------------------------- 表格合并


def test_table_horizontal_merge_becomes_col_span(sample):
    _docx, _store, parsed = sample
    table = _blocks_of(parsed, "table")[0]
    assert isinstance(table, TableBlock)
    assert table.column_count == 2, "解析器必须填写网格宽度（行边界还原的唯一依据）"
    assert parsed.locators[table.id]["tableColumns"] == 2
    assert parsed.locators[table.id]["tableRows"] == 2
    cells = table.cells
    assert [(cell.text, cell.col_span, cell.row_span) for cell in cells] == [
        (TABLE_H_HEADER, 2, 1),
        (TABLE_H_LEFT, 1, 1),
        (TABLE_H_RIGHT, 1, 1),
    ]
    # 首行 + 横向合并单元格 = 表头
    assert cells[0].is_header is True
    assert [cell.is_header for cell in cells[1:]] == [False, False]


def test_table_vertical_merge_becomes_row_span(sample):
    _docx, _store, parsed = sample
    table = _blocks_of(parsed, "table")[1]
    assert [(cell.text, cell.row_span, cell.col_span) for cell in table.cells] == [
        (TABLE_V_MERGED, 2, 1),
        (TABLE_V_RIGHT_TOP, 1, 1),
        (TABLE_V_RIGHT_BOTTOM, 1, 1),
    ]
    # 被纵向合并覆盖的续接单元格不再产生条目，因此只有 3 个单元格
    assert len(table.cells) == 3
    assert table.column_count == 2, "含 rowSpan 的表格也要给出网格宽度"
    assert [cell.is_header for cell in table.cells] == [True, True, False]
    assert parsed.locators[table.id]["tableColumns"] == 2
    assert parsed.locators[table.id]["tableRows"] == 2
    # 每个表格块都必须带 columnCount（扁平列表的行边界唯一依据）
    assert all(block.column_count is not None for block in _blocks_of(parsed, "table"))


# --------------------------------------------------------------------------- 图片


def test_image_is_stored_by_real_bytes_and_readable_back(sample):
    docx_path, store, parsed = sample
    images = _blocks_of(parsed, "image")
    assert len(images) == 1
    block = images[0]
    assert isinstance(block, ImageBlock)
    digest = hashlib.sha256(PNG_BYTES).hexdigest()
    assert block.asset_id == f"blobs/{digest}"
    assert store.read(block.asset_id) == PNG_BYTES, "受管资产能按 assetId 读回原字节"
    assert (block.width, block.height) == (
        support.IMAGE_WIDTH_PX,
        support.IMAGE_HEIGHT_PX,
    )

    assert len(parsed.assets) == 1
    asset = parsed.assets[0]
    assert asset.asset_id == block.asset_id
    assert asset.sha256 == digest == hashlib.sha256(store.read(block.asset_id)).hexdigest()
    assert asset.media_type == "image/png"
    # 原件本身没有被改写（解析只读原件）
    assert docx_path.read_bytes()[:2] == b"PK"


def test_image_extracted_through_renamed_relationship(tmp_path):
    """图片按关系取字节，不靠固定 rId：把图片关系重定向到哨兵 rId 后仍能取到。"""
    docx_path = build_rich_docx(tmp_path / "renamed.docx")
    original_id = rename_image_relationship(docx_path)
    store = AssetStore(tmp_path / "assets")
    parsed = parse_docx_rich(path=docx_path, file_name="renamed.docx", assets=store)
    images = _blocks_of(parsed, "image")
    assert len(images) == 1
    assert store.read(images[0].asset_id) == PNG_BYTES
    assert original_id != support.IMAGE_SENTINEL_RID


# --------------------------------------------------------------------------- OMML


def test_omml_kept_verbatim(sample):
    docx_path, _store, parsed = sample
    formulas = _blocks_of(parsed, "formula")
    assert len(formulas) == 2
    originals = source_omml_strings(docx_path)
    assert originals and len(originals) == 2
    assert [block.omml_xml for block in formulas] == originals, (
        "ommlXml 必须是原件 m:oMath 的逐字节序列化（原样保留，不改写、不做 LaTeX 往返）"
    )
    for block in formulas:
        assert block.latex is None
        assert "oMath" in block.omml_xml
        assert "<m:t>" in block.omml_xml


# --------------------------------------------------------------------------- 未知对象


def test_unknown_objects_are_reported_with_locator(sample):
    _docx, _store, parsed = sample
    assert [issue.code for issue in parsed.issues] == ["UNSUPPORTED_OBJECT", "UNSUPPORTED_OBJECT"]
    inside_paragraph, body_level = parsed.issues
    # 段落内 w:object：段落没有块，但问题仍带段落坐标
    assert inside_paragraph.block_id is None
    assert "object" in inside_paragraph.message
    assert inside_paragraph.source_locator == {"blockStart": 10, "blockEnd": 10}
    # 正文级 w:sdt：带 bodyIndex 定位，内容不静默消失
    assert body_level.block_id is None
    assert "sdt" in body_level.message
    assert isinstance(body_level.source_locator.get("bodyIndex"), int)
    assert body_level.source_locator["bodyIndex"] >= 1
    # 未知对象没有变成"空块"混进内容
    assert all(block.kind != "paragraph" or block.text for block in parsed.blocks)


# --------------------------------------------------------------------------- 共同材料分组


def test_shared_material_grouping_on_fixture(sample):
    _docx, _store, parsed = sample
    materials, unassigned = group_shared_materials(parsed.blocks, parsed.locators)
    assert [(material.id, [block.id for block in material.blocks]) for material in materials] == [
        ("material-1", ["p1", "p2", "p3", "p4", "t5", "t6"])
    ]
    assert unassigned == ("p7", "p8", "p8-1", "p8-2", "p9")
    # unassigned 只含题号段与题号之后的未分类块
    by_id = {block.id: block for block in parsed.blocks}
    assert is_question_number_block(by_id["p7"]) is True
    assert is_question_number_block(by_id["p8"]) is True
    assert is_question_number_block(by_id["p9"]) is False
    assert is_material_heading_block(by_id["p2"]) is True
    # 材料组内块顺序 = 文档顺序，且组内容与定位一致
    material_ids = [block.id for material in materials for block in material.blocks]
    assert material_ids == sorted(material_ids, key=lambda item: parsed.locators[item]["blockStart"])


def test_group_shared_materials_rules_are_conservative():
    blocks = [
        _plain_paragraph("a1", "阅读材料：甲"),
        _plain_paragraph("a2", "甲正文第一段"),
        _plain_paragraph("q1", "1. 第一小题"),
        _plain_paragraph("o1", "A. 选项甲"),
        _plain_paragraph("b1", "材料乙"),
        _plain_paragraph("b2", "乙正文"),
        _plain_paragraph("q2", "2. 第二小题"),
        _plain_paragraph("t1", "尾注：本卷共两题"),
    ]
    locators = {block.id: {"blockStart": index + 1, "blockEnd": index + 1} for index, block in enumerate(blocks)}
    materials, unassigned = group_shared_materials(blocks, locators)
    assert [(material.id, [block.id for block in material.blocks]) for material in materials] == [
        ("material-1", ["a1", "a2"]),
        ("material-2", ["b1", "b2"]),
    ]
    assert unassigned == ("q1", "o1", "q2", "t1")

    # 没有任何题号段、也没有材料锚点：一律 unassigned，不猜
    lonely = [_plain_paragraph("x1", "第一段"), _plain_paragraph("x2", "第二段")]
    lonely_locators = {block.id: {"blockStart": index + 1} for index, block in enumerate(lonely)}
    assert group_shared_materials(lonely, lonely_locators) == ((), ("x1", "x2"))

    # 有材料锚点但没有题号段：锚点起组并连续吸收后续块
    anchored = [_plain_paragraph("y1", "阅读下面的材料，完成各题。"), _plain_paragraph("y2", "材料正文")]
    anchored_locators = {block.id: {"blockStart": index + 1} for index, block in enumerate(anchored)}
    materials, unassigned = group_shared_materials(anchored, anchored_locators)
    assert [(material.id, [block.id for block in material.blocks]) for material in materials] == [
        ("material-1", ["y1", "y2"])
    ]
    assert unassigned == ()

    assert group_shared_materials([], {}) == ((), ())


def test_question_number_pattern_is_conservative():
    positive = ["1. 题干", "2、题干", "第3题 题干", "（4）题干", "(5) 题干", "10．题干", "6) 题干"]
    negative = ["2023 年真题", "2.5 米等于多少厘米", "A. 选项甲", "材料正文", "第一题", "第Ⅱ卷"]
    for text in positive:
        assert is_question_number_block(_plain_paragraph("p", text)) is True, text
    for text in negative:
        assert is_question_number_block(_plain_paragraph("p", text)) is False, text
    assert is_question_number_block(TableBlock(id="t", kind="table", cells=[])) is False


# --------------------------------------------------------------------------- 组装


def test_rich_content_from_blocks_assembles_and_revalidates(sample):
    docx_path, _store, parsed = sample
    materials, unassigned = group_shared_materials(parsed.blocks, parsed.locators)
    origin = RichOrigin(
        originalAssetId="asset-origin-1",
        originalSha256=support.sha256_of_file(docx_path),
        sourceLocator={"fileName": "sample.docx"},
    )
    rich = rich_content_from_blocks(
        blocks=parsed.blocks,
        materials=materials,
        assets=parsed.assets,
        origin=origin,
        stem_block_ids=unassigned,
    )
    assert isinstance(rich, RichContentV2)
    assert rich.version == 2
    assert [block.id for block in rich.stem_blocks] == list(unassigned)
    assert [material.id for material in rich.shared_materials] == ["material-1"]
    assert rich.option_blocks == {} and rich.answer_blocks == [] and rich.explanation_blocks == []
    assert [asset.asset_id for asset in rich.assets] == [parsed.assets[0].asset_id]
    # TableBlock.columnCount 透传（组装不重建块对象、不丢字段）
    tables_in_rich = [
        block
        for material in rich.shared_materials
        for block in material.blocks
        if block.kind == "table"
    ] + [block for block in rich.stem_blocks if block.kind == "table"]
    assert len(tables_in_rich) == 2
    assert all(block.column_count == 2 for block in tables_in_rich)
    # 线上 JSON 形状可原样再验证（by_alias）
    dumped = rich.model_dump(by_alias=True)
    assert "sharedMaterials" in dumped and "stemBlocks" in dumped
    assert RichContentV2.model_validate(dumped) == rich

    # 默认（不给 stem_block_ids）：全部非材料块按文档顺序作题干
    default = rich_content_from_blocks(
        blocks=parsed.blocks, materials=materials, assets=parsed.assets, origin=origin
    )
    assert [block.id for block in default.stem_blocks] == list(unassigned)


def test_rich_content_from_blocks_rejects_inconsistent_inputs(sample):
    docx_path, _store, parsed = sample
    materials, unassigned = group_shared_materials(parsed.blocks, parsed.locators)
    origin = RichOrigin(originalAssetId="a", originalSha256="0" * 64)

    with pytest.raises(AppError) as excinfo:
        rich_content_from_blocks(
            blocks=parsed.blocks,
            materials=materials,
            assets=parsed.assets,
            origin=origin,
            stem_block_ids=[*unassigned, "t5"],  # t5 已归入共享材料
        )
    assert excinfo.value.code == "INVALID_REQUEST"
    assert excinfo.value.status_code == 422
    assert "t5" in str(excinfo.value)

    with pytest.raises(AppError) as excinfo:
        rich_content_from_blocks(
            blocks=parsed.blocks,
            materials=materials,
            assets=parsed.assets,
            origin=origin,
            stem_block_ids=["not-a-block"],
        )
    assert excinfo.value.code == "INVALID_REQUEST"

    with pytest.raises(AppError):
        rich_content_from_blocks(
            blocks=[parsed.blocks[0], parsed.blocks[0]],
            materials=[],
            assets=[],
            origin=origin,
        )

    with pytest.raises(AppError):
        rich_content_from_blocks(
            blocks=[parsed.blocks[0]],
            materials=[
                SharedMaterial(id="bad", blocks=[_plain_paragraph("x", "不在块列表里")])
            ],
            assets=[],
            origin=origin,
        )

    with pytest.raises(AppError):
        rich_content_from_blocks(
            blocks=[parsed.blocks[0], _plain_paragraph("p1", "同 id 不同内容")],
            materials=[],
            assets=[],
            origin=origin,
        )

    with pytest.raises(AppError):
        rich_content_from_blocks(
            blocks=parsed.blocks, materials=materials, assets=[], origin="not-an-origin"  # type: ignore[arg-type]
        )


# --------------------------------------------------------------------------- 解析失败路径


def test_parse_docx_rich_error_paths(sample, tmp_path):
    docx_path, store, _parsed = sample
    with pytest.raises(AppError) as excinfo:
        parse_docx_rich(path=tmp_path / "missing.docx", file_name="missing.docx", assets=store)
    assert (excinfo.value.code, excinfo.value.status_code) == ("DOCUMENT_FILE_MISSING", 500)

    with pytest.raises(AppError) as excinfo:
        parse_docx_rich(path=docx_path, file_name="sample.pdf", assets=store)
    assert (excinfo.value.code, excinfo.value.status_code) == (
        "UNSUPPORTED_DOCUMENT_FORMAT",
        422,
    )

    with pytest.raises(AppError) as excinfo:
        parse_docx_rich(path=docx_path, file_name="sample.docx", assets=object())
    assert (excinfo.value.code, excinfo.value.status_code) == ("INVALID_REQUEST", 422)

    with pytest.raises(AppError) as excinfo:
        parse_docx_rich(
            path=docx_path,
            file_name="sample.docx",
            assets=store,
            origin="bad-origin",  # type: ignore[arg-type]
        )
    assert excinfo.value.code == "INVALID_REQUEST"


def test_corrupt_docx_fails_loudly(tmp_path):
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a zip archive at all")
    with pytest.raises(AppError) as excinfo:
        parse_docx_rich(
            path=broken, file_name="broken.docx", assets=AssetStore(tmp_path / "assets")
        )
    assert (excinfo.value.code, excinfo.value.status_code) == ("DOCUMENT_PARSE_FAILED", 422)


def test_tbl_header_row_is_marked_and_tab_break_text_matches_python_docx(tmp_path):
    """``w:tblHeader`` 行也是表头；段落文本口径与 python-docx ``Paragraph.text`` 一致。"""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    path = tmp_path / "header.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("制表符")
    paragraph.add_run("\t")
    paragraph.add_run("换行")
    paragraph.add_run("\n")
    paragraph.add_run("末行")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "第一行甲"
    table.cell(0, 1).text = "第一行乙"
    for cell in table.rows[1].cells:
        cell.text = "重复表头"
    row_properties = table.rows[1]._tr.get_or_add_trPr()  # noqa: SLF001
    row_properties.append(OxmlElement("w:tblHeader"))
    document.save(str(path))

    parsed = parse_docx_rich(
        path=path, file_name="header.docx", assets=AssetStore(tmp_path / "assets")
    )
    paragraph_block = _blocks_of(parsed, "paragraph")[0]
    assert paragraph_block.text == document.paragraphs[0].text
    assert "\t" in paragraph_block.text and "\n" in paragraph_block.text
    table_block = _blocks_of(parsed, "table")[0]
    assert table_block.column_count == 2
    assert [cell.is_header for cell in table_block.cells] == [True, True, True, True]
    assert [cell.text for cell in table_block.cells] == [
        "第一行甲",
        "第一行乙",
        "重复表头",
        "重复表头",
    ]
    # 第二行没有 gridSpan / vMerge 之外的结构偏差
    assert {cell.row_span for cell in table_block.cells} == {1}
    assert {cell.col_span for cell in table_block.cells} == {1}


def test_origin_hash_mismatch_is_a_visible_warning(sample):
    docx_path, store, _parsed = sample
    matched = parse_docx_rich(
        path=docx_path,
        file_name="sample.docx",
        assets=store,
        origin=RichOrigin(
            originalAssetId="asset-origin-1",
            originalSha256=support.sha256_of_file(docx_path),
        ),
    )
    assert matched.warnings == ()

    mismatched = parse_docx_rich(
        path=docx_path,
        file_name="sample.docx",
        assets=store,
        origin=RichOrigin(originalAssetId="asset-origin-1", originalSha256="a" * 64),
    )
    assert any("originalSha256" in warning for warning in mismatched.warnings)
    assert len(mismatched.blocks) == len(matched.blocks), "来源元数据不一致不影响内容解析"


# --------------------------------------------------------------------------- LaTeX → OMML


def test_omml_from_latex_produces_self_contained_math_node():
    from lxml import etree

    from app.services.rich_content.formulas import MATH_NAMESPACE

    xml = omml_from_latex(r"\frac{1}{2}")
    assert "oMath" in xml
    root = etree.fromstring(xml.encode("utf-8"))
    assert root.tag == f"{{{MATH_NAMESPACE}}}oMath"
    assert "xmlns:m=" in xml, "返回值必须自带命名空间声明，调用方可直接解析"


@pytest.mark.parametrize("latex", ["", "   ", r"\frac{1}{", r"\begin{itemize}"])
def test_omml_from_latex_fails_loudly(latex):
    with pytest.raises(AppError) as excinfo:
        omml_from_latex(latex)
    assert (excinfo.value.code, excinfo.value.status_code) == (
        "FORMULA_CONVERSION_FAILED",
        422,
    )
    assert str(excinfo.value).strip()


def test_formula_block_requires_a_payload_source():
    """契约允许两者其一；两者都空由渲染层拒绝（此处只锁定解析层产出的块形态）。"""
    block = FormulaBlock(id="f", kind="formula", latex="x", ommlXml=None)
    assert block.omml_xml is None
    assert block.latex == "x"
    cell = TableCell(text="x", isHeader=True, rowSpan=2, colSpan=3)
    assert (cell.is_header, cell.row_span, cell.col_span) == (True, 2, 3)
