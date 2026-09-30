"""富内容渲染（T10 / 验收项 A4）：学生版与教师版、共同材料只出一次、图片重建 relationship、
表格合并、OMML 原样、LaTeX 经 math2docx 转换与失败报错。

源 DOCX 在 ``tmp_path/source`` 另建并把图片关系重定向到哨兵 rId（``rId404``）：
渲染产物必须**自己**重建 relationship —— 引用的图片在新包里、哨兵 rId 不得出现、
所有 ``r:embed`` 都能在新包解析。产物一律用 python-docx 重新打开校验。
"""

from __future__ import annotations

import hashlib
from io import BytesIO
from types import SimpleNamespace

import pytest
from lxml import etree
from tests.rich_content_support import (
    ANSWER_TEXT,
    EXPLANATION_TEXT,
    IMAGE_SENTINEL_RID,
    MATERIAL_BODY,
    MATERIAL_HEADING,
    OPTION_A_TEXT,
    OPTION_B_TEXT,
    ORIGIN_ASSET_ID,
    PNG_BYTES,
    QUESTION_ONE,
    STANDALONE_LATEX,
    TABLE_H_HEADER,
    TABLE_V_MERGED,
    TITLE_TEXT,
    build_rich_content,
    build_rich_docx,
    document_text,
    document_xml,
    embed_rids,
    omml_elements,
    omml_texts,
    package_image_blobs,
    rename_image_relationship,
    source_image_part_rids,
)

from app.contracts.teaching_loop import (
    FormulaBlock,
    ImageBlock,
    ParagraphBlock,
    RichAsset,
    RichContentV2,
    RichOrigin,
    TableBlock,
    TableCell,
)
from app.core.exceptions import AppError
from app.services.assets.store import AssetStore
from app.services.rich_content import (
    group_shared_materials,
    omml_from_latex,
    parse_docx_rich,
    render_rich_document,
)

W_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
ORIGINAL_OMML = (
    '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
    "<m:r><m:t>原样公式</m:t></m:r></m:oMath>"
)


@pytest.fixture()
def env(tmp_path):
    """源 DOCX（独立临时目录 + 哨兵 rId）+ 资产根 + 解析结果 + 完整富内容。"""
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    docx_path = build_rich_docx(source_dir / "sample.docx")
    original_rid = rename_image_relationship(docx_path)
    store = AssetStore(tmp_path / "assets")
    parsed = parse_docx_rich(path=docx_path, file_name="sample.docx", assets=store)
    materials, _unassigned = group_shared_materials(parsed.blocks, parsed.locators)
    rich = build_rich_content(parsed, materials, docx_path=docx_path)
    return SimpleNamespace(
        tmp_path=tmp_path,
        docx_path=docx_path,
        original_rid=original_rid,
        store=store,
        parsed=parsed,
        materials=materials,
        rich=rich,
    )


def _render(env, variant: str = "student", *, rich=None, title: str | None = "课后练习"):
    return render_rich_document(
        rich=env.rich if rich is None else rich,
        assets=env.store,
        variant=variant,  # type: ignore[arg-type]
        title=title,
    )


def _open(data: bytes):
    from docx import Document

    return Document(BytesIO(data))


def _merge_marks(data: bytes) -> tuple[list[str], list[str]]:
    """渲染产物里的 gridSpan / vMerge 取值（vMerge 无 val 记 ``continue``）。"""
    root = etree.fromstring(document_xml(data))
    spans = [
        node.get(f"{{{W_NAMESPACE}}}val") or ""
        for node in root.iter(f"{{{W_NAMESPACE}}}gridSpan")
    ]
    merges = [
        node.get(f"{{{W_NAMESPACE}}}val") or "continue"
        for node in root.iter(f"{{{W_NAMESPACE}}}vMerge")
    ]
    return spans, merges


def _payload_image_block(payload: dict, *, block_id: str | None = None) -> dict:
    """在线上 JSON 里找图片块（共同材料与题干都可能有）。"""
    groups = [*payload["sharedMaterials"]]
    for material in groups:
        for block in material["blocks"]:
            if block["kind"] == "image" and (block_id is None or block["id"] == block_id):
                return block
    for block in payload["stemBlocks"]:
        if block["kind"] == "image" and (block_id is None or block["id"] == block_id):
            return block
    raise AssertionError("样本里没有图片块")


def _rich_with_extra_block(env, block) -> RichContentV2:
    payload = env.rich.model_dump(by_alias=True)
    payload["stemBlocks"] = list(payload["stemBlocks"]) + [block.model_dump(by_alias=True)]
    return RichContentV2.model_validate(payload)


def _table_only_rich(
    cells,
    *,
    block_id: str = "t-only",
    column_count: int | None = None,
    table_grids: dict | None = None,
) -> RichContentV2:
    """只含一个表格块的富内容；``column_count=None`` 时不写 columnCount（旧数据回退路径）。"""
    table = TableBlock(
        id=block_id, kind="table", cells=list(cells), column_count=column_count
    )
    origin = RichOrigin(
        originalAssetId=ORIGIN_ASSET_ID,
        originalSha256="0" * 64,
        sourceLocator={"tableGrids": table_grids} if table_grids else {},
    )
    return RichContentV2.model_validate(
        {
            "version": 2,
            "sharedMaterials": [],
            "stemBlocks": [table.model_dump(by_alias=True)],
            "optionBlocks": {},
            "answerBlocks": [],
            "explanationBlocks": [],
            "assets": [],
            "origin": origin.model_dump(by_alias=True),
        }
    )


# --------------------------------------------------------------------------- 变体


def test_student_variant_omits_answers_and_explanation(env):
    student = _render(env, "student")
    text = document_text(_open(student))
    assert ANSWER_TEXT not in text
    assert EXPLANATION_TEXT not in text
    for expected in (
        TITLE_TEXT,
        MATERIAL_HEADING,
        MATERIAL_BODY,
        QUESTION_ONE,
        OPTION_A_TEXT,
        OPTION_B_TEXT,
    ):
        assert expected in text, expected
    assert "课后练习" in text  # 标题


def test_teacher_variant_includes_answers_and_explanation(env):
    teacher = _render(env, "teacher")
    text = document_text(_open(teacher))
    assert ANSWER_TEXT in text
    assert EXPLANATION_TEXT in text
    assert QUESTION_ONE in text and OPTION_A_TEXT in text
    # 学生版 / 教师版的差别只在这两段
    assert ANSWER_TEXT not in document_text(_open(_render(env, "student")))


def test_unknown_variant_is_rejected(env):
    with pytest.raises(AppError) as excinfo:
        _render(env, "anonymous")
    assert (excinfo.value.code, excinfo.value.status_code) == ("INVALID_REQUEST", 422)


# --------------------------------------------------------------------------- 共同材料


def test_shared_materials_are_emitted_exactly_once(env):
    for variant in ("student", "teacher"):
        text = document_text(_open(_render(env, variant)))
        for marker in (MATERIAL_HEADING, MATERIAL_BODY, TABLE_H_HEADER, TABLE_V_MERGED):
            assert text.count(marker) == 1, f"{variant}: {marker}"
        # 共同材料在题干之前
        assert text.index(MATERIAL_HEADING) < text.index(QUESTION_ONE)


def test_duplicate_block_ids_are_rejected(env):
    payload = env.rich.model_dump(by_alias=True)
    payload["stemBlocks"] = list(payload["stemBlocks"]) + [
        {"id": "opt-a", "kind": "paragraph", "text": "与选项 A 撞 id"}
    ]
    duplicated = RichContentV2.model_validate(payload)
    with pytest.raises(AppError) as excinfo:
        _render(env, "teacher", rich=duplicated)
    assert excinfo.value.code == "INVALID_REQUEST"
    assert "opt-a" in str(excinfo.value)
    assert excinfo.value.details == {"fields": ["opt-a"]}


# --------------------------------------------------------------------------- 图片


def test_images_are_rebuilt_with_fresh_relationships(env):
    data = _render(env, "student")
    document = _open(data)
    rels = document.part.rels

    # 源包里有 2 条图片关系（原始 + 哨兵），渲染产物只重建被引用的那一条
    assert len(source_image_part_rids(env.docx_path)) == 2
    image_rels = [rid for rid, rel in rels.items() if rel.reltype.endswith("/image")]
    assert len(image_rels) == 1
    assert IMAGE_SENTINEL_RID not in rels, "不得复用源文档的 rId"

    # 所有 r:embed / r:id 都能在**新包**里解析
    for rid in embed_rids(data):
        assert rid in rels, rid
    (rid,) = image_rels
    assert rels[rid].target_part.blob == PNG_BYTES
    # 图片字节真实写进了新包
    assert package_image_blobs(data) == [PNG_BYTES]
    assert IMAGE_SENTINEL_RID not in document_xml(data).decode("utf-8")


def test_same_asset_renders_one_picture_part(env):
    payload = env.rich.model_dump(by_alias=True)
    image_block = _payload_image_block(payload)
    payload["stemBlocks"] = list(payload["stemBlocks"]) + [
        {**image_block, "id": "image-again"}
    ]
    duplicated = RichContentV2.model_validate(payload)
    data = _render(env, "student", rich=duplicated)
    document = _open(data)
    image_rels = [rid for rid, rel in document.part.rels.items() if rel.reltype.endswith("/image")]
    assert len(image_rels) == 1, "同一资产在同一次渲染内共享一个 part"
    assert len(embed_rids(data)) == 2, "但两处引用都要出现"
    assert package_image_blobs(data) == [PNG_BYTES]


def test_image_asset_resolution_paths(env):
    payload = env.rich.model_dump(by_alias=True)
    digest = hashlib.sha256(PNG_BYTES).hexdigest()
    image_block = _payload_image_block(payload)
    image_block["assetId"] = "asset-image-1"
    payload["assets"] = [{"assetId": "asset-image-1", "sha256": digest, "mediaType": "image/png"}]
    content_with_store_id = RichContentV2.model_validate(payload)
    data = _render(env, "student", rich=content_with_store_id)
    assert package_image_blobs(data) == [PNG_BYTES]

    payload["assets"] = []
    unresolvable = RichContentV2.model_validate(payload)
    with pytest.raises(AppError) as excinfo:
        _render(env, "student", rich=unresolvable)
    assert (excinfo.value.code, excinfo.value.status_code) == ("ASSET_NOT_FOUND", 422)
    assert image_block["id"] in str(excinfo.value)


def test_tampered_asset_is_not_rendered(env, tmp_path):
    tampered = AssetStore(tmp_path / "tampered-assets")
    parsed = parse_docx_rich(
        path=env.docx_path, file_name="sample.docx", assets=tampered
    )
    materials, _unassigned = group_shared_materials(parsed.blocks, parsed.locators)
    rich = build_rich_content(parsed, materials, docx_path=env.docx_path)
    image = next(block for block in parsed.blocks if isinstance(block, ImageBlock))
    blob_path = tampered.path_of(image.asset_id)
    blob_path.write_bytes(b"\x89PNG\r\n\x1a\n-tampered")
    with pytest.raises(AppError) as excinfo:
        render_rich_document(rich=rich, assets=tampered, variant="student")
    assert (excinfo.value.code, excinfo.value.status_code) == ("ASSET_CORRUPT", 500)


# --------------------------------------------------------------------------- 表格


def test_tables_are_real_and_merges_apply(env):
    data = _render(env, "student")
    document = _open(data)
    assert len(document.tables) == 2
    horizontal, vertical = document.tables
    assert (len(horizontal.rows), len(horizontal.columns)) == (2, 2)
    assert (len(vertical.rows), len(vertical.columns)) == (2, 2)

    # 横向合并：同一行两列指向同一个 tc
    assert horizontal.cell(0, 0)._tc is horizontal.cell(0, 1)._tc  # noqa: SLF001
    # 纵向合并：同一列两行指向同一个 tc
    assert vertical.cell(0, 0)._tc is vertical.cell(1, 0)._tc  # noqa: SLF001
    assert vertical.cell(0, 0).text == TABLE_V_MERGED
    assert vertical.cell(1, 0).text == TABLE_V_MERGED

    spans, merges = _merge_marks(data)
    assert "2" in spans
    assert "restart" in merges and "continue" in merges

    # 表头加粗，正文不加粗
    header_runs = [run for p in horizontal.cell(0, 0).paragraphs for run in p.runs]
    assert header_runs and all(run.bold for run in header_runs)
    body_runs = [run for p in horizontal.cell(1, 0).paragraphs for run in p.runs]
    assert body_runs and not any(run.bold for run in body_runs)


def test_parsed_tables_render_without_origin_hint(env):
    """解析产物自带 columnCount：即使 origin 不给 tableGrids 提示也精确还原。"""
    without_hint = build_rich_content(
        env.parsed, env.materials, docx_path=env.docx_path, with_grid_hint=False
    )
    assert without_hint.origin.source_locator.get("tableGrids") is None
    document = _open(_render(env, "student", rich=without_hint))
    assert len(document.tables) == 2
    horizontal, vertical = document.tables
    assert (len(horizontal.rows), len(horizontal.columns)) == (2, 2)
    assert (len(vertical.rows), len(vertical.columns)) == (2, 2)
    assert horizontal.cell(0, 0)._tc is horizontal.cell(0, 1)._tc  # noqa: SLF001
    assert vertical.cell(0, 0)._tc is vertical.cell(1, 0)._tc  # noqa: SLF001


def test_column_count_alone_reconstructs_rows(env):
    """契约字段 columnCount 是行边界的唯一依据：不给 origin hint 也能精确还原。"""
    plain = [TableCell(text=f"单元格{i}") for i in range(4)]  # 无表头、无合并
    two_columns = _table_only_rich(plain, column_count=2)
    document = _open(render_rich_document(rich=two_columns, assets=env.store, variant="student"))
    assert (len(document.tables[0].rows), len(document.tables[0].columns)) == (2, 2)
    assert [cell.text for cell in document.tables[0].rows[0].cells] == ["单元格0", "单元格1"]
    assert [cell.text for cell in document.tables[0].rows[1].cells] == ["单元格2", "单元格3"]

    four_columns = _table_only_rich(plain, column_count=4)
    document = _open(
        render_rich_document(rich=four_columns, assets=env.store, variant="student")
    )
    assert (len(document.tables[0].rows), len(document.tables[0].columns)) == (1, 4)

    # 含 rowSpan 的样本同样只靠 columnCount 还原
    merged_cells = [
        TableCell(text="纵向合并", is_header=True, row_span=2),
        TableCell(text="右上", is_header=True),
        TableCell(text="右下"),
    ]
    merged = _table_only_rich(merged_cells, column_count=2)
    document = _open(render_rich_document(rich=merged, assets=env.store, variant="student"))
    assert (len(document.tables[0].rows), len(document.tables[0].columns)) == (2, 2)
    assert document.tables[0].cell(0, 0)._tc is document.tables[0].cell(1, 0)._tc  # noqa: SLF001

    # columnCount 与单元格序列矛盾：报错，不猜
    broken = _table_only_rich(plain, column_count=3)
    with pytest.raises(AppError) as excinfo:
        render_rich_document(rich=broken, assets=env.store, variant="student")
    assert (excinfo.value.code, excinfo.value.status_code) == ("INVALID_REQUEST", 422)
    assert "t-only" in str(excinfo.value)


def test_hinted_table_grid_is_used_when_cells_are_ambiguous(env):
    """旧数据回退路径（``columnCount`` 缺失）：定位提示 → 表头信号 → 前缀和。"""
    plain = [
        TableCell(text="单元格0", isHeader=True),
        TableCell(text="单元格1", isHeader=True),
        TableCell(text="单元格2"),
        TableCell(text="单元格3"),
    ]
    # 无 hint：靠表头信号（前两格 isHeader）还原 2×2
    inferred = _table_only_rich(plain)
    document = _open(render_rich_document(rich=inferred, assets=env.store, variant="student"))
    assert (len(document.tables[0].rows), len(document.tables[0].columns)) == (2, 2)

    # 有 hint：按调用方给的列数还原成一行四列
    hinted = _table_only_rich(plain, table_grids={"t-only": 4})
    document = _open(render_rich_document(rich=hinted, assets=env.store, variant="student"))
    assert (len(document.tables[0].rows), len(document.tables[0].columns)) == (1, 4)

    # hint 与单元格序列矛盾：报错，不猜一个"看起来成功"的表格
    broken = _table_only_rich(plain, table_grids={"t-only": 5})
    with pytest.raises(AppError) as excinfo:
        render_rich_document(rich=broken, assets=env.store, variant="student")
    assert (excinfo.value.code, excinfo.value.status_code) == ("INVALID_REQUEST", 422)
    assert "t-only" in str(excinfo.value)


def test_empty_table_block_is_rejected(env):
    with pytest.raises(AppError) as excinfo:
        render_rich_document(
            rich=_table_only_rich([], block_id="t-empty"),
            assets=env.store,
            variant="student",
        )
    assert (excinfo.value.code, excinfo.value.status_code) == ("INVALID_REQUEST", 422)
    assert "t-empty" in str(excinfo.value)


# --------------------------------------------------------------------------- 公式


def test_omml_is_inserted_verbatim_and_latex_is_converted(env):
    data = _render(env, "teacher")
    elements = omml_elements(data)
    assert len(elements) == 3, "行内公式 + 独立公式 + LaTeX-only 公式"
    inline_text, standalone_text, latex_text = omml_texts(data)
    assert inline_text == "x2+1=2"
    assert standalone_text == "ab=c"
    # LaTeX-only 经 math2docx 得到与原件同一结构的公式
    assert latex_text == standalone_text
    assert omml_from_latex(STANDALONE_LATEX)


def test_original_omml_wins_over_latex(env):
    block = FormulaBlock(
        id="f-both", kind="formula", latex=r"\beta", ommlXml=ORIGINAL_OMML
    )
    data = _render(env, "teacher", rich=_rich_with_extra_block(env, block))
    texts = omml_texts(data)
    assert "原样公式" in texts, "有 ommlXml 时不得走 LaTeX 转换"
    assert len(texts) == 4


def test_latex_conversion_failure_reports_block_id(env):
    broken = FormulaBlock(id="f-broken", kind="formula", latex=r"\frac{1}{", ommlXml=None)
    with pytest.raises(AppError) as excinfo:
        _render(env, "student", rich=_rich_with_extra_block(env, broken))
    assert (excinfo.value.code, excinfo.value.status_code) == (
        "FORMULA_CONVERSION_FAILED",
        422,
    )
    assert "f-broken" in str(excinfo.value)
    assert excinfo.value.details == {"fields": ["f-broken"]}


def test_formula_without_payload_is_rejected(env):
    empty = FormulaBlock(id="f-empty", kind="formula", latex=" ", ommlXml=None)
    with pytest.raises(AppError) as excinfo:
        _render(env, "student", rich=_rich_with_extra_block(env, empty))
    assert excinfo.value.code == "INVALID_REQUEST"
    assert "f-empty" in str(excinfo.value)


def test_broken_omml_xml_is_rejected(env):
    broken = FormulaBlock(id="f-xml", kind="formula", latex=None, ommlXml="<m:oMath>")
    with pytest.raises(AppError) as excinfo:
        _render(env, "student", rich=_rich_with_extra_block(env, broken))
    assert excinfo.value.code == "INVALID_REQUEST"
    assert "f-xml" in str(excinfo.value)


def test_non_math_omml_root_is_rejected(env):
    wrong_root = FormulaBlock(
        id="f-wrong",
        kind="formula",
        latex=None,
        ommlXml=(
            '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            "<m:oMathParaPr/><m:oMath/><m:oMath/></m:oMathPara>"
        ),
    )
    with pytest.raises(AppError) as excinfo:
        _render(env, "student", rich=_rich_with_extra_block(env, wrong_root))
    assert excinfo.value.code == "INVALID_REQUEST"
    assert "f-wrong" in str(excinfo.value)


# --------------------------------------------------------------------------- 产物自洽


def test_rendered_document_reopens_with_all_stem_content(env):
    data = _render(env, "teacher")
    document = _open(data)
    text = document_text(document)
    for block in env.rich.stem_blocks:
        if isinstance(block, ParagraphBlock):
            assert block.text in text, block.id
    assert len(document.tables) == 2
    # 内容之外的一切都是新包自己的部件（每个 r:embed 都有对应关系）
    rels = document.part.rels
    assert all(rid in rels for rid in embed_rids(data))


def test_render_accepts_payload_dict_and_requires_assets(env):
    payload = env.rich.model_dump(by_alias=True)
    data = render_rich_document(rich=payload, assets=env.store, variant="student")
    assert QUESTION_ONE in document_text(_open(data))
    with pytest.raises(AppError) as excinfo:
        render_rich_document(rich=payload, assets=object(), variant="student")  # type: ignore[arg-type]
    assert excinfo.value.code == "INVALID_REQUEST"
    with pytest.raises(AppError) as excinfo:
        render_rich_document(rich="not-rich", assets=env.store, variant="student")  # type: ignore[arg-type]
    assert excinfo.value.code == "INVALID_REQUEST"


def test_asset_ref_type_used_by_contract(env):
    """解析产物用 blob_key 作 assetId；RichAsset 同时给出 sha256 与 mediaType（供登记）。"""
    image = next(block for block in env.parsed.blocks if isinstance(block, ImageBlock))
    asset = next(asset for asset in env.rich.assets if asset.asset_id == image.asset_id)
    assert isinstance(asset, RichAsset)
    assert asset.sha256 == hashlib.sha256(env.store.read(image.asset_id)).hexdigest()
    assert asset.media_type == "image/png"
