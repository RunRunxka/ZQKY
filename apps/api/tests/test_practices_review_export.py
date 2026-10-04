import asyncio
import io
import sqlite3
import pytest
from openpyxl import load_workbook
from app.contracts import b4
from app.core.exceptions import AppError
from tests.practices_support import PracticesScene, package_parts, marked_png


@pytest.fixture
async def scene(tmp_path):
    value=await PracticesScene.create(tmp_path)
    yield value
    value.integrity()


def reviewed(scene, *, count=1):
    questions=[scene.question(rich=scene.rich()) for _ in range(count)]
    # Distinct surfaces with the same shared material for whole-package dedup.
    if count>1:
        rich=scene.rich();rich["stemBlocks"][0]["text"]="第二完整题干"
        questions[1]=scene.question(rich=rich)
    return scene.review(scene.save(scene.create_set(count=count),[scene.item(q,ordinal=i+1) for i,q in enumerate(questions)]))


@pytest.mark.parametrize("table,sql",[
    ("revision","UPDATE practice_revisions SET title_snapshot='改' WHERE id=?"),
    ("selection","UPDATE practice_selections SET reason_json='{}' WHERE practice_revision_id=?"),
    ("item","UPDATE practice_items SET question_no='改' WHERE practice_revision_id=?"),
    ("link","DELETE FROM practice_item_knowledge WHERE practice_revision_id=?"),
    ("selection-insert","INSERT INTO practice_selections SELECT 'new',practice_revision_id,'extra',99,question_id,question_revision_id,question_content_hash,content_snapshot_json,metadata_snapshot_json,rich_assets_json,reason_json,source_snapshot_json,answer_state FROM practice_selections WHERE practice_revision_id=?"),
    ("item-insert","INSERT INTO practice_items SELECT 'new',practice_revision_id,selection_id,'extra',NULL,'99',99,is_scored,max_score_units,content_json,source_locator_json FROM practice_items WHERE practice_revision_id=?"),
    ("link-insert","INSERT INTO practice_item_knowledge SELECT item_id,practice_revision_id,'extra','extra','extra','math','primary' FROM practice_item_knowledge WHERE practice_revision_id=? LIMIT 1")])
async def test_reviewed_database_sealed_all_children(scene,table,sql):
    practice=reviewed(scene)
    with pytest.raises(sqlite3.IntegrityError):
        with scene.catalog.write_transaction() as conn: conn.execute(sql,(practice.current_revision.practice_revision_id,))


async def test_real_docx_whole_zip_student_teacher_material_formula_image_table(scene):
    practice=reviewed(scene,count=2)
    student=await scene.export(practice)
    assert (await scene.finish(student)).state=="succeeded"
    artifact,data=scene.artifact(practice)
    parts=package_parts(data)
    joined=b"\n".join(parts.values())
    assert b"ANS_PRIVATE" not in joined and b"EXPL_PRIVATE" not in joined
    document=parts["word/document.xml"].decode()
    assert document.count("共享材料唯一出现")==1
    assert "完整题干" in document and "第二完整题干" in document and "选项甲" in document
    assert "m:f" in document and "w:gridSpan" in document and "单元乙" in document
    assert any(x.startswith("word/media/") for x in parts)
    assert artifact.sha256==__import__("hashlib").sha256(data).hexdigest()
    teacher=await scene.export(practice,variant="teacher")
    assert (await scene.finish(teacher)).state=="succeeded"
    _,teacherdata=scene.artifact(practice,teacher.export_id)
    teach=b"\n".join(package_parts(teacherdata).values())
    assert b"ANS_PRIVATE" in teach and b"EXPL_PRIVATE" in teach
    scene.questions.archive_question(practice.current_revision.items[0].question_id)
    assert (await scene.export(practice,submission="same-fixed" )).reused is True


async def test_missing_answer_teacher_label_and_idempotent_job(scene):
    practice=scene.review(scene.save(scene.create_set(),[scene.item(scene.question(answer=False))]))
    receipt=await scene.export(practice,variant="teacher")
    repeated=await scene.export(practice,variant="teacher")
    assert repeated.replayed and repeated.export_id==receipt.export_id
    assert (await scene.export(practice,variant="teacher",submission="another")).reused
    assert scene.count("practice_exports")==1
    assert (await scene.finish(receipt)).state=="succeeded"
    _,data=scene.artifact(practice)
    assert "未提供答案" in package_parts(data)["word/document.xml"].decode()


async def test_bad_asset_failure_and_retry_original_frozen_input(scene):
    practice=reviewed(scene)
    receipt=await scene.export(practice)
    sha=practice.current_revision.items[0].content.assets[0].sha256
    path=scene.assets.path_of("blobs/"+sha)
    original=path.read_bytes();path.write_bytes(b"corrupt")
    failed=await scene.finish(receipt)
    assert failed.state=="failed" and scene.count("export_artifacts")==0
    frozen=failed.frozen_input
    path.write_bytes(original)
    retried=scene.store.retry(failed.job_id)
    assert retried.frozen_input==frozen
    assert (await scene.finish(receipt)).state=="succeeded"
    assert scene.count("export_artifacts")==1


async def test_publish_failure_rolls_metadata_with_job_then_retry(scene,monkeypatch):
    practice=reviewed(scene)
    receipt=await scene.export(practice)
    original=scene.service.file_assets.create_in
    def fail(conn,**kw):
        original(conn,**kw)
        raise RuntimeError("injected after asset metadata")
    baseline=scene.count("file_assets")
    monkeypatch.setattr(scene.service.file_assets,"create_in",fail)
    assert (await scene.finish(receipt)).state=="failed"
    assert scene.count("file_assets")==baseline and scene.count("export_artifacts")==0
    monkeypatch.setattr(scene.service.file_assets,"create_in",original)
    scene.store.retry(receipt.job.job_id)
    assert (await scene.finish(receipt)).state=="succeeded"


async def test_cancelled_queued_no_publication(scene):
    practice=reviewed(scene)
    receipt=await scene.export(practice)
    scene.store.request_cancel(receipt.job.job_id)
    assert (await scene.finish(receipt)).state=="cancelled"
    assert scene.count("export_artifacts")==0


async def test_template_freezes_actual_current_assessment_leaves_roster(scene):
    practice=reviewed(scene)
    conversion=scene.conversion(practice)
    with scene.catalog.write_transaction() as conn:
        conn.execute("UPDATE students SET name='=SECRET()',student_no='0000123' WHERE id='s000'")
    # Existing conversion participant snapshot remains the server-created name A.
    receipt=await scene.export(practice,variant="score_template",assessment_id=conversion.assessment_id)
    from app.contracts.assessments import ParticipantAddRequest
    detail=scene.assessments.get_assessment(conversion.assessment_id)
    scene.assessments.add_participants(conversion.assessment_id,ParticipantAddRequest(submissionId="supplement",expectedRevision=detail.assessment.revision,
        participants=[dict(studentId="s000",classId="class",attendance="absent",attemptNo=2)]))
    assert (await scene.finish(receipt)).state=="succeeded"
    _,data=scene.artifact(practice)
    workbook=load_workbook(io.BytesIO(data))
    sheet=workbook["成绩"]
    assert sheet.max_row==2 and sheet.cell(2,1).value=="00000" and sheet.cell(2,2).value=="A"
    assert sheet.cell(2,1).data_type=="s" and sheet.cell(2,5).value is None
    mapping=list(workbook["固定映射"].values)
    assert any(conversion.paper_revision_id in row for row in mapping)
    second=await scene.export(practice,variant="score_template",submission="new-snapshot",assessment_id=conversion.assessment_id)
    assert not second.reused and second.input_hash!=receipt.input_hash
    assert (await scene.finish(second)).state=="succeeded"
    _,data2=scene.artifact(practice,second.export_id)
    sheet2=load_workbook(io.BytesIO(data2))["成绩"]
    assert sheet2.max_row==3 and sheet2.cell(3,2).data_type=="s"
    assert sheet2.cell(3,2).value=="=SECRET()" and sheet2.cell(3,1).value=="0000123"
    with pytest.raises(AppError): await scene.export(practice,variant="score_template",submission="wrong",assessment_id="assessment")


async def test_multileaf_reordered_docx_template_and_paper_exact_numbers_scores(scene):
    rich=scene.rich()
    q=scene.question(rich=rich)
    simple=scene.question("后调到第一题")
    nodes=[dict(nodeKey="root",parentNodeKey=None,questionNo="16",ordinal=1,isScored=False,maxScore=None,knowledgePointIds=[],sourceBlockIds=["stem"]),
        dict(nodeKey="one",parentNodeKey="root",questionNo="16(1)",ordinal=2,isScored=True,maxScore="0.75",knowledgePointIds=["k1"],sourceBlockIds=["f","t"]),
        dict(nodeKey="two",parentNodeKey="root",questionNo="16(2)",ordinal=3,isScored=True,maxScore="0.50",knowledgePointIds=["k2"],sourceBlockIds=["i","a","b"])]
    simple_item=scene.item(simple,ordinal=2)
    simple_item["itemStructure"]["nodes"][0]["questionNo"]="自定义-17"
    practice=scene.review(scene.save(scene.create_set(count=2),[scene.item(q,ordinal=7,nodes=nodes),simple_item]))
    numbers=[x.question_no for x in practice.current_revision.items]
    assert numbers==["自定义-17","16","16(1)","16(2)"]
    assert set(practice.current_revision.items[3].content.option_blocks)=={"A","B"}
    receipt=await scene.export(practice)
    assert (await scene.finish(receipt)).state=="succeeded"
    _,data=scene.artifact(practice,receipt.export_id)
    document=package_parts(data)["word/document.xml"].decode()
    assert all("题号 "+number in document for number in numbers)
    assert "满分 0.75" in document and "满分 0.50" in document and "不计分" in document
    conversion=scene.conversion(practice)
    with scene.catalog.read_connection() as conn:
        leaves=[dict(x) for x in conn.execute("SELECT question_no,max_score_units FROM paper_items WHERE paper_revision_id=? AND is_scored=1 ORDER BY ordinal",(conversion.paper_revision_id,))]
    assert leaves==[{"question_no":"自定义-17","max_score_units":125},{"question_no":"16(1)","max_score_units":75},{"question_no":"16(2)","max_score_units":50}]
    template=await scene.export(practice,variant="score_template",assessment_id=conversion.assessment_id)
    assert (await scene.finish(template)).state=="succeeded"
    _,xlsx=scene.artifact(practice,template.export_id)
    assert list(load_workbook(io.BytesIO(xlsx))["成绩"].values)[0][4:]==tuple(x["question_no"] for x in leaves)


async def test_review_archive_during_lock_free_asset_preflight_no_publish(scene,monkeypatch):
    q=scene.question(rich=scene.rich())
    practice=scene.save(scene.create_set(),[scene.item(q)])
    original=scene.service.read_question_asset
    def archive(asset_id):
        assert not scene.coordinator.busy
        with scene.coordinator.publication(operation="question.archive"):
            scene.questions.archive_question(q.question_id)
        return original(asset_id)
    monkeypatch.setattr(scene.service,"read_question_asset",archive)
    with pytest.raises(AppError) as error:scene.review(practice)
    assert error.value.code=="PRACTICE_REFERENCE_CHANGED"
    current=scene.service.get_practice(practice.practice_set_id)
    assert current.revision==1 and current.current_revision.state=="draft"


async def test_teacher_only_image_never_in_student_entire_zip(scene):
    import hashlib
    rich=scene.rich()
    payload=marked_png("TEACHER_ONLY_ASSET");sha=hashlib.sha256(payload).hexdigest();aid="blobs/"+sha
    scene.source_assets[aid]=(payload,"image/png")
    rich["assets"].append(dict(assetId=aid,sha256=sha,mediaType="image/png"))
    rich["answerBlocks"].append(dict(id="teacher-image",kind="image",assetId=aid,width=10,height=10))
    practice=scene.review(scene.save(scene.create_set(),[scene.item(scene.question(rich=rich))]))
    student=await scene.export(practice)
    assert (await scene.finish(student)).state=="succeeded"
    _,data=scene.artifact(practice,student.export_id)
    assert b"TEACHER_ONLY_ASSET" not in b"\n".join(package_parts(data).values())
    teacher=await scene.export(practice,variant="teacher")
    assert (await scene.finish(teacher)).state=="succeeded"
    _,data=scene.artifact(practice,teacher.export_id)
    assert b"TEACHER_ONLY_ASSET" in b"\n".join(package_parts(data).values())


@pytest.mark.parametrize("field",["sha","media","duplicate-block"])
async def test_bad_rich_declaration_cannot_save_or_review(scene,field):
    rich=scene.rich()
    if field=="sha":rich["assets"][0]["sha256"]="ff"*32
    if field=="media":rich["assets"][0]["mediaType"]="image/jpeg"
    if field=="duplicate-block":rich["answerBlocks"][0]["id"]="stem"
    question=scene.question(rich=rich)
    practice=scene.create_set()
    with pytest.raises(AppError):
        scene.save(practice,[scene.item(question)])
    assert scene.service.get_practice(practice.practice_set_id).revision==0


@pytest.mark.parametrize("same_bytes",[True,False])
async def test_material_image_identity_sha_not_source_alias(scene,same_bytes):
    import copy
    import hashlib
    first=scene.rich()
    first["sharedMaterials"][0]["blocks"].append(dict(id="material-image",kind="image",assetId=first["assets"][0]["assetId"],width=10,height=10))
    second=copy.deepcopy(first)
    second["stemBlocks"][0]["text"]="独立第二题题面"
    payload=scene.source_assets[first["assets"][0]["assetId"]][0] if same_bytes else marked_png("different-SHA")
    sha=hashlib.sha256(payload).hexdigest()
    # Both bare SHA and managed keys are supported source aliases.
    aid=sha;scene.source_assets[aid]=(payload,"image/png")
    second["assets"].append(dict(assetId=aid,sha256=sha,mediaType="image/png"))
    second["sharedMaterials"][0]["blocks"][1]["assetId"]=aid
    q1=scene.question(rich=first);q2=scene.question(rich=second)
    practice=scene.review(scene.save(scene.create_set(count=2),[scene.item(q1),scene.item(q2,ordinal=2)]))
    receipt=await scene.export(practice)
    assert (await scene.finish(receipt)).state=="succeeded"
    _,data=scene.artifact(practice,receipt.export_id)
    document=package_parts(data)["word/document.xml"].decode()
    assert document.count("共享材料唯一出现")== (1 if same_bytes else 2)
