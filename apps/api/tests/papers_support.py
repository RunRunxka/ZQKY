"""T40 原卷测试台（非 ``test_`` 前缀，不参与 pytest 收集）。

- ``build_paper_docx``：程序化生成 DOCX 样本（不读正式数据目录、不联网）：共同材料
  2 段、真实 PNG 图片、2×2 横向合并表格、两级小题（``16.`` + ``16(1)``/``16(2)``）、
  独立计分题（含行内公式）、独立公式段；可选追加未知对象（``w:object``）、无分值子题
  （``19.（5 分）`` + 独立 ``（1）``/``（2）``）与重复题号；
- ``build_no_question_docx`` / ``build_empty_docx`` / ``broken_docx_bytes``：坏文件与
  "没有题号"的边界样本；
- ``PapersHarness``：真 ``TeachingCatalog`` / ``KnowledgeCatalog`` / ``AssetStore`` /
  ``FileAssetsRepository`` / ``PublicationCoordinator`` + 受控模型替身 + 真实任务引擎；
- ``items_payload``/``blocks_payload``/``issues_payload``：由已落库草稿构造整表替换请求，
  供确认链测试使用；``FakeProvider``/``FakeResolver``：模型替身（不触网、不含凭证）。
"""

from __future__ import annotations

import asyncio
import hashlib
import io
import sqlite3
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

from app.contracts.papers import PaperProposalJobRequest
from app.core.exceptions import AppError
from app.main import create_app
from app.providers.llm.base import FINISH_STOP, LLMConfig, LLMResponse
from app.repositories.assets.file_assets import FileAssetsRepository
from app.repositories.jobs.repository import JobRecord, JobStore
from app.repositories.knowledge.catalog import KnowledgeCatalog
from app.repositories.knowledge.points import (
    KnowledgePointRepository,
    SubjectRepository,
)
from app.repositories.teaching.catalog import TeachingCatalog
from app.schemas.model_config import ModelProtocol
from app.services.assets.store import AssetStore
from app.services.jobs.engine import JobEngine
from app.services.model_runtime import ChatModelHandle
from app.services.papers.service import PaperService, build_paper_service
from app.services.publication import PublicationCoordinator
from tests.conftest import make_settings
from tests.rich_content_support import EMU_PER_PIXEL, minimal_png

LOCAL_PROFILE = "chat-model-local"
#: 写入替身凭证的哨兵值：冻结输入/任务快照/响应里出现它即视为泄漏
FAKE_API_KEY = "sk-test-should-never-be-persisted"
SUBJECT_ID = "math"
OTHER_SUBJECT_ID = "physics"

M_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

TITLE_TEXT = "智启课源 期中数学试卷"
MATERIAL_HEADING = "阅读材料：二次函数"
MATERIAL_BODY = "材料正文：二次函数 y=x^2 的图象开口向上，顶点在原点。"
TABLE_HEADER = "得分表"
TABLE_LEFT = "题号"
TABLE_RIGHT = "满分"
ITEM_ROOT_TEXT = "16.（12 分）阅读材料，完成下列小题。"
ITEM_CHILD_ONE = "16(1)（4 分）求抛物线的顶点坐标。"
ITEM_CHILD_TWO = "16(2)（8 分）证明该函数在 x>0 时单调递增。"
ITEM_SEVENTEEN_PREFIX = "17.（6 分）已知 "
ITEM_SEVENTEEN_SUFFIX = "，求它的值。"
ITEM_EIGHTEEN = "18.（3 分）写出二次函数的对称轴。"
INLINE_LATEX = r"x^{2}+1=2"
STANDALONE_LATEX = r"\frac{a}{b}=c"
BARE_ROOT_TEXT = "19.（5 分）回答下列问题。"
BARE_CHILD_ONE = "（1）第一小问。"
BARE_CHILD_TWO = "（2）第二小问。"
DUPLICATE_ROOT_TEXT = "16.（12 分）重复出现的题号。"

PNG_BYTES = minimal_png()
IMAGE_WIDTH_PX = 192
IMAGE_HEIGHT_PX = 96

#: 计分叶子与满分（样本口径：4 + 8 + 6 + 3 = 21）
SAMPLE_LEAF_SCORES = {"16(1)": "4", "16(2)": "8", "17": "6", "18": "3"}
SAMPLE_TOTAL_UNITS = 2100
TERMINAL_STATES = frozenset({"succeeded", "failed", "cancelled", "interrupted"})
_UNSET = object()


# --------------------------------------------------------------------------- DOCX 样本


def build_paper_docx(
    path: Path,
    *,
    with_unknown_object: bool = False,
    with_unscored_children: bool = False,
    with_duplicate_number: bool = False,
) -> Path:
    """生成 T40 样本 DOCX；返回写入的路径（调用方用 ``tmp_path``）。"""
    import math2docx
    from docx import Document
    from docx.oxml import OxmlElement, parse_xml
    from docx.shared import Emu

    target = Path(path)
    document = Document()
    document.add_paragraph(TITLE_TEXT)
    document.add_paragraph(MATERIAL_HEADING)
    document.add_paragraph(MATERIAL_BODY)
    document.add_picture(
        io.BytesIO(PNG_BYTES),
        width=Emu(IMAGE_WIDTH_PX * EMU_PER_PIXEL),
        height=Emu(IMAGE_HEIGHT_PX * EMU_PER_PIXEL),
    )

    table = document.add_table(rows=2, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = TABLE_HEADER
    table.cell(1, 0).text = TABLE_LEFT
    table.cell(1, 1).text = TABLE_RIGHT

    document.add_paragraph(ITEM_ROOT_TEXT)
    document.add_paragraph(ITEM_CHILD_ONE)
    document.add_paragraph(ITEM_CHILD_TWO)

    mixed = document.add_paragraph()
    mixed.add_run(ITEM_SEVENTEEN_PREFIX)
    math2docx.add_math(mixed, INLINE_LATEX)
    mixed.add_run(ITEM_SEVENTEEN_SUFFIX)

    standalone = document.add_paragraph()
    math2docx.add_math(standalone, STANDALONE_LATEX)
    math_element = standalone._p[-1]  # noqa: SLF001 - 把公式包进 m:oMathPara
    wrapper = OxmlElement("m:oMathPara")
    standalone._p.replace(math_element, wrapper)  # noqa: SLF001
    wrapper.append(math_element)

    document.add_paragraph(ITEM_EIGHTEEN)

    if with_unscored_children:
        document.add_paragraph(BARE_ROOT_TEXT)
        document.add_paragraph(BARE_CHILD_ONE)
        document.add_paragraph(BARE_CHILD_TWO)

    if with_duplicate_number:
        document.add_paragraph(DUPLICATE_ROOT_TEXT)

    if with_unknown_object:
        object_paragraph = document.add_paragraph()
        object_run = object_paragraph.add_run()
        object_run._r.append(  # noqa: SLF001 - 自造嵌入对象（未知对象样本）
            parse_xml(
                f'<w:object xmlns:w="{W_NAMESPACE}" xmlns:v="urn:schemas-microsoft-com:vml"'
                f' xmlns:o="urn:schemas-microsoft-com:office:office"'
                f' xmlns:r="{R_NAMESPACE}">'
                '<v:shape id="ole1" style="width:40pt;height:20pt" type="#_x0000_t75">'
                '<v:imagedata r:id="rIdMissing" o:title="embedded"/>'
                "</v:shape>"
                '<o:OLEObject Type="Embed" ProgID="Package" ShapeID="ole1"/>'
                "</w:object>"
            )
        )

    document.save(str(target))
    return target


def build_no_question_docx(path: Path) -> Path:
    """没有题号的 DOCX：只有标题与一张表（用于 NO_ITEMS_DETECTED 边界）。"""
    from docx import Document

    target = Path(path)
    document = Document()
    document.add_paragraph("这是一份没有题号的文档。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "甲"
    table.cell(0, 1).text = "乙"
    document.save(str(target))
    return target


def build_empty_docx(path: Path) -> Path:
    """只有空段落的 DOCX：零内容块（整文件不可解析）。"""
    from docx import Document

    target = Path(path)
    document = Document()
    document.add_paragraph("")
    document.save(str(target))
    return target


def broken_docx_bytes() -> bytes:
    return b"this is not a zip / docx file"


def sha256_of_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


# --------------------------------------------------------------------------- 模型替身


class FakeProvider:
    """Provider 替身：``handler(request)`` 优先，其次按序弹出 ``replies``。

    ``handler`` 可返回字符串 / ``LLMResponse`` / 异常 / ``None``（默认空建议）。
    """

    def __init__(self, replies: Sequence[Any] | None = None, handler=None) -> None:
        self.replies = list(replies or [])
        self.handler = handler
        self.calls: list[Any] = []

    async def complete(self, config: LLMConfig, request: Any, *, transport: Any = None) -> LLMResponse:
        self.calls.append(request)
        value = self.handler(request) if self.handler is not None else (
            self.replies.pop(0) if self.replies else None
        )
        if value is None:
            return LLMResponse(text='{"items":[]}', finishReason=FINISH_STOP)
        if isinstance(value, BaseException):
            raise value
        if isinstance(value, LLMResponse):
            return value
        return LLMResponse(text=str(value), finishReason=FINISH_STOP)


def make_handle(
    profile_id: str = LOCAL_PROFILE,
    *,
    provider: Any | None = None,
    api_key: str | None = FAKE_API_KEY,
) -> ChatModelHandle:
    config = LLMConfig(
        protocol=ModelProtocol.openai_chat,
        baseUrl="http://127.0.0.1:11434/v1",
        modelId="qwen2.5:7b",
        apiKey=api_key,
        apiFormat="openai_chat",
        connectionId="conn-1",
        modelProfileId=profile_id,
    )
    return ChatModelHandle(
        profile_id=profile_id,
        model_id="qwen2.5:7b",
        provider=provider if provider is not None else FakeProvider(),
        config=config,
        max_output_tokens=2048,
    )


class FakeResolver:
    """``(profileId) -> ChatModelHandle`` 替身：可换映射、可抛错、记录调用。"""

    def __init__(
        self,
        profiles: Mapping[str, ChatModelHandle] | None = None,
        *,
        error: BaseException | None = None,
    ) -> None:
        self.profiles = dict(profiles or {})
        self.error = error
        self.calls: list[str] = []

    def __call__(self, profile_id: str) -> ChatModelHandle:
        self.calls.append(profile_id)
        if self.error is not None:
            raise self.error
        handle = self.profiles.get(profile_id)
        if handle is None:
            raise AppError("模型配置不存在。", code="MODEL_PROFILE_NOT_FOUND", status_code=404)
        return handle


# --------------------------------------------------------------------------- 请求构造


def items_payload(
    items: Sequence[Any],
    *,
    knowledge: Mapping[str, Sequence[str]] | None = None,
    drop_questions: Sequence[str] = (),
    score_overrides: Mapping[str, str | None] | None = None,
    parent_overrides: Mapping[str, str | None] | None = None,
) -> list[dict[str, Any]]:
    """由 ``PaperItemView`` 列表构造整表替换的 items（默认原样回写）。"""
    knowledge = dict(knowledge or {})
    score_overrides = dict(score_overrides or {})
    parent_overrides = dict(parent_overrides or {})
    dropped = set(drop_questions)
    payload: list[dict[str, Any]] = []
    for item in items:
        if item.question_no in dropped:
            continue
        entry: dict[str, Any] = {
            "itemId": item.item_id,
            "questionNo": item.question_no,
            "ordinal": item.ordinal,
            "isScored": item.is_scored,
            "content": dict(item.content),
            "sourceLocator": dict(item.source_locator),
            "knowledge": [
                {"knowledgePointId": point_id, "role": "primary"}
                for point_id in knowledge.get(item.question_no, ())
            ],
        }
        if item.parent_item_id is not None:
            entry["parentItemId"] = item.parent_item_id
        if item.question_no in parent_overrides:
            entry["parentItemId"] = parent_overrides[item.question_no]
        if item.max_score is not None:
            entry["maxScore"] = item.max_score
        if item.question_no in score_overrides:
            override = score_overrides[item.question_no]
            if override is None:
                entry.pop("maxScore", None)
                entry["isScored"] = False
            else:
                entry["maxScore"] = override
                entry["isScored"] = True
        payload.append(entry)
    return payload


def items_payload_from_json(
    items: Sequence[Mapping[str, Any]],
    *,
    knowledge: Mapping[str, Sequence[str]] | None = None,
) -> list[dict[str, Any]]:
    """由 HTTP 视图的 camelCase items 构造整表替换 payload（HTTP 端到端测试用）。"""
    knowledge = dict(knowledge or {})
    payload: list[dict[str, Any]] = []
    for item in items:
        entry: dict[str, Any] = {
            "itemId": item["itemId"],
            "questionNo": item["questionNo"],
            "ordinal": item["ordinal"],
            "isScored": item["isScored"],
            "content": item.get("content") or {},
            "sourceLocator": item.get("sourceLocator") or {},
            "knowledge": [
                {"knowledgePointId": point_id, "role": "primary"}
                for point_id in knowledge.get(item["questionNo"], ())
            ],
        }
        if item.get("parentItemId"):
            entry["parentItemId"] = item["parentItemId"]
        if item.get("maxScore") is not None:
            entry["maxScore"] = item["maxScore"]
        payload.append(entry)
    return payload


def blocks_payload(blocks: Sequence[Any]) -> list[dict[str, Any]]:
    """由 ``PaperSourceBlockView`` 列表构造块处置请求（默认保持原样）。"""
    return [
        {
            "blockId": block.block_id,
            "disposition": block.disposition,
            **({"itemId": block.item_id} if block.item_id else {}),
            **({"excludeReason": block.exclude_reason} if block.exclude_reason else {}),
        }
        for block in blocks
    ]


def issues_payload(
    issues: Sequence[Any],
    *,
    resolve_all: bool = True,
    target_block_id: str | None = None,
) -> list[dict[str, Any]]:
    """把问题清单构造成**结构化**处置请求（B3/G0 起 resolution 不再接受任意 JSON）。

    默认用 ``supplement_text`` 补录到问题自身定位的块，或调用方给出的段落块；
    内容损失类问题必须补录（不能仅以 exclude 放行），因此这里统一走补录路径。
    """
    payload: list[dict[str, Any]] = []
    for issue in issues:
        if not resolve_all:
            break
        target = issue.block_id or target_block_id
        assert target, "issues_payload 需要问题自身带 block_id 或显式 target_block_id"
        payload.append(
            {
                "issueId": issue.issue_id,
                "status": "resolved",
                "resolution": {
                    "kind": "supplement_text",
                    "targetBlockId": target,
                    "text": "教师已人工核对并补录（测试替身）",
                },
            }
        )
    return payload


# --------------------------------------------------------------------------- 测试台


class PapersHarness:
    """T40 测试台：真四库组件 + 受控模型替身 + 真实任务引擎（临时目录，不碰正式数据）。"""

    def __init__(
        self,
        tmp_path: Path,
        *,
        provider: Any | None = None,
        resolver: Any = _UNSET,
        model_resolver: Any = _UNSET,
    ) -> None:
        self.settings = make_settings(tmp_path / "data")
        self.catalog = TeachingCatalog(self.settings.teaching_root / "teaching.sqlite3")
        self.catalog.migrate()
        self.knowledge = KnowledgeCatalog(self.settings.knowledge_root / "knowledge.sqlite3")
        self.knowledge.migrate()
        self.assets = AssetStore(self.settings.assets_root)
        self.file_assets = FileAssetsRepository(self.catalog)
        self.coordinator = PublicationCoordinator()
        self.points = KnowledgePointRepository()
        self.subjects = SubjectRepository()
        self.provider = provider if provider is not None else FakeProvider()
        chosen = resolver if resolver is not _UNSET else model_resolver
        if chosen is _UNSET:
            chosen = FakeResolver({LOCAL_PROFILE: make_handle(provider=self.provider)})
        self.model_resolver = chosen
        #: 最近一次 run_proposal_job 里调度任务的异常（发布失败/取消等场景断言用）
        self.last_job_exception: BaseException | None = None

    # ---------------------------------------------------------------- 装配

    def service(self, *, engine: Any = _UNSET, resolver: Any = _UNSET, knowledge: Any = _UNSET) -> PaperService:
        return build_paper_service(
            self.catalog,
            asset_store=self.assets,
            file_assets=self.file_assets,
            knowledge_catalog=self.knowledge if knowledge is _UNSET else knowledge,
            coordinator=self.coordinator,
            model_resolver=self.model_resolver if resolver is _UNSET else resolver,
            job_engine=None if engine is _UNSET else engine,
        )

    def new_engine(self) -> JobEngine:
        """在**当前事件循环**里新建任务引擎（semaphore 与循环绑定，不能跨 loop 复用）。"""
        store = JobStore(
            self.catalog,
            domain="teaching",
            table="workflow_jobs",
            kinds=frozenset({"paper_import", "paper_mapping", "analysis", "lesson_generation", "export"}),
        )
        return JobEngine({"teaching": store}, heavy_limit=2, model_limit=1)

    def create_app_with_service(self, *, provider: Any | None = None) -> tuple[Any, PaperService]:
        """走 ``create_app`` 的真实装配路径（CTRL 装配前的本地挂载）+ 受控模型替身。"""
        app = create_app(make_settings(self.settings.data_dir))
        handle = make_handle(provider=provider if provider is not None else FakeProvider())
        service = build_paper_service(
            app.state.teaching,
            asset_store=app.state.asset_store,
            file_assets=app.state.file_assets,
            knowledge_catalog=app.state.knowledge,
            coordinator=app.state.publication_coordinator,
            model_resolver=FakeResolver({LOCAL_PROFILE: handle}),
            job_engine=app.state.job_engine,
        )
        app.state.paper_service = service
        install_papers_router(app)
        return app, service

    # ---------------------------------------------------------------- 知识点

    def add_point(
        self,
        code: str,
        name: str,
        *,
        subject_id: str = SUBJECT_ID,
        status: str = "active",
        aliases: Sequence[str] = (),
    ):
        with self.knowledge.write_transaction() as conn:
            self.subjects.ensure_subject(conn, subject_id=subject_id, name=subject_id)
            return self.points.create_point(
                conn,
                subject_id=subject_id,
                code=code,
                name=name,
                status=status,
                aliases=list(aliases),
            )

    def archive_point(self, point_id: str) -> None:
        with self.knowledge.write_transaction() as conn:
            point = self.points.require_point(conn, point_id)
            self.points.set_status(
                conn, point_id, expected_revision=point.revision, archived=True
            )

    def rename_point(self, point_id: str, *, name: str) -> str:
        """追加内容修订并返回新的修订 id（用于"名称快照取当前修订"断言）。"""
        with self.knowledge.write_transaction() as conn:
            point = self.points.require_point(conn, point_id)
            updated = self.points.update_point(
                conn, point_id, expected_revision=point.revision, name=name
            )
            return updated.revision_id

    # ---------------------------------------------------------------- 数据库直读（断言用）

    def raw_rows(self, sql: str, params: Sequence[Any] = ()) -> list[dict[str, Any]]:
        connection = sqlite3.connect(str(self.catalog.db_path))
        connection.row_factory = sqlite3.Row
        try:
            return [dict(row) for row in connection.execute(sql, list(params)).fetchall()]
        finally:
            connection.close()

    def raw_execute(self, sql: str, params: Sequence[Any] = ()) -> None:
        """绕过服务直写（触发器保护断言用）；失败时抛 ``sqlite3.IntegrityError``。"""
        connection = sqlite3.connect(str(self.catalog.db_path))
        connection.execute("PRAGMA foreign_keys = ON")
        try:
            connection.execute(sql, list(params))
            connection.commit()
        finally:
            connection.close()

    def count(self, table: str, where: str = "", params: Sequence[Any] = ()) -> int:
        clause = f" WHERE {where}" if where else ""
        return self.raw_rows(f"SELECT COUNT(*) AS n FROM {table}{clause}", params)[0]["n"]

    # ---------------------------------------------------------------- 流程

    def import_docx(
        self,
        service: PaperService,
        path: Path,
        *,
        file_name: str | None = None,
        subject_id: str = SUBJECT_ID,
        title: str | None = None,
    ):
        return service.create_import(
            file_name=file_name or Path(path).name,
            content=Path(path).read_bytes(),
            media_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            subject_id=subject_id,
            title=title,
        )

    def run_proposal_job(
        self,
        paper_id: str,
        *,
        expected_revision: int,
        profile_id: str = LOCAL_PROFILE,
        provider: Any | None = None,
        timeout_seconds: float = 20.0,
    ) -> tuple[PaperService, JobRecord]:
        """在独立事件循环里建任务并等待执行器收尾；返回 ``(service, job 记录)``。

        监听 ``engine.schedule`` 取到调度任务并 await 它，以便在引擎收敛完成后再断言：
        B3/G0 起，发布事务失败（如过期 stale）由引擎用**本次执行开始时的原 JobLease**
        在新短事务里收敛终态（取消优先 → ``cancelled``；失权 → 零写入），不再把任务
        留在 ``running`` 等重启收敛；被引擎吞掉/收敛的异常仍记录在
        ``self.last_job_exception``（无异常时为 ``None``，终态与错误码以
        ``record.state`` / ``record.error_code`` 为准）。
        """
        selected_provider = provider if provider is not None else self.provider
        resolver = FakeResolver({profile_id: make_handle(profile_id, provider=selected_provider)})

        async def main() -> tuple[PaperService, JobRecord]:
            engine = self.new_engine()
            scheduled: list[Any] = []
            original = engine.schedule

            def spy(domain, job_id, executor, *, uses_model=False):
                task = original(domain, job_id, executor, uses_model=uses_model)
                scheduled.append(task)
                return task

            engine.schedule = spy  # type: ignore[method-assign]
            service = self.service(engine=engine, resolver=resolver)
            view = await service.create_proposal_job(
                paper_id,
                PaperProposalJobRequest(
                    modelProfileId=profile_id, expectedRevision=expected_revision
                ),
            )
            store = engine.store("teaching")
            waited = 0.0
            while not scheduled and waited < timeout_seconds:
                await asyncio.sleep(0.01)
                waited += 0.01
            if not scheduled:
                raise AssertionError("任务没有被调度")
            results = await asyncio.gather(*scheduled, return_exceptions=True)
            self.last_job_exception = next(
                (value for value in results if isinstance(value, BaseException)), None
            )
            return service, store.get(view.job_id)

        return asyncio.run(main())

    def rerun_proposal_job(
        self,
        job_id: str,
        *,
        profile_id: str = LOCAL_PROFILE,
        provider: Any | None = None,
        timeout_seconds: float = 20.0,
    ) -> tuple[PaperService, JobRecord]:
        """重试路径：``retry`` 回 ``queued`` 后**沿用原冻结输入/模型指纹**重跑执行器。

        恢复不重新冻结（不换模型、不换版本），发布事务内仍会复核草稿版本。
        """
        selected = provider if provider is not None else self.provider
        resolver = FakeResolver({profile_id: make_handle(profile_id, provider=selected)})

        async def main() -> tuple[PaperService, JobRecord]:
            from app.services.papers.proposals import ProposalRunner

            engine = self.new_engine()
            store = engine.store("teaching")
            store.retry(job_id)
            service = self.service(engine=engine, resolver=resolver)
            runner = ProposalRunner(
                resolve_model=service._resolve_model_for_job,  # noqa: SLF001 - 真实执行器
                publish_proposal=service._publish_proposal,  # noqa: SLF001
            )
            task = engine.schedule("teaching", job_id, runner, uses_model=True)
            results = await asyncio.gather(task, return_exceptions=True)
            self.last_job_exception = next(
                (value for value in results if isinstance(value, BaseException)), None
            )
            return service, store.get(job_id)

        return asyncio.run(main())


def install_papers_router(app) -> None:
    """挂载 T40 路由；通配 501 占位必须留在最后（否则抢先匹配）。"""
    from app.api.v1 import papers as papers_route

    paths = {getattr(route, "path", None) for route in app.router.routes}
    if "/api/v1/papers" not in paths:
        app.include_router(papers_route.router, prefix="/api/v1")
    routes = list(app.router.routes)
    catch = [
        route for route in routes if getattr(route, "name", "") == "feature_not_implemented"
    ]
    if catch:
        app.router.routes[:] = [route for route in routes if route not in catch] + catch


__all__ = [
    "FAKE_API_KEY",
    "FakeProvider",
    "FakeResolver",
    "LOCAL_PROFILE",
    "OTHER_SUBJECT_ID",
    "PNG_BYTES",
    "PapersHarness",
    "SAMPLE_LEAF_SCORES",
    "SAMPLE_TOTAL_UNITS",
    "SUBJECT_ID",
    "TERMINAL_STATES",
    "blocks_payload",
    "broken_docx_bytes",
    "build_empty_docx",
    "build_no_question_docx",
    "build_paper_docx",
    "install_papers_router",
    "issues_payload",
    "items_payload",
    "items_payload_from_json",
    "make_handle",
    "sha256_of_bytes",
]
