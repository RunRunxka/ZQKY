"""T10 富内容测试样本构造器（非 ``test_`` 前缀，不参与 pytest 收集）。

- ``build_rich_docx``：程序化生成 DOCX 样本（不读正式教材目录、不联网）：
  共同材料 2 段（标题 + 阅读材料 + 材料正文）、一张真实 PNG、2×2 横向合并表格、
  2×2 纵向合并表格、两个小题题干、一个行内公式段、一个独立公式段（``m:oMathPara``）、
  段落内未知对象（``w:object``）、正文级未知对象（``w:sdt``）与一个空段落；
- ``minimal_png``：手写的最小合法 PNG（真彩 2×3，无第三方依赖）；
- ``rename_image_relationship``：把样本的图片关系重定向到哨兵 rId（``rId404``），
  用来证明渲染器不会复用源文档的 rId/relationship；
- ``source_omml_strings``：直接从 ``word/document.xml`` 取每个 ``m:oMath`` 的逐字节
  序列化，作为"OMML 原样"断言的参照；
- ``document_text`` / ``distinct_cell_texts`` / ``package_entries`` 等：渲染产物断言助手。
"""

from __future__ import annotations

import hashlib
import io
import re
import struct
import zipfile
import zlib
from pathlib import Path

from lxml import etree

from app.contracts.teaching_loop import RichContentV2, RichOrigin
from app.services.rich_content import rich_content_from_blocks

M_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
IMAGE_RELATIONSHIP_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
)

#: 96 dpi 下 1 像素的 EMU
EMU_PER_PIXEL = 9525

#: 图片尺寸（EMU 取整像素，避免"取整口径"含义不清）
IMAGE_WIDTH_PX = 192
IMAGE_HEIGHT_PX = 96
PNG_WIDTH = 2
PNG_HEIGHT = 3

#: 渲染器测试用的哨兵 rId（源文档不会存在于渲染产物里）
IMAGE_SENTINEL_RID = "rId404"

TITLE_TEXT = "春雨课堂练习"
MATERIAL_HEADING = "阅读材料：春雨"
MATERIAL_BODY = "材料正文：春雨贵如油，雨中的村庄安静下来。"
QUESTION_ONE = "1. 第一小题：材料中雨后的村庄给你怎样的感受？请简要作答。"
QUESTION_TWO_PREFIX = "2. 第二小题：已知 "
QUESTION_TWO_SUFFIX = "，求它的值。"
INLINE_LATEX = r"x^{2}+1=2"
STANDALONE_LATEX = r"\frac{a}{b}=c"
TABLE_H_HEADER = "横向合并表头"
TABLE_H_LEFT = "左下格"
TABLE_H_RIGHT = "右下格"
TABLE_V_MERGED = "纵向合并"
TABLE_V_RIGHT_TOP = "右上格"
TABLE_V_RIGHT_BOTTOM = "右下格"
OPTION_A_TEXT = "A. 选项甲"
OPTION_B_TEXT = "B. 选项乙"
ANSWER_TEXT = "答案：选项乙（学生版不得出现这一段）"
EXPLANATION_TEXT = "解析：因为材料写了春雨贵如油（学生版不得出现这一段）"
SDT_INNER_TEXT = "内容控件里的文字"
ORIGIN_ASSET_ID = "asset-origin-0001"


def minimal_png(
    *, width: int = PNG_WIDTH, height: int = PNG_HEIGHT, rgb: tuple[int, int, int] = (64, 128, 192)
) -> bytes:
    """手写最小合法 PNG（8 位真彩，无 filter）；python-docx 可读出真实像素尺寸。"""

    def chunk(tag: bytes, payload: bytes) -> bytes:
        return (
            struct.pack(">I", len(payload))
            + tag
            + payload
            + struct.pack(">I", zlib.crc32(tag + payload) & 0xFFFFFFFF)
        )

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(raw, 9))
        + chunk(b"IEND", b"")
    )


PNG_BYTES = minimal_png()


def build_rich_docx(path: Path) -> Path:
    """生成 T10 样本 DOCX；返回写入的路径（调用方用 ``tmp_path``）。"""
    import math2docx
    from docx import Document
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import qn
    from docx.shared import Emu

    target = Path(path)
    document = Document()
    document.add_paragraph(TITLE_TEXT)
    document.add_paragraph(MATERIAL_HEADING)
    document.add_paragraph(MATERIAL_BODY)
    document.add_picture(
        io.BytesIO(PNG_BYTES),
        width=Emu(IMAGE_WIDTH_PX * EMU_PER_PIXEL),
        height=Emu(IMAGE_HEIGHT_PX * EMU_PER_PIXEL),
    )

    horizontal = document.add_table(rows=2, cols=2)
    merged_header = horizontal.cell(0, 0).merge(horizontal.cell(0, 1))
    merged_header.text = TABLE_H_HEADER
    horizontal.cell(1, 0).text = TABLE_H_LEFT
    horizontal.cell(1, 1).text = TABLE_H_RIGHT

    vertical = document.add_table(rows=2, cols=2)
    merged_column = vertical.cell(0, 0).merge(vertical.cell(1, 0))
    merged_column.text = TABLE_V_MERGED
    vertical.cell(0, 1).text = TABLE_V_RIGHT_TOP
    vertical.cell(1, 1).text = TABLE_V_RIGHT_BOTTOM

    document.add_paragraph(QUESTION_ONE)

    mixed = document.add_paragraph()
    mixed.add_run(QUESTION_TWO_PREFIX)
    math2docx.add_math(mixed, INLINE_LATEX)
    mixed.add_run(QUESTION_TWO_SUFFIX)

    standalone = document.add_paragraph()
    math2docx.add_math(standalone, STANDALONE_LATEX)
    math_element = standalone._p[-1]  # noqa: SLF001 - 需要把公式包进 m:oMathPara
    wrapper = OxmlElement("m:oMathPara")
    standalone._p.replace(math_element, wrapper)  # noqa: SLF001
    wrapper.append(math_element)

    object_paragraph = document.add_paragraph()
    object_run = object_paragraph.add_run()
    object_run._r.append(  # noqa: SLF001 - 自造嵌入对象（未知对象样本）
        parse_xml(
            f'<w:object xmlns:w="{W_NAMESPACE}" xmlns:v="urn:schemas-microsoft-com:vml"'
            f' xmlns:o="urn:schemas-microsoft-com:office:office"'
            f' xmlns:r="{R_NAMESPACE}">'
            '<v:shape id="ole1" style="width:40pt;height:20pt" type="#_x0000_t75">'
            '<v:imagedata r:id="rIdMissing" o:title="embedded"/>'
            "</v:shape>"
            '<o:OLEObject Type="Embed" ProgID="Package" ShapeID="ole1"/>'
            "</w:object>"
        )
    )

    document.add_paragraph("")

    body = document.element.body
    sdt = parse_xml(
        f'<w:sdt xmlns:w="{W_NAMESPACE}">'
        "<w:sdtPr/><w:sdtContent>"
        f"<w:p><w:r><w:t>{SDT_INNER_TEXT}</w:t></w:r></w:p>"
        "</w:sdtContent></w:sdt>"
    )
    section_properties = body.find(qn("w:sectPr"))
    if section_properties is not None:
        section_properties.addprevious(sdt)
    else:  # pragma: no cover - python-docx 默认模板总有 sectPr
        body.append(sdt)

    document.save(str(target))
    return target


def rename_image_relationship(path: Path, *, sentinel: str = IMAGE_SENTINEL_RID) -> str:
    """把样本第一个图片关系改成哨兵 rId，并让 blip 指向它；返回被替换的原始 rId。

    只重写 ``word/_rels/document.xml.rels`` 与 ``word/document.xml`` 两个条目，其余字节
    原样保留。哨兵关系是**新增**的（原关系保留但不再被引用），因此源文档里有 2 条图片
    关系、渲染产物应当只重建真正被引用的那一条。
    """
    source = Path(path)
    with zipfile.ZipFile(source) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    rels_name = "word/_rels/document.xml.rels"
    document_name = "word/document.xml"
    rels = entries[rels_name].decode("utf-8")
    matched = re.search(
        rf'<Relationship Id="([^"]+)"[^>]*Type="{re.escape(IMAGE_RELATIONSHIP_TYPE)}"'
        r'[^>]*Target="([^"]+)"',
        rels,
    )
    if matched is None:
        raise AssertionError("样本里没有图片关系，无法重命名")
    original_id, target_part = matched.group(1), matched.group(2)
    # 关系元素里属性的顺序由 python-docx 决定，这里按实际出现顺序重取 Target
    element = re.search(
        rf'<Relationship [^>]*Id="{re.escape(original_id)}"[^>]*/>', rels
    )
    if element is None:  # pragma: no cover - 上述匹配成功时必然存在
        raise AssertionError("图片关系元素匹配失败")
    target_matched = re.search(r'Target="([^"]+)"', element.group(0))
    if target_matched is not None:
        target_part = target_matched.group(1)
    entries[rels_name] = rels.replace(
        "</Relationships>",
        f'<Relationship Id="{sentinel}" Type="{IMAGE_RELATIONSHIP_TYPE}"'
        f' Target="{target_part}"/></Relationships>',
    ).encode("utf-8")

    xml = entries[document_name].decode("utf-8")
    embed_attribute = f'r:embed="{original_id}"'
    if embed_attribute not in xml:  # pragma: no cover - 样本必然带图片
        raise AssertionError("document.xml 里没有指向图片关系的 r:embed")
    entries[document_name] = xml.replace(
        embed_attribute, f'r:embed="{sentinel}"'
    ).encode("utf-8")

    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return original_id


def source_omml_strings(path: Path) -> list[str]:
    """原件 ``word/document.xml`` 里每个 ``m:oMath`` 的逐字节序列化（文档顺序）。"""
    with zipfile.ZipFile(Path(path)) as archive:
        xml = archive.read("word/document.xml")
    root = etree.fromstring(xml)
    return [
        etree.tostring(node, encoding="unicode")
        for node in root.iter(f"{{{M_NAMESPACE}}}oMath")
    ]


def source_image_part_rids(path: Path) -> set[str]:
    """原件里指向 image part 的关系 id 集合（含未被引用的关系）。"""
    with zipfile.ZipFile(Path(path)) as archive:
        rels = archive.read("word/_rels/document.xml.rels").decode("utf-8")
    return set(
        re.findall(
            rf'<Relationship Id="([^"]+)"[^>]*Type="{re.escape(IMAGE_RELATIONSHIP_TYPE)}"',
            rels,
        )
    )


def sha256_of_file(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def build_rich_content(
    parsed,
    materials,
    *,
    docx_path: Path,
    with_grid_hint: bool = True,
):
    """由解析结果组装完整 ``RichContentV2``：共同材料 + 题干 + 选项 + 答案 + 解析 + LaTeX 公式。

    走 ``model_dump(by_alias=True)`` → ``model_validate``，因此同时验证线上 JSON 形状。
    ``with_grid_hint=False`` 时不写 ``origin.sourceLocator.tableGrids``（用于验证渲染器
    对表格网格的自主重建）。
    """
    material_ids = {block.id for material in materials for block in material.blocks}
    stem_ids = [block.id for block in parsed.blocks if block.id not in material_ids]
    locator: dict = {"fileName": docx_path.name, "parserVersion": "zqky-rich-v1"}
    if with_grid_hint:
        locator["tableGrids"] = {
            block.id: parsed.locators[block.id]["tableColumns"]
            for block in parsed.blocks
            if block.kind == "table"
        }
    origin = RichOrigin(
        originalAssetId=ORIGIN_ASSET_ID,
        originalSha256=sha256_of_file(docx_path),
        sourceLocator=locator,
    )
    base = rich_content_from_blocks(
        blocks=parsed.blocks,
        materials=materials,
        assets=parsed.assets,
        origin=origin,
        stem_block_ids=stem_ids,
    )
    payload = base.model_dump(by_alias=True)
    payload["optionBlocks"] = {
        "A": [{"id": "opt-a", "kind": "paragraph", "text": OPTION_A_TEXT}],
        "B": [{"id": "opt-b", "kind": "paragraph", "text": OPTION_B_TEXT}],
    }
    payload["answerBlocks"] = [{"id": "ans-1", "kind": "paragraph", "text": ANSWER_TEXT}]
    payload["explanationBlocks"] = [
        {"id": "exp-1", "kind": "paragraph", "text": EXPLANATION_TEXT}
    ]
    payload["stemBlocks"] = list(payload["stemBlocks"]) + [
        {"id": "f-latex", "kind": "formula", "latex": STANDALONE_LATEX, "ommlXml": None}
    ]
    return RichContentV2.model_validate(payload)


# --------------------------------------------------------------------------- 渲染产物助手


def package_entries(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def package_image_blobs(data: bytes) -> list[bytes]:
    return [value for name, value in package_entries(data).items() if name.startswith("word/media/")]


def document_xml(data: bytes) -> bytes:
    return package_entries(data)["word/document.xml"]


def embed_rids(data: bytes) -> list[str]:
    """渲染产物 ``word/document.xml`` 里出现的 r:embed / r:id 值。"""
    xml = document_xml(data).decode("utf-8")
    return re.findall(r'r:(?:embed|id)="([^"]+)"', xml)


def omml_elements(data: bytes) -> list:
    root = etree.fromstring(document_xml(data))
    return list(root.iter(f"{{{M_NAMESPACE}}}oMath"))


def omml_texts(data: bytes) -> list[str]:
    return [
        "".join(node.text or "" for node in element.iter(f"{{{M_NAMESPACE}}}t"))
        for element in omml_elements(data)
    ]


def distinct_cell_texts(document) -> list[str]:
    """文档里每个物理表格单元格的文本（合并单元格不重复计数）。"""
    from docx.oxml.ns import qn
    from docx.table import Table

    texts: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag != qn("w:tbl"):
            continue
        table = Table(child, document)
        seen: list = []
        for cell in table._cells:  # noqa: SLF001 - 网格中的重复引用需要按 tc 去重
            if any(cell._tc is existing._tc for existing in seen):  # noqa: SLF001
                continue
            seen.append(cell)
        texts.extend(cell.text for cell in seen)
    return texts


def document_text(document) -> str:
    """整篇可见文本（含表格单元格，合并单元格只算一次）。"""
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    pieces: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            pieces.append(Paragraph(child, document).text)
    pieces.extend(distinct_cell_texts(document))
    return "\n".join(pieces)
