"""富内容 v2 → DOCX 渲染（``student`` / ``teacher`` 两个变体）。

渲染口径（可测、不猜）：

- 学生版（``variant="student"``）**不含** ``answerBlocks`` 与 ``explanationBlocks``；
  教师版（``variant="teacher"``）两者都出（答案 → 解析，顺序固定）；
- ``sharedMaterials`` **只输出一次**，按出现顺序、放在正文之前：``RichContentV2`` 是**单题**
  视图，容器里没有"第一个引用它的题目"，因此统一前置一次，不做"每题重复输出"；
  同一块 id 在材料与题干/选项/答案之间重复出现一律 422（既不静默去重也不重复输出）；
- ``optionBlocks`` 按 key 顺序输出（key 只是结构信息，不额外打印标签，避免与题干里的
  选项文字重复）；
- **每次渲染都新建 ``Document()``**：所有 relationship/part 都属于新包，绝不复用源文档的
  rId 或 part；同一次渲染内同一资产共享一个 part（python-docx 按图片字节散列去重）；
- 图片：``ImageBlock.assetId`` 是受管键（``blobs/<sha256>``）时直接 ``AssetStore.read``；
  否则在 ``rich.assets`` 里按 ``assetId`` 找到 ``RichAsset``，再按 ``blobs/<sha256>`` 读。
  ``AssetStore.read`` 会重算散列，指纹不符一律 ``ASSET_CORRUPT``，不返回可疑内容；
- 公式：优先原始 ``ommlXml``（``parse_xml`` 原样插入，不改写、不做 LaTeX 往返）；只有
  ``latex`` 时经 ``formulas.omml_from_latex`` 转换，失败报 422 ``FORMULA_CONVERSION_FAILED``，
  消息与 ``details.fields`` 带块 id；
- 表格：``TableBlock.cells`` 是**行优先的扁平列表**，行边界由 ``columnCount``（网格宽度）
  给出。网格重建优先级：
  1. ``block.columnCount``（解析器一律填写，契约冻结字段）—— 直接采用；与单元格序列矛盾时
     报 422 ``INVALID_REQUEST``，不猜一个"看起来成功"的表格；
  2. 旧数据缺 ``columnCount``：先用 ``rich.origin.sourceLocator["tableGrids"][blockId]``
     （B1 早期约定的定位提示，可选），再用启发式 —— 表头行（``isHeader`` 起始连续段）的
     覆盖列宽、然后全部前缀和（降序），取第一个能完整铺成矩形网格且 rowSpan 不越界的列数；
  3. 合并单元格用 ``cell.merge`` 实现（横向 → ``w:gridSpan``，纵向 → ``w:vMerge``），
     表头单元格加粗。

已知限制：缺 ``columnCount`` 且无表头、无合并的表格，扁平列表无法唯一确定行数（启发式可能
与原表不一致）；DOCX 解析器因此强制填写 ``columnCount``，消费方也应优先使用它。
"""

from __future__ import annotations

import io
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from typing import Any, Literal

from app.contracts.teaching_loop import (
    ContentBlock,
    FormulaBlock,
    ImageBlock,
    ParagraphBlock,
    RichContentV2,
    TableBlock,
    TableCell,
    error_details,
)
from app.core.exceptions import AppError
from app.services.assets.store import AssetStore, is_managed_blob_key
from app.services.rich_content.formulas import MATH_NAMESPACE, omml_from_latex

RenderVariant = Literal["student", "teacher"]
RENDER_VARIANTS: tuple[str, ...] = ("student", "teacher")

#: 96 dpi 下 1 像素的 EMU（914400 / 96）
EMU_PER_PIXEL = 9525


def _invalid(message: str, *, fields: Sequence[str] | None = None) -> AppError:
    details = error_details(fields=list(fields)) if fields else None
    return AppError(message, code="INVALID_REQUEST", status_code=422, details=details)


def _math_tag(local_name: str) -> str:
    return f"{{{MATH_NAMESPACE}}}{local_name}"


@dataclass(frozen=True)
class _PlacedCell:
    """一次网格铺放：单元格落在 (row, column) 的左上角。"""

    cell: TableCell
    row: int
    column: int


# --------------------------------------------------------------------------- 表格重建


def _simulate_grid(cells: Sequence[TableCell], columns: int) -> list[list[_PlacedCell]] | None:
    """按列数把扁平单元格序列铺成矩形网格；不可行返回 None。"""
    if columns < 1:
        return None
    remaining = list(cells)
    rows: list[list[_PlacedCell]] = []
    #: 网格列 -> 后续仍需覆盖的行数（纵向合并的延续）
    pending: dict[int, int] = {}
    guard = len(cells) * 2 + sum(cell.row_span for cell in cells) + 4
    while (remaining or pending) and guard > 0:
        guard -= 1
        covered = set(pending)
        row: list[_PlacedCell] = []
        fresh: dict[int, int] = {}
        column = 0
        while column < columns:
            if column in covered:
                column += 1
                continue
            if not remaining:
                return None
            cell = remaining[0]
            col_span = max(1, cell.col_span)
            row_span = max(1, cell.row_span)
            if column + col_span > columns:
                return None
            if any(covered_column in covered for covered_column in range(column, column + col_span)):
                return None
            remaining.pop(0)
            row.append(_PlacedCell(cell=cell, row=len(rows), column=column))
            for covered_column in range(column, column + col_span):
                covered.add(covered_column)
                if row_span > 1:
                    fresh[covered_column] = max(
                        fresh.get(covered_column, 0), row_span - 1
                    )
            column += col_span
        if not row and not covered:
            return None
        rows.append(row)
        pending = {col: left - 1 for col, left in pending.items() if left - 1 > 0}
        for col, extra in fresh.items():
            pending[col] = max(pending.get(col, 0), extra)
    if guard <= 0 or remaining or pending:
        return None
    return rows


def _candidate_columns(cells: Sequence[TableCell], *, hint: int | None) -> list[int]:
    if hint is not None:
        return [hint]
    candidates: list[int] = []
    header_total = 0
    for cell in cells:
        if not cell.is_header:
            break
        header_total += max(1, cell.col_span)
    if header_total >= 1:
        candidates.append(header_total)
    prefix = 0
    prefixes: list[int] = []
    for cell in cells:
        prefix += max(1, cell.col_span)
        prefixes.append(prefix)
    candidates.extend(reversed(prefixes))
    unique: list[int] = []
    for value in candidates:
        if value >= 1 and value not in unique:
            unique.append(value)
    return unique


def _table_grid(
    block: TableBlock, *, hint: int | None
) -> tuple[int, list[list[_PlacedCell]]]:
    if not block.cells:
        raise _invalid(f"表格块 {block.id} 没有任何单元格，无法渲染。", fields=[block.id])
    # 1) 契约字段 columnCount 是精确依据；与单元格序列矛盾时直接报错（不猜）
    if block.column_count is not None:
        rows = _simulate_grid(block.cells, block.column_count)
        if rows:
            return block.column_count, rows
        raise _invalid(
            f"表格块 {block.id} 的 columnCount={block.column_count} 与单元格序列"
            "（rowSpan/colSpan）不一致，无法还原为矩形网格。",
            fields=[block.id],
        )
    # 2) 旧数据：定位提示 → 表头行覆盖列宽 → 前缀和（降序）启发式
    for columns in _candidate_columns(block.cells, hint=hint):
        rows = _simulate_grid(block.cells, columns)
        if rows:
            return columns, rows
    raise _invalid(
        f"表格块 {block.id} 的单元格结构无法还原为矩形网格"
        "（缺 columnCount，且各行覆盖列数与 rowSpan/colSpan 不一致）。",
        fields=[block.id],
    )


def _table_hint(rich: RichContentV2, block_id: str) -> int | None:
    locator = rich.origin.source_locator
    if not isinstance(locator, Mapping):
        return None
    grids = locator.get("tableGrids")
    if not isinstance(grids, Mapping):
        return None
    value = grids.get(block_id)
    if isinstance(value, bool) or not isinstance(value, int) or value < 1:
        return None
    return value


# --------------------------------------------------------------------------- 渲染


def _content_sections(content: RichContentV2, variant: str) -> list[list[ContentBlock]]:
    sections: list[list[ContentBlock]] = [
        [block for material in content.shared_materials for block in material.blocks],
        list(content.stem_blocks),
    ]
    for key in content.option_blocks:
        sections.append(list(content.option_blocks[key]))
    if variant == "teacher":
        sections.append(list(content.answer_blocks))
        sections.append(list(content.explanation_blocks))
    return sections


def _assert_unique_block_ids(content: RichContentV2) -> None:
    seen: set[str] = set()
    for section in _content_sections(content, "teacher"):
        for block in section:
            if block.id in seen:
                raise _invalid(
                    f"富内容里块 id 重复出现：{block.id}（同一块不能在多个位置输出）。",
                    fields=[block.id],
                )
            seen.add(block.id)


def _as_rich_content(value: Any) -> RichContentV2:
    if isinstance(value, RichContentV2):
        return value
    if isinstance(value, Mapping):
        try:
            return RichContentV2.model_validate(dict(value))
        except Exception as exc:  # noqa: BLE001 - 结构不合法一律 422，不返回假成功
            raise _invalid(f"富内容结构不合法：{exc.__class__.__name__}。") from exc
    raise _invalid(f"rich 必须是 RichContentV2，收到 {type(value).__name__}。")


class _Renderer:
    def __init__(self, *, document: Any, assets: AssetStore, rich: RichContentV2) -> None:
        self.document = document
        self.assets = assets
        self.rich = rich

    def render_blocks(self, blocks: Sequence[ContentBlock]) -> None:
        for block in blocks:
            if isinstance(block, ParagraphBlock):
                self.document.add_paragraph(block.text)
            elif isinstance(block, TableBlock):
                self.render_table(block)
            elif isinstance(block, FormulaBlock):
                self.render_formula(block)
            elif isinstance(block, ImageBlock):
                self.render_image(block)
            else:
                raise _invalid(f"不支持的内容块类型：{type(block).__name__}。")

    # -------------------------------------------------------------- 表格

    def render_table(self, block: TableBlock) -> None:
        columns, rows = _table_grid(block, hint=_table_hint(self.rich, block.id))
        table = self.document.add_table(rows=len(rows), cols=columns)
        try:
            table.style = "Table Grid"
        except (KeyError, ValueError):  # 样式缺失不影响结构正确性
            pass
        for row in rows:
            for placed in row:
                cell = table.cell(placed.row, placed.column)
                cell.text = placed.cell.text
                if placed.cell.is_header:
                    self._bold_cell(cell)
                col_span = max(1, placed.cell.col_span)
                row_span = max(1, placed.cell.row_span)
                if col_span > 1 or row_span > 1:
                    cell.merge(
                        table.cell(
                            placed.row + row_span - 1, placed.column + col_span - 1
                        )
                    )

    @staticmethod
    def _bold_cell(cell: Any) -> None:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    # -------------------------------------------------------------- 公式

    def render_formula(self, block: FormulaBlock) -> None:
        from docx.oxml import parse_xml

        omml_xml = (block.omml_xml or "").strip()
        latex = (block.latex or "").strip()
        if not omml_xml and not latex:
            raise _invalid(
                f"公式块 {block.id} 既没有 ommlXml 也没有 latex。", fields=[block.id]
            )
        if omml_xml:
            try:
                element = parse_xml(omml_xml)
            except Exception as exc:  # noqa: BLE001 - 存量 OMML 损坏必须可见
                raise _invalid(
                    f"公式块 {block.id} 的 ommlXml 不是合法 XML。", fields=[block.id]
                ) from exc
            element = self._require_math_element(element, block_id=block.id)
        else:
            try:
                converted = omml_from_latex(latex)
            except AppError as exc:
                raise AppError(
                    f"公式块 {block.id} 的 LaTeX 转换失败：{exc}",
                    code="FORMULA_CONVERSION_FAILED",
                    status_code=422,
                    details=error_details(fields=[block.id]),
                ) from exc
            try:
                element = parse_xml(converted)
            except Exception as exc:  # noqa: BLE001 - 转换器产出异常必须可见
                raise _invalid(
                    f"公式块 {block.id} 的 LaTeX 转换结果不是合法 XML。", fields=[block.id]
                ) from exc
        paragraph = self.document.add_paragraph()
        paragraph._p.append(element)

    @staticmethod
    def _require_math_element(element: Any, *, block_id: str) -> Any:
        if element.tag == _math_tag("oMath"):
            return element
        if element.tag == _math_tag("oMathPara"):
            children = [child for child in element if child.tag == _math_tag("oMath")]
            if len(children) == 1:
                return children[0]
        raise _invalid(
            f"公式块 {block_id} 的 ommlXml 根节点不是 m:oMath。", fields=[block_id]
        )

    # -------------------------------------------------------------- 图片

    def render_image(self, block: ImageBlock) -> None:
        from docx.shared import Emu

        blob_key = self._resolve_blob_key(block)
        data = self.assets.read(blob_key)
        stream = io.BytesIO(data)
        width = Emu(block.width * EMU_PER_PIXEL) if block.width > 0 else None
        height = Emu(block.height * EMU_PER_PIXEL) if block.height > 0 else None
        try:
            self.document.add_picture(stream, width=width, height=height)
        except Exception as exc:  # noqa: BLE001 - 不支持的图片格式必须可见
            raise AppError(
                f"图片块 {block.id} 写入失败：{exc.__class__.__name__}（不支持的图片格式或数据损坏）。",
                code="RICH_IMAGE_UNSUPPORTED",
                status_code=422,
                details=error_details(fields=[block.id]),
            ) from exc

    def _resolve_blob_key(self, block: ImageBlock) -> str:
        if is_managed_blob_key(block.asset_id):
            return block.asset_id
        for asset in self.rich.assets:
            if asset.asset_id == block.asset_id:
                return f"blobs/{asset.sha256}"
        raise AppError(
            f"图片块 {block.id} 引用的资产无法解析为受管键：{block.asset_id}。",
            code="ASSET_NOT_FOUND",
            status_code=422,
            details=error_details(fields=[block.id]),
        )


def render_rich_document(
    *,
    rich: RichContentV2,
    assets: AssetStore,
    variant: RenderVariant,
    title: str | None = None,
) -> bytes:
    """把富内容渲染成新的 ``.docx`` 字节；产物自洽，可直接用 python-docx 重新打开。

    - 每次调用都新建文档与 relationship，不复用任何源文档的 rId/part；
    - 学生版跳过 ``answerBlocks`` / ``explanationBlocks``；教师版全部输出；
    - 失败一律抛 ``AppError``（422/500），不返回"看起来成功"的空文档。
    """
    if variant not in RENDER_VARIANTS:
        raise _invalid(f"variant 必须是 {list(RENDER_VARIANTS)} 之一，收到：{variant!r}。")
    content = _as_rich_content(rich)
    if not callable(getattr(assets, "read", None)):
        raise _invalid("assets 必须是 AssetStore（需要 read）。")
    _assert_unique_block_ids(content)

    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - 依赖缺失时明确失败
        raise AppError(
            "DOCX 渲染依赖 python-docx 未安装。",
            code="DOCUMENT_PARSER_UNAVAILABLE",
            status_code=500,
        ) from exc

    document = Document()
    heading = (title or "").strip()
    if heading:
        document.add_heading(heading, level=1)
    renderer = _Renderer(document=document, assets=assets, rich=content)
    for section in _content_sections(content, variant):
        renderer.render_blocks(section)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


__all__ = ["EMU_PER_PIXEL", "RENDER_VARIANTS", "RenderVariant", "render_rich_document"]
