"""Practice export composition only; DOCX rendering stays in T10."""
import copy
import io
import zipfile
from app.contracts.teaching_loop import RichContentV2, canonical_hash
from app.services.rich_content.renderer_docx import render_rich_document
from .common import invalid

DOCX_MEDIA = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _identity(blocks, assets):
    result=[]
    declarations={x["assetId"]:x for x in assets}
    for block in blocks:
        value={k:v for k,v in block.items() if k!="id"}
        if block["kind"]=="image":
            declaration=declarations[block["assetId"]]
            value.pop("assetId")
            value.update(sha256=declaration["sha256"],mediaType=declaration["mediaType"])
        result.append(value)
    return result


def combine(selections, title):
    result = dict(version=2, sharedMaterials=[], stemBlocks=[], optionBlocks={}, answerBlocks=[], explanationBlocks=[], assets=[],
                  origin=copy.deepcopy(selections[0]["rich"]["origin"]))
    materials = set()
    declared = {}
    for index, selection in enumerate(selections, 1):
        rich = selection["rich"]
        def namespace(blocks, section):
            return [dict(copy.deepcopy(block), id=f"selection-{index}:{section}:{n}") for n, block in enumerate(blocks)]
        for material in rich["sharedMaterials"]:
            identity = canonical_hash(_identity(material["blocks"],rich["assets"]))
            if identity not in materials:
                materials.add(identity)
                result["sharedMaterials"].append(dict(id=identity, blocks=namespace(material["blocks"], "material:"+identity)))
        score = selection['maxScoreUnits']
        number=selection["ordinal"]
        result["stemBlocks"].append(dict(id=f"title-{index}", kind="paragraph", text=f"题组 {number}（整题满分 {score//100}.{score%100:02d}）"))
        for n,node in enumerate(selection["nodes"]):
            score_text = "不计分" if not node["isScored"] else f"满分 {node['maxScoreUnits']//100}.{node['maxScoreUnits']%100:02d}"
            result["stemBlocks"].append(dict(id=f"node-{index}-{n}",kind="paragraph",text=f"题号 {node['questionNo']}：{score_text}"))
        result["stemBlocks"].extend(namespace(rich["stemBlocks"], "stem"))
        for key, blocks in rich["optionBlocks"].items():
            result["stemBlocks"].append(dict(id=f"option-title-{index}-{key}", kind="paragraph", text=f"{key}."))
            result["stemBlocks"].extend(namespace(blocks, "option:"+key))
        result["answerBlocks"].append(dict(id=f"answer-title-{index}", kind="paragraph",
            text=f"题组 {number} 答案" if rich["answerBlocks"] else f"题组 {number} 未提供答案"))
        result["answerBlocks"].extend(namespace(rich["answerBlocks"], "answer"))
        if rich["explanationBlocks"]:
            result["explanationBlocks"].append(dict(id=f"explanation-title-{index}", kind="paragraph", text=f"题组 {number} 解析"))
            result["explanationBlocks"].extend(namespace(rich["explanationBlocks"], "explanation"))
        for asset in rich["assets"]:
            if asset["assetId"] in declared and declared[asset["assetId"]] != asset:
                raise invalid("同一图片资产声明冲突。", "assets")
            declared[asset["assetId"]] = asset
    result["assets"] = list(declared.values())
    return RichContentV2.model_validate(result)


def render(frozen, assets):
    if frozen["variant"] == "score_template":
        return template(frozen), XLSX_MEDIA, "成绩模板.xlsx"
    rich = combine(frozen["selections"], frozen["title"])
    payload = render_rich_document(rich=rich, assets=assets, variant=frozen["variant"], title=frozen["title"])
    # Verify the final package is complete; tests independently inspect every part for leakage.
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        if archive.testzip() is not None:
            raise invalid("导出 ZIP 校验失败。", "variant", "PRACTICE_EXPORT_INVALID")
    return payload, DOCX_MEDIA, ("学生练习.docx" if frozen["variant"] == "student" else "教师练习.docx")


def template(frozen):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "成绩"
    leaves = frozen["leaves"]
    sheet.append(["学号", "姓名", "出勤", "人次", *[x["questionNo"] for x in leaves]])
    for participant in frozen["participants"]:
        values = [participant["studentNoSnapshot"] or "", participant["nameSnapshot"], participant["attendance"], participant["attemptNo"]]
        sheet.append([*values, *[None for _ in leaves]])
        # Formula-looking names/numbers are text as well as leading-zero identifiers.
        for column in (1, 2):
            cell = sheet.cell(sheet.max_row, column)
            cell.data_type = "s"
            cell.number_format = "@"
    meta = workbook.create_sheet("固定映射")
    for field in ("practiceRevisionId", "assessmentId", "paperRevisionId", "frozenAt"):
        meta.append([field, frozen[field]])
    meta.append(["列", "paperItemId", "完整题号", "满分整数单位"])
    for n, leaf in enumerate(leaves, 5):
        meta.append([n, leaf["itemId"], leaf["questionNo"], leaf["maxScoreUnits"]])
    meta.append(["参与者快照", "participantId", "studentId", "classId", "attemptNo", "attendance"])
    for p in frozen["participants"]:
        meta.append(["", p["participantId"], p["studentId"], p["classId"], p["attemptNo"], p["attendance"]])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
