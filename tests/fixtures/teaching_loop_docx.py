"""固定浏览器原卷样本：真实DOCX原文、共同材料、公式、合并表格和图片。"""
from io import BytesIO
from docx import Document
from docx.oxml import parse_xml
from docx.shared import Inches
from tests.rich_content_support import minimal_png

def build_complete_paper() -> bytes:
    document = Document()
    document.add_paragraph('阅读下面材料，回答第1至3题：保留每个小题的原始得分。')
    document.add_paragraph('1. （2分）计算 1 + 1。')
    formula = document.add_paragraph()
    formula._p.append(parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x+1</m:t></m:r></m:oMath>'))
    document.add_paragraph('2. （3分）比较两个有理数。')
    document.add_paragraph('3. （5分）根据表格与图片完成计算。')
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = '原文合并表头'
    table.cell(1, 0).text = '甲'; table.cell(1, 1).text = '2'
    table.cell(2, 0).text = '乙'; table.cell(2, 1).text = '3'
    document.add_picture(BytesIO(minimal_png()), width=Inches(0.5))
    stream = BytesIO(); document.save(stream)
    return stream.getvalue()
